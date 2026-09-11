"""Question 3 endpoint solver and audit pipeline.

Question 3 inherits the fixed-radius, variable-property coupled model from
``q2_solver``.  This module owns the long-time boundary extension, monitors
the complete radial state at accepted BDF steps and dense midpoints, locates
the first moisture-threshold crossing, writes the strict 60-second result
workbook, and keeps all numerical evidence in ``results/q3``.

The implementation deliberately does not modify the question 2 solver or its
results.  The public q2 functions are imported read-only for geometry,
properties, the conservative RHS, and the analytic log(C) Jacobian.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import platform
import shutil
import sys
import time
from copy import copy as copy_style
from pathlib import Path
from typing import Any, Callable

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.optimize import brentq

import q2_solver as q2


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BOUNDARY = ROOT.parent / "附件1.xlsx"
DEFAULT_TEMPLATE = ROOT.parent / "A题" / "A题" / "附件" / "附件3" / "result3.xlsx"
OFFICIAL_RESULT3 = DEFAULT_TEMPLATE
DEFAULT_OUTPUT = ROOT / "results" / "q3"
Q2_RESULTS = ROOT / "results" / "q2"
Q2_RESTART = Q2_RESULTS / "q2_restart.npz"
Q2_FIELDS = Q2_RESULTS / "q2_fields.npz"
Q2_MASS_BALANCE = Q2_RESULTS / "q2_mass_balance.csv"

THRESHOLD_C = 0.15
NOMINAL_PLATEAU_T_C = 50.0
NOMINAL_PLATEAU_C = 0.05
OUTPUT_STEP_S = 60
PAPER_STEP_S = 6 * 3600
PLATEAU_START_S = 14400
DEFAULT_START_S = 0
DEFAULT_MAX_END_S = 168 * 3600
DEFAULT_CHUNK_S = 600
EVENT_TIME_TOL_S = 1.0
MONOTONICITY_TOL_C = 1.0e-8
CONVERGENCE_REL_TIME = 1.0e-4
CONVERGENCE_ABS_TIME_S = 30.0
FIELD_THRESHOLD_T_C = 5.0e-5
FIELD_THRESHOLD_C = 5.0e-5
MASS_BALANCE_REL_TOL = 1.0e-4
EVENT_CONCENTRATION_TOL = 1.0e-9
PLATFORM_T_STD_C = 0.5
PLATFORM_C_STD = 1.0e-3


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    return value


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(_json_safe(value), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def platform_diagnostic(
    boundary: np.ndarray,
    start_s: float = 10800.0,
    end_s: float = 14400.0,
) -> dict[str, Any]:
    """Diagnose whether the final hour is a usable environment plateau.

    The slope is reported in physical units per hour.  The tests are explicit
    engineering screening rules, not a statistical proof about future days.
    """
    boundary = np.asarray(boundary, dtype=float)
    mask = (boundary[:, 0] >= start_s - 1e-10) & (boundary[:, 0] <= end_s + 1e-10)
    tail = boundary[mask]
    if len(tail) < 3:
        raise ValueError("At least three final-hour boundary samples are required.")
    if not np.isclose(tail[0, 0], start_s) or not np.isclose(tail[-1, 0], end_s):
        raise ValueError("Final-hour boundary interval is incomplete.")
    t_hour = (tail[:, 0] - start_s) / 3600.0
    rows: dict[str, Any] = {}
    checks: dict[str, bool] = {}
    limits = {
        "temperature_C": PLATFORM_T_STD_C,
        "moisture_kg_per_kg": PLATFORM_C_STD,
    }
    for index, key in ((1, "temperature_C"), (2, "moisture_kg_per_kg")):
        values = tail[:, index]
        mean = float(np.mean(values))
        std = float(np.std(values, ddof=1))
        slope_per_hour = float(np.polyfit(t_hour, values, 1)[0])
        jump = float(values[-1] - mean)
        slope_ok = abs(slope_per_hour) <= std
        std_ok = std <= limits[key]
        jump_ok = abs(jump) <= 2.0 * std
        checks[f"{key}.std"] = std_ok
        checks[f"{key}.slope"] = slope_ok
        checks[f"{key}.jump"] = jump_ok
        rows[key] = {
            "sample_count": int(len(values)),
            "mean": mean,
            "sample_std": std,
            "linear_slope_per_hour": slope_per_hour,
            "last_value": float(values[-1]),
            "last_minus_mean": jump,
            "std_limit": float(limits[key]),
            "slope_limit_abs": std,
            "jump_limit_abs": float(2.0 * std),
            "pass_std": std_ok,
            "pass_slope": slope_ok,
            "pass_jump": jump_ok,
        }
    return {
        "interval_s": [float(start_s), float(end_s)],
        "sample_count": int(len(tail)),
        "checks": checks,
        "pass": bool(all(checks.values())),
        "temperature_C": rows["temperature_C"],
        "moisture_kg_per_kg": rows["moisture_kg_per_kg"],
        "decision": (
            "platform supports the stated nominal post-14400 s boundary; final-hour means are retained for sensitivity"
            if all(checks.values())
            else "do not publish the plateau extension; boundary model review required"
        ),
        "interpretation": "Engineering screening only; not a proof of future multi-day stationarity.",
    }


def make_plateau_boundary(
    diagnostic: dict[str, Any],
    max_end_s: int,
    *,
    temperature_C: float = NOMINAL_PLATEAU_T_C,
    moisture_kg_per_kg: float = NOMINAL_PLATEAU_C,
) -> np.ndarray:
    """Return a constant post-4-hour boundary after the platform check.

    The nominal Q3 continuation is the stated stable setting (50 C,
    0.05 kg/kg).  The final-hour mean is retained as evidence of platform
    behavior and is used only by the explicit boundary sensitivity case.
    """
    if max_end_s <= PLATEAU_START_S:
        raise ValueError("max_end_s must be later than the 4-hour boundary switch.")
    if not diagnostic.get("pass", False):
        raise ValueError("Cannot build a baseline plateau boundary after a failed diagnostic.")
    if not np.isfinite(temperature_C) or temperature_C <= -273.15:
        raise ValueError("Plateau temperature must be a finite physical value.")
    if not np.isfinite(moisture_kg_per_kg) or moisture_kg_per_kg < 0.0:
        raise ValueError("Plateau moisture must be a finite nonnegative value.")
    return np.array(
        [
            [float(PLATEAU_START_S), float(temperature_C), float(moisture_kg_per_kg)],
            [float(max_end_s), float(temperature_C), float(moisture_kg_per_kg)],
        ],
        dtype=float,
    )


def interleaved_state(temperature_C: np.ndarray, moisture: np.ndarray) -> np.ndarray:
    temperature_C = np.asarray(temperature_C, dtype=float)
    moisture = np.asarray(moisture, dtype=float)
    if temperature_C.shape != moisture.shape or temperature_C.ndim != 1:
        raise ValueError("Temperature and moisture restart fields must be equal-length vectors.")
    state = np.empty(2 * len(temperature_C), dtype=float)
    state[0::2] = temperature_C
    state[1::2] = moisture
    return state


def load_q2_restart(path: str | Path = Q2_RESTART) -> dict[str, Any]:
    """Read and validate the question 2 full-grid restart as a read-only input."""
    path = Path(path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Question 2 restart does not exist: {path}")
    with np.load(path, allow_pickle=False) as payload:
        required = {
            "final_time_s",
            "radius_nodes_m",
            "control_volume_weights_m2",
            "temperature_final_C",
            "moisture_final_kg_per_kg",
            "n_intervals",
            "h_W_m2K",
            "hm_m_s",
        }
        missing = required.difference(payload.files)
        if missing:
            raise ValueError(f"Question 2 restart misses fields: {sorted(missing)}")
        result = {key: payload[key].copy() for key in payload.files}
    n = int(result["n_intervals"])
    if len(result["radius_nodes_m"]) != n + 1:
        raise ValueError("Question 2 restart radius and n_intervals disagree.")
    if len(result["temperature_final_C"]) != n + 1 or len(result["moisture_final_kg_per_kg"]) != n + 1:
        raise ValueError("Question 2 restart fields do not match n_intervals.")
    q2.material_properties(result["moisture_final_kg_per_kg"], result["temperature_final_C"])
    if np.any(result["moisture_final_kg_per_kg"] <= 0.0):
        raise ValueError("Question 2 restart contains nonpositive moisture.")
    return result


def _physical_from_log(z_state: np.ndarray, n: int) -> tuple[np.ndarray, np.ndarray]:
    z_state = np.asarray(z_state, dtype=float)
    with np.errstate(over="raise", invalid="raise", under="ignore"):
        try:
            moisture = np.exp(z_state[1::2])
        except FloatingPointError as exc:
            raise q2.Q2DomainError("log(C) trial overflowed in Q3.") from exc
    if not np.isfinite(moisture).all() or np.any(moisture <= 0.0):
        raise q2.Q2DomainError("Q3 transformed state did not map to positive C.")
    physical = np.asarray(z_state, dtype=float).copy()
    physical[1::2] = moisture
    if not np.isfinite(physical[0::2]).all() or np.any(physical[0::2] + 273.15 <= 0.0):
        raise q2.Q2DomainError("Q3 transformed state has nonphysical temperature.")
    return physical, moisture


def _to_log_state(state: np.ndarray, n: int) -> np.ndarray:
    state = np.asarray(state, dtype=float).copy()
    if state.size != 2 * (n + 1):
        raise ValueError("State size and grid disagree.")
    q2.material_properties(state[1::2], state[0::2])
    with np.errstate(divide="raise", invalid="raise"):
        try:
            state[1::2] = np.log(state[1::2])
        except FloatingPointError as exc:
            raise q2.Q2DomainError("Restart moisture cannot be represented in log(C).") from exc
    return state


def _sample_state(
    z_values: np.ndarray,
    n: int,
    volumes: np.ndarray,
    sample_nodes: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Convert a matrix of log states to sampled fields and audit quantities."""
    if z_values.ndim != 2:
        raise ValueError("Expected a state matrix with shape (state, time).")
    physical, moisture = _physical_from_log(z_values, n)
    temperature = physical[0::2, :].T
    moisture_t = moisture.T
    mean = moisture_t @ volumes / (q2.RADIUS_M * q2.RADIUS_M / 2.0)
    max_c = np.max(moisture_t, axis=1)
    center_gap = max_c - moisture_t[:, 0]
    monotonicity = np.max(np.diff(moisture_t, axis=1), axis=1)
    return (
        temperature[:, sample_nodes],
        moisture_t[:, sample_nodes],
        np.column_stack([mean, max_c, center_gap, monotonicity]),
        physical,
    )


def _locate_crossing(
    dense: Callable[[np.ndarray], np.ndarray],
    lo: float,
    hi: float,
    n: int,
    threshold: float,
    use_center: bool = False,
    tol_s: float = EVENT_TIME_TOL_S,
) -> tuple[float, np.ndarray, dict[str, Any]]:
    """Bisection on a dense BDF interpolant for a threshold crossing."""
    def value(t: float) -> float:
        physical, moisture = _physical_from_log(np.asarray(dense(np.array([t]))).reshape(-1), n)
        if use_center:
            return float(moisture[0] - threshold)
        return float(np.max(moisture) - threshold)

    vlo = value(lo)
    vhi = value(hi)
    if vlo < 0.0 or vhi > 0.0:
        raise ValueError("Crossing bracket does not have the required downward sign change.")
    while hi - lo > tol_s:
        mid = 0.5 * (lo + hi)
        vmid = value(mid)
        if vmid > 0.0:
            lo = mid
        else:
            hi = mid
    # Keep the <=1 s bracket as an audit item, then solve the same dense event
    # function to a concentration residual.  This reports the critical root,
    # while strict-before/after evidence below handles the problem's strict
    # inequality wording.
    event_time = float(brentq(value, lo, hi, xtol=1.0e-8, rtol=1.0e-12))
    event_log_state = np.asarray(dense(np.array([event_time]))).reshape(-1)
    event_physical, event_moisture = _physical_from_log(event_log_state, n)
    return event_time, event_physical, {
        "bracket_s": [float(lo), float(hi)],
        "bracket_width_s": float(hi - lo),
        "center_C": float(event_moisture[0]),
        "max_C": float(np.max(event_moisture)),
        "center_minus_max_C": float(event_moisture[0] - np.max(event_moisture)),
    }


def integrate_segment(
    state: np.ndarray,
    n: int,
    start_s: int,
    end_s: int,
    boundary: np.ndarray,
    *,
    rtol: float,
    atol: float,
    max_step: float,
    chunk_seconds: int = DEFAULT_CHUNK_S,
    output_step_s: int = OUTPUT_STEP_S,
    detect_event: bool = False,
    event_kind: str = "grid_max",
    h: float = q2.H_BASE,
    hm: float = q2.HM_BASE,
) -> dict[str, Any]:
    """Integrate one continuous boundary segment with full-grid audits."""
    if end_s <= start_s:
        raise ValueError("Segment end must be later than segment start.")
    if boundary[0, 0] > start_s or boundary[-1, 0] < end_s:
        raise ValueError("Segment boundary does not cover requested interval.")
    if event_kind not in {"grid_max", "center"}:
        raise ValueError("event_kind must be 'grid_max' or 'center'.")
    rhs_log, jac_log, geometry = q2.make_log_rhs(boundary, n, h=h, hm=hm)
    dr, nodes, faces, volumes = geometry
    sample_nodes = np.arange(0, n + 1, n // 20, dtype=int)
    if len(sample_nodes) != 21 or sample_nodes[-1] != n:
        raise ArithmeticError("Q3 requires exactly 21 output nodes.")
    transformed_state = _to_log_state(state, n)
    output_times_global = np.arange(start_s, end_s + 1, output_step_s, dtype=int)
    out_times: list[int] = []
    out_T: list[np.ndarray] = []
    out_C: list[np.ndarray] = []
    out_audit: list[np.ndarray] = []
    flux_times: list[np.ndarray] = []
    flux_rates: list[np.ndarray] = []
    monitor = {
        "max_positive_neighbor_jump_kg_per_kg": -math.inf,
        "max_center_gap_kg_per_kg": -math.inf,
        "min_moisture_kg_per_kg": math.inf,
        "max_moisture_kg_per_kg": -math.inf,
        "checked_accepted_states": 0,
        "checked_midpoint_states": 0,
        "checked_output_states": 0,
        "segment_records": [],
    }
    nfev = njev = nlu = 0
    positivity_retries = 0
    cursor = float(start_s)
    final_state = np.asarray(state, dtype=float).copy()
    initial_mean = float(np.dot(state[1::2], volumes) / (q2.RADIUS_M * q2.RADIUS_M / 2.0))
    event: dict[str, Any] | None = None
    started = time.perf_counter()

    def threshold_event(t: float, z_state: np.ndarray) -> float:
        _physical, moisture = _physical_from_log(z_state, n)
        control = moisture[0] if event_kind == "center" else np.max(moisture)
        return float(control - THRESHOLD_C)

    threshold_event.terminal = True  # type: ignore[attr-defined]
    threshold_event.direction = -1.0  # type: ignore[attr-defined]

    while cursor < float(end_s) - 1e-12:
        segment_end = min(cursor + float(chunk_seconds), float(end_s))
        step_used = float(max_step)
        while True:
            try:
                solution = solve_ivp(
                    rhs_log,
                    (cursor, segment_end),
                    transformed_state,
                    method="BDF",
                    jac=jac_log,
                    jac_sparsity=q2.jacobian_sparsity(n),
                    rtol=rtol,
                    atol=atol,
                    max_step=step_used,
                    dense_output=True,
                    events=threshold_event if detect_event else None,
                )
                break
            except q2.Q2DomainError:
                if step_used <= max_step / 64.0:
                    raise
                step_used /= 2.0
                positivity_retries += 1
        if not solution.success:
            raise RuntimeError(f"BDF failed on Q3 segment {cursor}--{segment_end}: {solution.message}")
        nfev += solution.nfev
        njev += solution.njev or 0
        nlu += solution.nlu or 0
        actual_end = float(solution.t[-1])
        _accepted_physical_for_flux, accepted_moisture_for_flux = _physical_from_log(solution.y, n)
        environment_moisture_for_flux = np.interp(solution.t, boundary[:, 0], boundary[:, 2])
        flux_times.append(np.asarray(solution.t, dtype=float))
        flux_rates.append(q2.RADIUS_M * hm * (accepted_moisture_for_flux[-1, :] - environment_moisture_for_flux))

        # Audit accepted states and dense midpoint states before extracting
        # only the 21 requested output radii.
        _, _, accepted_audit, accepted_physical = _sample_state(solution.y, n, volumes, sample_nodes)
        if len(solution.t) > 1:
            midpoint_times = 0.5 * (solution.t[:-1] + solution.t[1:])
            midpoint_states = solution.sol(midpoint_times)
            _, _, midpoint_audit, midpoint_physical = _sample_state(midpoint_states, n, volumes, sample_nodes)
        else:
            midpoint_audit = np.empty((0, 4), dtype=float)
            midpoint_physical = np.empty((2 * (n + 1), 0), dtype=float)
        for kind, audit_values in (("accepted", accepted_audit), ("midpoint", midpoint_audit)):
            if len(audit_values) == 0:
                continue
            monitor["max_positive_neighbor_jump_kg_per_kg"] = max(
                float(monitor["max_positive_neighbor_jump_kg_per_kg"]),
                float(np.max(audit_values[:, 3])),
            )
            monitor["max_center_gap_kg_per_kg"] = max(
                float(monitor["max_center_gap_kg_per_kg"]),
                float(np.max(audit_values[:, 2])),
            )
            physical_values = accepted_physical if kind == "accepted" else midpoint_physical
            moisture_values = physical_values[1::2, :]
            monitor["min_moisture_kg_per_kg"] = min(float(monitor["min_moisture_kg_per_kg"]), float(np.min(moisture_values)))
            monitor["max_moisture_kg_per_kg"] = max(float(monitor["max_moisture_kg_per_kg"]), float(np.max(moisture_values)))
            monitor[f"checked_{kind}_states"] += int(len(audit_values))
        monitor["segment_records"].append(
            {
                "start_s": float(cursor),
                "requested_end_s": float(segment_end),
                "actual_end_s": actual_end,
                "accepted_states": int(len(solution.t)),
                "midpoint_states": int(len(midpoint_audit)),
                "max_step_used_s": float(step_used),
                "event_detected": bool(solution.t_events and len(solution.t_events[0]) > 0),
            }
        )

        # Save requested 60-second output points reached in this solve.
        eligible = output_times_global[(output_times_global >= math.ceil(cursor - 1e-8)) & (output_times_global <= math.floor(actual_end + 1e-8))]
        if len(eligible):
            z_values = solution.sol(eligible.astype(float))
            T_sample, C_sample, audit_values, _ = _sample_state(z_values, n, volumes, sample_nodes)
            for index, tt in enumerate(eligible):
                if out_times and int(tt) <= out_times[-1]:
                    continue
                out_times.append(int(tt))
                out_T.append(T_sample[index].copy())
                out_C.append(C_sample[index].copy())
                out_audit.append(audit_values[index].copy())
            if len(audit_values):
                monitor["max_positive_neighbor_jump_kg_per_kg"] = max(
                    float(monitor["max_positive_neighbor_jump_kg_per_kg"]),
                    float(np.max(audit_values[:, 3])),
                )
                monitor["max_center_gap_kg_per_kg"] = max(
                    float(monitor["max_center_gap_kg_per_kg"]),
                    float(np.max(audit_values[:, 2])),
                )
                monitor["checked_output_states"] += int(len(audit_values))

        if detect_event and solution.t_events and len(solution.t_events[0]) > 0:
            raw_event_time = float(solution.t_events[0][0])
            # Refine the event with a one-second-or-smaller bracket on the
            # dense BDF interpolant.  The solver event itself is retained as
            # a diagnostic, but the stored state/time use this explicit root.
            bracket_lo = float(solution.t[-2]) if len(solution.t) >= 2 else float(cursor)
            while bracket_lo < raw_event_time - 1e-12:
                try:
                    _grid_lo_state, _grid_lo_moisture = _physical_from_log(np.asarray(solution.sol(np.array([bracket_lo]))).reshape(-1), n)
                    grid_lo_value = float(np.max(_grid_lo_moisture) - THRESHOLD_C)
                    if grid_lo_value >= 0.0:
                        break
                except Exception:
                    pass
                earlier = np.where(solution.t < bracket_lo - 1e-12)[0]
                if len(earlier) == 0:
                    break
                bracket_lo = float(solution.t[earlier[-1]])
            try:
                grid_time, event_physical, grid_meta = _locate_crossing(
                    solution.sol,
                    bracket_lo,
                    raw_event_time,
                    n,
                    THRESHOLD_C,
                    use_center=False,
                    tol_s=EVENT_TIME_TOL_S,
                )
                event_time = float(grid_time)
                event_log_state = np.asarray(solution.sol(np.array([event_time]))).reshape(-1)
            except ValueError:
                event_time = raw_event_time
                event_log_state = np.asarray(solution.y_events[0][0]).reshape(-1)
                event_physical, _ = _physical_from_log(event_log_state, n)
                grid_meta = {"bracket_s": [bracket_lo, raw_event_time], "bracket_width_s": raw_event_time - bracket_lo}
            event_physical, event_moisture = _physical_from_log(event_log_state, n)
            center_meta: dict[str, Any] | None = None
            try:
                center_time, _center_state, center_meta = _locate_crossing(
                    solution.sol,
                    bracket_lo,
                    raw_event_time,
                    n,
                    THRESHOLD_C,
                    use_center=True,
                    tol_s=EVENT_TIME_TOL_S,
                )
                center_meta["time_s"] = float(center_time)
            except ValueError:
                center_meta = None
            # Continue for 0.5 s after the exact event so strict-before and
            # strict-after evidence is based on the same numerical model.
            strict_before_time = max(float(cursor), event_time - 0.5)
            strict_after_time = event_time + 0.5
            before_state = np.asarray(solution.sol(np.array([strict_before_time]))).reshape(-1)
            try:
                after_solution = solve_ivp(
                    rhs_log,
                    (event_time, strict_after_time),
                    event_log_state,
                    method="BDF",
                    jac=jac_log,
                    jac_sparsity=q2.jacobian_sparsity(n),
                    rtol=rtol,
                    atol=atol,
                    max_step=min(step_used, 0.5),
                    dense_output=True,
                )
                if not after_solution.success:
                    raise RuntimeError(after_solution.message)
                after_state = np.asarray(after_solution.y[:, -1]).reshape(-1)
            except Exception:
                after_state = event_log_state.copy()
            _before_physical, before_moisture = _physical_from_log(before_state, n)
            _after_physical, after_moisture = _physical_from_log(after_state, n)
            event = {
                "time_s": event_time,
                "state": event_physical,
                "center_C": float(event_moisture[0]),
                "max_C": float(np.max(event_moisture)),
                "center_minus_max_C": float(event_moisture[0] - np.max(event_moisture)),
                "event_residual_C": float(np.max(event_moisture) - THRESHOLD_C),
                "event_kind": event_kind,
                "grid_event_time_s": float(event_time),
                "grid_event_bracket_s": grid_meta.get("bracket_s"),
                "grid_event_bracket_width_s": float(grid_meta.get("bracket_width_s", math.inf)),
                "center_event_time_s": None if center_meta is None else float(center_meta["time_s"]),
                "center_event_bracket_s": None if center_meta is None else center_meta.get("bracket_s"),
                "center_event_bracket_width_s": None if center_meta is None else float(center_meta.get("bracket_width_s", math.inf)),
                "strict_before_time_s": float(strict_before_time),
                "strict_after_time_s": float(strict_after_time),
                "strict_before_max_C": float(np.max(before_moisture)),
                "strict_after_max_C": float(np.max(after_moisture)),
                "strict_before_pass": bool(np.max(before_moisture) > THRESHOLD_C),
                "strict_after_pass": bool(np.max(after_moisture) < THRESHOLD_C),
            }
            final_state = event_physical
            cursor = event_time
            transformed_state = event_log_state.copy()
            break

        transformed_state = solution.y[:, -1].copy()
        final_state, _ = _physical_from_log(transformed_state, n)
        cursor = actual_end

    return {
        "start_s": int(start_s),
        "end_s_requested": int(end_s),
        "final_time_s": float(event["time_s"] if event else end_s),
        "state_final": final_state,
        "times_s": np.asarray(out_times, dtype=int),
        "temperature_C": np.asarray(out_T, dtype=float).reshape((-1, 21)) if out_T else np.empty((0, 21)),
        "moisture_kg_per_kg": np.asarray(out_C, dtype=float).reshape((-1, 21)) if out_C else np.empty((0, 21)),
        "audit": np.asarray(out_audit, dtype=float).reshape((-1, 4)) if out_audit else np.empty((0, 4)),
        "flux_times_s": np.concatenate(flux_times) if flux_times else np.empty(0, dtype=float),
        "flux_rates_kg_m_s": np.concatenate(flux_rates) if flux_rates else np.empty(0, dtype=float),
        "monitor": monitor,
        "event": event,
        "radius_nodes_m": nodes,
        "radius_output_cm": nodes[sample_nodes] * 100.0,
        "control_volume_weights_m2": volumes,
        "n_intervals": int(n),
        "dr_m": float(dr),
        "runtime_s": float(time.perf_counter() - started),
        "nfev": int(nfev),
        "njev": int(njev),
        "nlu": int(nlu),
        "positivity_retries": int(positivity_retries),
        "rtol": float(rtol),
        "atol": float(atol),
        "max_step_s": float(max_step),
        "h_W_m2K": float(h),
        "hm_m_s": float(hm),
        "initial_mean_moisture": initial_mean,
    }


def _merge_segments(first: dict[str, Any], second: dict[str, Any]) -> dict[str, Any]:
    """Merge two split integrations while preserving the 4-hour state split."""
    if not np.allclose(first["state_final"], second.get("state_initial", first["state_final"]), rtol=0.0, atol=0.0):
        # The second state is not required to be stored; this check is kept by
        # run_case before calling this function.  The branch protects against
        # accidental use with an unrelated continuation.
        raise ValueError("Segment states are not contiguous.")
    merged = dict(second)
    merged["times_s"] = np.concatenate([first["times_s"], second["times_s"]])
    merged["temperature_C"] = np.vstack([first["temperature_C"], second["temperature_C"]])
    merged["moisture_kg_per_kg"] = np.vstack([first["moisture_kg_per_kg"], second["moisture_kg_per_kg"]])
    merged["audit"] = np.vstack([first["audit"], second["audit"]])
    merged["monitor"] = {
        "max_positive_neighbor_jump_kg_per_kg": max(
            first["monitor"]["max_positive_neighbor_jump_kg_per_kg"],
            second["monitor"]["max_positive_neighbor_jump_kg_per_kg"],
        ),
        "max_center_gap_kg_per_kg": max(
            first["monitor"]["max_center_gap_kg_per_kg"],
            second["monitor"]["max_center_gap_kg_per_kg"],
        ),
        "min_moisture_kg_per_kg": min(
            first["monitor"]["min_moisture_kg_per_kg"],
            second["monitor"]["min_moisture_kg_per_kg"],
        ),
        "max_moisture_kg_per_kg": max(
            first["monitor"]["max_moisture_kg_per_kg"],
            second["monitor"]["max_moisture_kg_per_kg"],
        ),
        "checked_accepted_states": first["monitor"]["checked_accepted_states"] + second["monitor"]["checked_accepted_states"],
        "checked_midpoint_states": first["monitor"]["checked_midpoint_states"] + second["monitor"]["checked_midpoint_states"],
        "checked_output_states": first["monitor"]["checked_output_states"] + second["monitor"]["checked_output_states"],
        "segment_records": first["monitor"]["segment_records"] + second["monitor"]["segment_records"],
    }
    merged["runtime_s"] = float(first["runtime_s"] + second["runtime_s"])
    merged["nfev"] = int(first["nfev"] + second["nfev"])
    merged["njev"] = int(first["njev"] + second["njev"])
    merged["nlu"] = int(first["nlu"] + second["nlu"])
    merged["positivity_retries"] = int(first["positivity_retries"] + second["positivity_retries"])
    return merged


def run_case(
    boundary: np.ndarray,
    diagnostic: dict[str, Any],
    *,
    n: int,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_before_s: float,
    max_step_after_s: float,
    chunk_before_s: int = DEFAULT_CHUNK_S,
    chunk_after_s: int = DEFAULT_CHUNK_S,
    initial_state: np.ndarray | None = None,
    h: float = q2.H_BASE,
    hm: float = q2.HM_BASE,
    plateau_temperature_C: float = NOMINAL_PLATEAU_T_C,
    plateau_moisture_kg_per_kg: float = NOMINAL_PLATEAU_C,
) -> dict[str, Any]:
    """Run the full 0-to-end process, splitting exactly at 4 hours."""
    if not diagnostic.get("pass", False):
        raise ValueError("Cannot run the baseline without a passing platform diagnostic.")
    if max_end_s <= PLATEAU_START_S:
        raise ValueError("Q3 max_end_s must exceed 14400 s.")
    if initial_state is None:
        initial_state = q2.initial_state(n)
        first_start = 0
    else:
        initial_state = np.asarray(initial_state, dtype=float)
        first_start = 0
    # Formal run from t=0 provides full-grid monotonicity evidence through the
    # measured boundary interval.  A restart-only run is available through
    # ``run_from_restart`` for quick endpoint exploration.
    first = integrate_segment(
        initial_state,
        n,
        0,
        PLATEAU_START_S,
        boundary,
        rtol=rtol,
        atol=atol,
        max_step=max_step_before_s,
        chunk_seconds=chunk_before_s,
        detect_event=False,
        h=h,
        hm=hm,
    )
    plateau_boundary = make_plateau_boundary(
        diagnostic,
        max_end_s,
        temperature_C=plateau_temperature_C,
        moisture_kg_per_kg=plateau_moisture_kg_per_kg,
    )
    second = integrate_segment(
        first["state_final"],
        n,
        PLATEAU_START_S,
        max_end_s,
        plateau_boundary,
        rtol=rtol,
        atol=atol,
        max_step=max_step_after_s,
        chunk_seconds=chunk_after_s,
        detect_event=True,
        event_kind="grid_max",
        h=h,
        hm=hm,
    )
    # integrate_segment includes the split endpoint in both pieces.  Keep the
    # first copy and drop the duplicate 14400-s row from the continuation.
    second_mask = second["times_s"] > PLATEAU_START_S
    merged = dict(second)
    merged["times_s"] = np.concatenate([first["times_s"], second["times_s"][second_mask]])
    merged["temperature_C"] = np.vstack([first["temperature_C"], second["temperature_C"][second_mask]])
    merged["moisture_kg_per_kg"] = np.vstack([first["moisture_kg_per_kg"], second["moisture_kg_per_kg"][second_mask]])
    merged["audit"] = np.vstack([first["audit"], second["audit"][second_mask]])
    merged["flux_times_s"] = np.concatenate([first["flux_times_s"], second["flux_times_s"]])
    merged["flux_rates_kg_m_s"] = np.concatenate([first["flux_rates_kg_m_s"], second["flux_rates_kg_m_s"]])
    merged["initial_mean_moisture"] = first["initial_mean_moisture"]
    merged["monitor"] = {
        "max_positive_neighbor_jump_kg_per_kg": max(first["monitor"]["max_positive_neighbor_jump_kg_per_kg"], second["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
        "max_center_gap_kg_per_kg": max(first["monitor"]["max_center_gap_kg_kg"] if "max_center_gap_kg_kg" in first["monitor"] else first["monitor"]["max_center_gap_kg_per_kg"], second["monitor"]["max_center_gap_kg_per_kg"]),
        "min_moisture_kg_per_kg": min(first["monitor"]["min_moisture_kg_per_kg"], second["monitor"]["min_moisture_kg_per_kg"]),
        "max_moisture_kg_per_kg": max(first["monitor"]["max_moisture_kg_per_kg"], second["monitor"]["max_moisture_kg_per_kg"]),
        "checked_accepted_states": first["monitor"]["checked_accepted_states"] + second["monitor"]["checked_accepted_states"],
        "checked_midpoint_states": first["monitor"]["checked_midpoint_states"] + second["monitor"]["checked_midpoint_states"],
        "checked_output_states": first["monitor"]["checked_output_states"] + second["monitor"]["checked_output_states"],
        "segment_records": first["monitor"]["segment_records"] + second["monitor"]["segment_records"],
    }
    merged["event"] = second["event"]
    merged["first_segment"] = first
    merged["second_segment"] = second
    merged["final_time_s"] = second["final_time_s"]
    merged["state_final"] = second["state_final"]
    merged["runtime_s"] = float(first["runtime_s"] + second["runtime_s"])
    merged["nfev"] = int(first["nfev"] + second["nfev"])
    merged["njev"] = int(first["njev"] + second["njev"])
    merged["nlu"] = int(first["nlu"] + second["nlu"])
    merged["positivity_retries"] = int(first["positivity_retries"] + second["positivity_retries"])
    merged["boundary_plateau"] = plateau_boundary
    merged["h_W_m2K"] = float(h)
    merged["hm_m_s"] = float(hm)
    return merged


def run_from_restart(
    restart: dict[str, Any],
    diagnostic: dict[str, Any],
    *,
    measured_boundary: np.ndarray,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_after_s: float,
    chunk_after_s: int = DEFAULT_CHUNK_S,
    h: float = q2.H_BASE,
    hm: float = q2.HM_BASE,
    plateau_temperature_C: float = NOMINAL_PLATEAU_T_C,
    plateau_moisture_kg_per_kg: float = NOMINAL_PLATEAU_C,
) -> dict[str, Any]:
    """Continue from Q2's full-grid restart without inventing boundary data."""
    start_s = int(float(restart["final_time_s"]))
    n = int(restart["n_intervals"])
    if start_s > PLATEAU_START_S:
        raise ValueError("Q2 restart must end at or before the 4-hour switch.")
    if start_s < measured_boundary[0, 0] or PLATEAU_START_S > measured_boundary[-1, 0]:
        raise ValueError("Measured boundary does not cover the restart continuation interval.")
    plateau_boundary = make_plateau_boundary(
        diagnostic,
        max_end_s,
        temperature_C=plateau_temperature_C,
        moisture_kg_per_kg=plateau_moisture_kg_per_kg,
    )
    restart_state = interleaved_state(
        restart["temperature_final_C"], restart["moisture_final_kg_per_kg"]
    )
    first = integrate_segment(
        restart_state,
        n,
        start_s,
        PLATEAU_START_S,
        measured_boundary,
        rtol=rtol,
        atol=atol,
        max_step=max_step_after_s,
        chunk_seconds=chunk_after_s,
        detect_event=False,
        h=h,
        hm=hm,
    ) if start_s < PLATEAU_START_S else None
    state_at_switch = first["state_final"] if first is not None else interleaved_state(restart["temperature_final_C"], restart["moisture_final_kg_per_kg"])
    second = integrate_segment(
        state_at_switch,
        n,
        PLATEAU_START_S,
        max_end_s,
        plateau_boundary,
        rtol=rtol,
        atol=atol,
        max_step=max_step_after_s,
        chunk_seconds=chunk_after_s,
        detect_event=True,
        event_kind="grid_max",
        h=h,
        hm=hm,
    )
    if first is None:
        second["first_segment"] = None
        second["second_segment"] = second.copy()
        second["boundary_plateau"] = plateau_boundary
        return second
    second_mask = second["times_s"] > PLATEAU_START_S
    second["times_s"] = np.concatenate([first["times_s"], second["times_s"][second_mask]])
    second["temperature_C"] = np.vstack([first["temperature_C"], second["temperature_C"][second_mask]])
    second["moisture_kg_per_kg"] = np.vstack([first["moisture_kg_per_kg"], second["moisture_kg_per_kg"][second_mask]])
    second["audit"] = np.vstack([first["audit"], second["audit"][second_mask]])
    second["flux_times_s"] = np.concatenate([first["flux_times_s"], second["flux_times_s"]])
    second["flux_rates_kg_m_s"] = np.concatenate([first["flux_rates_kg_m_s"], second["flux_rates_kg_m_s"]])
    second["initial_mean_moisture"] = first["initial_mean_moisture"]
    second["first_segment"] = first
    second["second_segment"] = second.copy()
    second["boundary_plateau"] = plateau_boundary
    second["runtime_s"] = float(first["runtime_s"] + second["runtime_s"])
    second["nfev"] = int(first["nfev"] + second["nfev"])
    second["njev"] = int(first["njev"] + second["njev"])
    second["nlu"] = int(first["nlu"] + second["nlu"])
    second["positivity_retries"] = int(first["positivity_retries"] + second["positivity_retries"])
    return second


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy_style(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy_style(source.alignment)


def write_result3_workbook(
    destination: str | Path,
    times_s: np.ndarray,
    moisture: np.ndarray,
    template: str | Path = DEFAULT_TEMPLATE,
) -> None:
    """Write exactly the template's 60-second sequence through the last full minute."""
    destination = Path(destination)
    times_s = np.asarray(times_s, dtype=int)
    moisture = np.asarray(moisture, dtype=float)
    if times_s.ndim != 1 or moisture.shape != (len(times_s), 21):
        raise ValueError("result3 workbook expects (number_of_minutes, 21) moisture values.")
    if len(times_s) and (times_s[0] != 60 or np.any(np.diff(times_s) != 60)):
        raise ValueError("result3 workbook times must be 60, 120, ... with no extra endpoint row.")
    wb = load_workbook(template)
    if wb.sheetnames != ["Sheet1"]:
        wb.close()
        raise ValueError(f"Unexpected result3 template sheets: {wb.sheetnames}")
    ws = wb["Sheet1"]
    if ws.cell(1, 1).value is None:
        wb.close()
        raise ValueError("result3 A1 header is missing.")
    style_time = copy_style(ws.cell(2, 1)._style) if ws.max_row >= 2 else None
    style_value = copy_style(ws.cell(2, 2)._style) if ws.max_row >= 2 and ws.max_column >= 2 else None
    style_header = copy_style(ws.cell(1, 2)._style) if ws.max_column >= 2 else None
    # Clear existing cells in the used rectangle before creating the exact
    # output dimensions.  No extra t=0 or non-minute event row is inserted.
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, len(times_s) + 1), min_col=1, max_col=22):
        for cell in row:
            cell.value = None
    # The supplied template contains a trailing placeholder row.  Clearing
    # its values is insufficient because openpyxl still reports it as part of
    # the used range; remove all rows after the requested data block so the
    # published workbook has an exact structural contract.
    desired_max_row = len(times_s) + 1
    if ws.max_row > desired_max_row:
        ws.delete_rows(desired_max_row + 1, ws.max_row - desired_max_row)
    for j in range(21):
        cell = ws.cell(1, j + 2, j / 10.0)
        if style_header is not None:
            cell._style = copy_style(style_header)
        cell.number_format = "0.0"
    for row_index, second in enumerate(times_s, start=2):
        time_cell = ws.cell(row_index, 1, int(second))
        if style_time is not None:
            time_cell._style = copy_style(style_time)
        for j in range(21):
            cell = ws.cell(row_index, j + 2, float(f"{moisture[row_index - 2, j]:.4f}"))
            if style_value is not None:
                cell._style = copy_style(style_value)
            cell.number_format = "0.0000"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = max(float(ws.column_dimensions["A"].width or 12), 25.0)
    for column in range(2, 23):
        ws.column_dimensions[ws.cell(1, column).column_letter].width = 12.0
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    wb.close()


def validate_result3_workbook(
    path: str | Path,
    times_s: np.ndarray,
    moisture: np.ndarray,
    template: str | Path = DEFAULT_TEMPLATE,
) -> dict[str, Any]:
    """Read every official workbook cell back, including the exact time schema."""
    path = Path(path)
    times_s = np.asarray(times_s, dtype=int)
    moisture = np.asarray(moisture, dtype=float)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if wb.sheetnames != ["Sheet1"]:
            raise AssertionError(f"Unexpected workbook sheets: {wb.sheetnames}")
        ws = wb["Sheet1"]
        if ws.max_row != len(times_s) + 1 or ws.max_column != 22:
            raise AssertionError(f"Workbook dimensions {ws.max_row}x{ws.max_column} do not match expected {len(times_s)+1}x22.")
        template_wb = load_workbook(template, read_only=True, data_only=True)
        try:
            expected_a1 = template_wb["Sheet1"].cell(1, 1).value
        finally:
            template_wb.close()
        if ws.cell(1, 1).value != expected_a1:
            raise AssertionError("result3 A1 header was not preserved.")
        headers = np.asarray([ws.cell(1, j + 2).value for j in range(21)], dtype=float)
        np.testing.assert_allclose(headers, np.arange(21) / 10.0, atol=1e-12, rtol=0.0)
        observed_times = np.empty(len(times_s), dtype=int)
        observed = np.empty_like(moisture)
        for row_index, row in enumerate(ws.iter_rows(min_row=2, max_row=len(times_s) + 1, max_col=22)):
            observed_times[row_index] = int(row[0].value)
            observed[row_index] = np.asarray([cell.value for cell in row[1:]], dtype=float)
        np.testing.assert_array_equal(observed_times, times_s)
        np.testing.assert_array_equal(observed, np.round(moisture, 4))
        return {
            "path": str(path.resolve()),
            "sheet": "Sheet1",
            "rows": int(len(times_s)),
            "columns": 21,
            "time_schema": "60,120,...,floor(t_dry/60)*60; no non-minute endpoint row",
            "four_decimal_match": True,
        }
    finally:
        wb.close()


def write_csv(path: Path, header: list[str], values: np.ndarray) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(np.asarray(values).tolist())


def _event_metadata(result: dict[str, Any]) -> dict[str, Any]:
    event = result.get("event")
    if event is None:
        return {
            "status": "NOT_REACHED_WITHIN_MAX_END",
            "max_end_s": int(result["second_segment"]["end_s_requested"]),
        }
    event_state = np.asarray(event["state"], dtype=float)
    moisture = event_state[1::2]
    return {
        "status": "EVENT_LOCATED",
        "definition": "first t with max_r C(r,t) <= 0.15 kg/kg; strict below is reported after the event",
        "time_s": float(event["time_s"]),
        "time_h": float(event["time_s"] / 3600.0),
        "center_C": float(moisture[0]),
        "max_C": float(np.max(moisture)),
        "event_residual_C": float(event.get("event_residual_C", np.max(moisture) - THRESHOLD_C)),
        "event_residual_within_tolerance": bool(abs(float(event.get("event_residual_C", np.max(moisture) - THRESHOLD_C))) <= EVENT_CONCENTRATION_TOL),
        "min_C": float(np.min(moisture)),
        "center_minus_max_C": float(moisture[0] - np.max(moisture)),
        "grid_event_bracket_s": event.get("grid_event_bracket_s"),
        "grid_event_bracket_width_s": event.get("grid_event_bracket_width_s"),
        "center_event_time_s": event.get("center_event_time_s"),
        "center_event_bracket_s": event.get("center_event_bracket_s"),
        "center_event_bracket_width_s": event.get("center_event_bracket_width_s"),
        "strict_before_time_s": event.get("strict_before_time_s"),
        "strict_after_time_s": event.get("strict_after_time_s"),
        "strict_before_max_C": event.get("strict_before_max_C"),
        "strict_after_max_C": event.get("strict_after_max_C"),
        "strict_before_pass": bool(event.get("strict_before_pass", False)),
        "strict_after_pass": bool(event.get("strict_after_pass", False)),
        "state_file": "q3_event.npz",
    }


def compute_mass_balance(
    times_s: np.ndarray,
    mean_moisture: np.ndarray,
    surface_moisture: np.ndarray,
    boundary_moisture: np.ndarray,
    initial_mean: float,
    hm: float = q2.HM_BASE,
) -> np.ndarray:
    """Trapezoidal effective boundary-flux balance on output times."""
    outward_rate = q2.RADIUS_M * hm * (surface_moisture - boundary_moisture)
    cumulative = np.zeros(len(times_s), dtype=float)
    if len(times_s) > 1:
        dt = np.diff(times_s.astype(float))
        cumulative[1:] = np.cumsum(0.5 * (outward_rate[:-1] + outward_rate[1:]) * dt)
    residual = mean_moisture + cumulative / (q2.RADIUS_M * q2.RADIUS_M / 2.0) - initial_mean
    return residual


def _load_q2_history() -> dict[str, np.ndarray]:
    with np.load(Q2_FIELDS, allow_pickle=False) as payload:
        fields = {key: payload[key].copy() for key in payload.files}
    if fields["time_s"][0] != 0 or fields["time_s"][-1] != 10800:
        raise ValueError("Q2 history does not cover 0--10800 s.")
    return fields


def _field_difference(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
    if not np.array_equal(a["times_s"], b["times_s"]):
        raise ValueError("Q3 cases have different output times.")
    dt = np.abs(a["temperature_C"] - b["temperature_C"])
    dc = np.abs(a["moisture_kg_per_kg"] - b["moisture_kg_per_kg"])
    return {
        "temperature_max_abs": float(np.max(dt)),
        "moisture_max_abs": float(np.max(dc)),
        "temperature_pass": bool(np.max(dt) < FIELD_THRESHOLD_T_C),
        "moisture_pass": bool(np.max(dc) < FIELD_THRESHOLD_C),
        "time_s_at_max_temperature": int(a["times_s"][np.unravel_index(np.argmax(dt), dt.shape)[0]]),
        "time_s_at_max_moisture": int(a["times_s"][np.unravel_index(np.argmax(dc), dc.shape)[0]]),
    }


def _event_time_difference(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
    ea = a.get("event")
    eb = b.get("event")
    if ea is None or eb is None:
        return {"pass": False, "reason": "one case did not locate the event"}
    ta = float(ea["time_s"])
    tb = float(eb["time_s"])
    diff = abs(tb - ta)
    relative = diff / max(abs(tb), 1.0)
    return {
        "coarse_time_s": ta,
        "fine_time_s": tb,
        "abs_difference_s": diff,
        "relative_difference": relative,
        "relative_threshold": CONVERGENCE_REL_TIME,
        "absolute_threshold_s": CONVERGENCE_ABS_TIME_S,
        "pass": bool(relative < CONVERGENCE_REL_TIME and diff < CONVERGENCE_ABS_TIME_S),
    }


def build_output_arrays(
    result: dict[str, Any],
    q2_history: dict[str, np.ndarray],
    boundary: np.ndarray,
) -> dict[str, np.ndarray]:
    """Combine measured-period Q2 outputs and Q3 continuation outputs."""
    times_q3 = result["times_s"]
    mask_early = q2_history["time_s"] <= PLATEAU_START_S
    early_times = q2_history["time_s"][mask_early].astype(int)
    early_T = q2_history["temperature_C"][mask_early]
    early_C = q2_history["moisture_kg_per_kg"][mask_early]
    # q2 history has 0..10800; q3 formal run has 0..14400 and continuation.
    # Use q3 for all measured-period rows so the audited case is internally
    # self-consistent; q2 history is retained as a compatibility cross-check.
    out_mask = times_q3 >= 0
    times = times_q3[out_mask]
    T = result["temperature_C"][out_mask]
    C = result["moisture_kg_per_kg"][out_mask]
    output = {
        "times_s": times.astype(int),
        "temperature_C": T,
        "moisture_kg_per_kg": C,
        "mean_moisture": result["audit"][out_mask, 0],
        "max_moisture": result["audit"][out_mask, 1],
        "center_gap": result["audit"][out_mask, 2],
        "monotonicity": result["audit"][out_mask, 3],
    }
    # This comparison is intentionally diagnostic only; it does not replace
    # the full-grid result with 21-node Q2 output.
    q2_n = None
    try:
        with np.load(Q2_RESTART, allow_pickle=False) as restart_payload:
            q2_n = int(restart_payload["n_intervals"])
    except (FileNotFoundError, KeyError, ValueError):
        q2_n = None
    if len(early_times) and q2_n == int(result["n_intervals"]):
        lookup = {int(t): i for i, t in enumerate(times)}
        common = np.array([t for t in early_times if int(t) in lookup], dtype=int)
        if len(common):
            qi = np.array([int(np.where(early_times == t)[0][0]) for t in common])
            oi = np.array([lookup[int(t)] for t in common])
            output["q2_history_temperature_max_abs"] = np.array([np.max(np.abs(T[oi] - early_T[qi]))])
            output["q2_history_moisture_max_abs"] = np.array([np.max(np.abs(C[oi] - early_C[qi]))])
        output["q2_compatibility_status"] = "SAME_GRID_COMPARISON"
    else:
        output["q2_compatibility_status"] = "SKIPPED_GRID_MISMATCH"
        output["q2_compatibility_q2_n_intervals"] = np.array([-1 if q2_n is None else q2_n], dtype=int)
    env_T = np.interp(times.astype(float), boundary[:, 0], boundary[:, 1])
    env_C = np.interp(times.astype(float), boundary[:, 0], boundary[:, 2])
    plateau_start = times >= PLATEAU_START_S
    if np.any(plateau_start):
        # Measured boundary interpolation is used up to 14400; the q3
        # continuation output contains only post-switch rows after the split.
        env_T[plateau_start] = result["boundary_plateau"][0, 1]
        env_C[plateau_start] = result["boundary_plateau"][0, 2]
    output["environment_temperature_C"] = env_T
    output["environment_moisture_kg_per_kg"] = env_C
    flux_times = np.asarray(result.get("flux_times_s", []), dtype=float)
    flux_rates = np.asarray(result.get("flux_rates_kg_m_s", []), dtype=float)
    if len(flux_times) >= 2 and len(flux_times) == len(flux_rates):
        order = np.argsort(flux_times, kind="stable")
        flux_times = flux_times[order]
        flux_rates = flux_rates[order]
        cumulative_flux = np.zeros(len(flux_times), dtype=float)
        if len(flux_times) > 1:
            dt = np.diff(flux_times)
            cumulative_flux[1:] = np.cumsum(0.5 * (flux_rates[:-1] + flux_rates[1:]) * dt)
        cumulative_at_output = np.interp(times.astype(float), flux_times, cumulative_flux)
        initial_mean = float(result.get("initial_mean_moisture", output["mean_moisture"][0]))
        output["mass_balance_residual"] = output["mean_moisture"] + cumulative_at_output / (q2.RADIUS_M * q2.RADIUS_M / 2.0) - initial_mean
        output["mass_balance_method"] = "accepted_BDF_nodes_with_piecewise_boundary_flux"
    else:
        output["mass_balance_residual"] = compute_mass_balance(
            times,
            output["mean_moisture"],
            C[:, -1],
            env_C,
            float(output["mean_moisture"][0]),
        )
        output["mass_balance_method"] = "60_second_output_trapezoid_fallback"
    return output


def save_result_files(
    output: Path,
    result: dict[str, Any],
    fields: dict[str, np.ndarray],
    platform: dict[str, Any],
    validation: dict[str, Any],
) -> None:
    output.mkdir(parents=True, exist_ok=True)
    times = fields["times_s"]
    radius_cm = result["radius_output_cm"]
    np.savez_compressed(
        output / "q3_fields.npz",
        time_s=times,
        radius_cm=radius_cm,
        temperature_C=fields["temperature_C"],
        moisture_kg_per_kg=fields["moisture_kg_per_kg"],
        mean_moisture=fields["mean_moisture"],
        max_moisture=fields["max_moisture"],
        center_gap=fields["center_gap"],
        monotonicity=fields["monotonicity"],
        environment_temperature_C=fields["environment_temperature_C"],
        environment_moisture_kg_per_kg=fields["environment_moisture_kg_per_kg"],
        mass_balance_residual=fields["mass_balance_residual"],
    )
    event = result.get("event")
    if event is not None:
        event_state = np.asarray(event["state"], dtype=float)
        np.savez_compressed(
            output / "q3_event.npz",
            time_s=np.array(event["time_s"]),
            radius_nodes_m=result["radius_nodes_m"],
            temperature_C=event_state[0::2],
            moisture_kg_per_kg=event_state[1::2],
            control_volume_weights_m2=result["control_volume_weights_m2"],
        )
    radii_header = [f"r_{value:.1f}_cm" for value in radius_cm]
    write_csv(output / "q3_temperature_full.csv", ["time_s"] + radii_header, np.column_stack([times, fields["temperature_C"]]))
    write_csv(output / "q3_moisture_full.csv", ["time_s"] + radii_header, np.column_stack([times, fields["moisture_kg_per_kg"]]))
    write_csv(output / "q3_mass_balance.csv", ["time_s", "mean_moisture", "max_moisture", "center_gap", "monotonicity", "effective_balance_residual"], np.column_stack([times, fields["mean_moisture"], fields["max_moisture"], fields["center_gap"], fields["monotonicity"], fields["mass_balance_residual"]]))
    write_csv(output / "q3_boundary_used.csv", ["time_s", "temperature_C", "moisture_kg_per_kg"], np.column_stack([times, fields["environment_temperature_C"], fields["environment_moisture_kg_per_kg"]]))
    _write_json(output / "q3_platform_diagnostic.json", platform)
    _write_json(output / "q3_validation.json", validation)


def plot_results(output: Path, fields: dict[str, np.ndarray], event: dict[str, Any] | None, radius_cm: np.ndarray) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    available = {font.name for font in font_manager.fontManager.ttflist}
    font_name = next((name for name in ("Microsoft YaHei", "SimHei", "SimSun", "Noto Sans CJK SC") if name in available), None)
    if font_name:
        plt.rcParams["font.family"] = font_name
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["pdf.fonttype"] = 42
    plt.rcParams["font.size"] = 9
    output.mkdir(parents=True, exist_ok=True)
    times = fields["times_s"].astype(float)
    max_c = fields["max_moisture"]
    fig, axes = plt.subplots(1, 2, figsize=(9.0, 3.6), constrained_layout=True)
    axes[0].plot(times / 3600.0, max_c, color="#174A70", linewidth=1.1, label="全域最大含水率")
    axes[0].axhline(THRESHOLD_C, color="#963E3E", linestyle="--", linewidth=0.9, label="阈值 0.15")
    if event is not None:
        axes[0].axvline(float(event["time_s"]) / 3600.0, color="#D88A41", linestyle=":", linewidth=0.9, label="临界时刻")
    axes[0].set(xlabel="时间（h）", ylabel="干基含水率（kg/kg）")
    axes[0].grid(alpha=0.18)
    axes[0].legend(frameon=False, fontsize=7)
    if event is not None:
        event_time = float(event["time_s"])
        window = (times >= max(0.0, event_time - 6 * 3600)) & (times <= event_time + 6 * 3600)
        axes[1].plot(times[window] / 3600.0, max_c[window], color="#174A70", linewidth=1.1)
        axes[1].axhline(THRESHOLD_C, color="#963E3E", linestyle="--", linewidth=0.9)
        axes[1].axvline(event_time / 3600.0, color="#D88A41", linestyle=":", linewidth=0.9)
        axes[1].set(xlabel="时间（h）", ylabel="阈值附近最大含水率（kg/kg）")
        axes[1].grid(alpha=0.18)
    else:
        axes[1].plot(times / 3600.0, fields["center_gap"], color="#2584A6", linewidth=1.0)
        axes[1].set(xlabel="时间（h）", ylabel="中心与全域最大值差（kg/kg）")
        axes[1].grid(alpha=0.18)
    fig.savefig(output / "q3_moisture_max.pdf", bbox_inches="tight")
    fig.savefig(output / "q3_moisture_max.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    profiles = [int(t) for t in fields["times_s"] if int(t) % PAPER_STEP_S == 0]
    if event is not None:
        profiles.append(-1)
    fig, ax = plt.subplots(figsize=(5.2, 3.7), constrained_layout=True)
    colors = ["#174A70", "#2584A6", "#77AABD", "#D88A41", "#963E3E", "#6A4C93"]
    for index, tt in enumerate(profiles):
        if tt == -1:
            full_state = np.asarray(event["state"], dtype=float)[1::2]
            sample_indices = np.arange(0, len(full_state), max(1, (len(full_state) - 1) // 20), dtype=int)
            state = full_state[sample_indices[: len(radius_cm)]]
            label = "烘干结束"
        else:
            row = int(np.where(fields["times_s"] == tt)[0][0])
            state = fields["moisture_kg_per_kg"][row]
            label = f"{tt / 3600:g} h"
        ax.plot(radius_cm, state, color=colors[index % len(colors)], linewidth=1.0, label=label)
    ax.axhline(THRESHOLD_C, color="#963E3E", linestyle="--", linewidth=0.8)
    ax.set(xlabel="径向距离（cm）", ylabel="干基含水率（kg/kg）")
    ax.grid(alpha=0.18)
    ax.legend(frameon=False, fontsize=7, ncol=2)
    fig.savefig(output / "q3_moisture_profiles.pdf", bbox_inches="tight")
    fig.savefig(output / "q3_moisture_profiles.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def run_radau_crosscheck(
    state: np.ndarray,
    n: int,
    boundary: np.ndarray,
    start_s: int,
    end_s: int,
    *,
    rtol: float,
    atol: float,
    max_step: float,
    h: float = q2.H_BASE,
    hm: float = q2.HM_BASE,
) -> dict[str, Any]:
    """Independent medium-grid Radau continuation for a fixed interval."""
    rhs_log, jac_log, geometry = q2.make_log_rhs(boundary, n, h=h, hm=hm)
    dr, nodes, faces, volumes = geometry
    transformed = _to_log_state(state, n)
    started = time.perf_counter()
    solution = solve_ivp(
        rhs_log,
        (start_s, end_s),
        transformed,
        method="Radau",
        jac=jac_log,
        jac_sparsity=q2.jacobian_sparsity(n),
        rtol=rtol,
        atol=atol,
        max_step=max_step,
    )
    if not solution.success:
        raise RuntimeError(f"Radau cross-check failed: {solution.message}")
    final, moisture = _physical_from_log(solution.y[:, -1], n)
    return {
        "status": "PASSED",
        "n_intervals": int(n),
        "start_s": int(start_s),
        "end_s": int(end_s),
        "temperature_final_C": final[0::2],
        "moisture_final_kg_per_kg": moisture,
        "nfev": int(solution.nfev),
        "njev": int(solution.njev or 0),
        "nlu": int(solution.nlu or 0),
        "runtime_s": float(time.perf_counter() - started),
        "h_W_m2K": float(h),
        "hm_m_s": float(hm),
    }


def sensitivity_endpoint(
    boundary: np.ndarray,
    diagnostic: dict[str, Any],
    *,
    n: int,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_before_s: float,
    max_step_after_s: float,
    h: float,
    hm: float,
    plateau_temperature_C: float = NOMINAL_PLATEAU_T_C,
    plateau_moisture_kg_per_kg: float = NOMINAL_PLATEAU_C,
) -> dict[str, Any]:
    """Run one explicit sensitivity case and return its endpoint evidence."""
    started = time.perf_counter()
    case = run_case(
        boundary,
        diagnostic,
        n=n,
        max_end_s=max_end_s,
        rtol=rtol,
        atol=atol,
        max_step_before_s=max_step_before_s,
        max_step_after_s=max_step_after_s,
        h=h,
        hm=hm,
        plateau_temperature_C=plateau_temperature_C,
        plateau_moisture_kg_per_kg=plateau_moisture_kg_per_kg,
    )
    event = case.get("event")
    return {
        "status": "EVENT_LOCATED" if event is not None else "NOT_REACHED_WITHIN_MAX_END",
        "n_intervals": int(n),
        "h_W_m2K": float(h),
        "hm_m_s": float(hm),
        "plateau_temperature_C": float(plateau_temperature_C),
        "plateau_moisture_kg_per_kg": float(plateau_moisture_kg_per_kg),
        "time_s": None if event is None else float(event["time_s"]),
        "time_h": None if event is None else float(event["time_s"] / 3600.0),
        "center_C": None if event is None else float(event["center_C"]),
        "max_C": None if event is None else float(event["max_C"]),
        "runtime_s": float(time.perf_counter() - started),
        "monotonicity_max_positive_neighbor_jump_kg_per_kg": float(case["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
    }


def _compare_case_fields(coarse: dict[str, Any], fine: dict[str, Any]) -> dict[str, Any]:
    """Compare common 60-second fields without requiring equal event lengths."""
    coarse_times = np.asarray(coarse["times_s"], dtype=float)
    fine_times = np.asarray(fine["times_s"], dtype=float)
    if len(coarse_times) == 0 or len(fine_times) == 0:
        return {"pass": False, "reason": "one case has no output samples"}
    common = coarse_times[coarse_times <= fine_times[-1] + 1e-9]
    if len(common) == 0:
        return {"pass": False, "reason": "cases have no common time range"}
    def interpolate(values: np.ndarray) -> np.ndarray:
        return np.column_stack([np.interp(common, fine_times, values[:, j]) for j in range(values.shape[1])])
    fine_t = interpolate(np.asarray(fine["temperature_C"], dtype=float))
    fine_c = interpolate(np.asarray(fine["moisture_kg_per_kg"], dtype=float))
    coarse_t = np.asarray(coarse["temperature_C"], dtype=float)[: len(common)]
    coarse_c = np.asarray(coarse["moisture_kg_per_kg"], dtype=float)[: len(common)]
    dt = np.abs(coarse_t - fine_t)
    dc = np.abs(coarse_c - fine_c)
    return {
        "common_start_s": int(common[0]),
        "common_end_s": int(common[-1]),
        "common_rows": int(len(common)),
        "temperature_max_abs": float(np.max(dt)),
        "moisture_max_abs": float(np.max(dc)),
        "temperature_threshold": FIELD_THRESHOLD_T_C,
        "moisture_threshold": FIELD_THRESHOLD_C,
        "temperature_pass": bool(np.max(dt) < FIELD_THRESHOLD_T_C),
        "moisture_pass": bool(np.max(dc) < FIELD_THRESHOLD_C),
        "pass": bool(np.max(dt) < FIELD_THRESHOLD_T_C and np.max(dc) < FIELD_THRESHOLD_C),
    }


def run_convergence(
    boundary: np.ndarray,
    diagnostic: dict[str, Any],
    *,
    output: Path,
    coarse_n: int,
    fine_n: int,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_before_s: float,
    max_step_after_s: float,
) -> tuple[dict[str, Any] | None, dict[str, Any]]:
    """Run spatial and time refinement and persist only compact evidence."""
    if coarse_n >= fine_n or coarse_n % 20 or fine_n % 20:
        raise ValueError("Convergence grids must be increasing multiples of 20.")
    diagnostics_dir = output / "diagnostics"
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    records: dict[str, Any] = {
        "coarse_n": int(coarse_n),
        "fine_n": int(fine_n),
        "relative_time_threshold": CONVERGENCE_REL_TIME,
        "absolute_time_threshold_s": CONVERGENCE_ABS_TIME_S,
        "temperature_field_threshold_C": FIELD_THRESHOLD_T_C,
        "moisture_field_threshold_kg_per_kg": FIELD_THRESHOLD_C,
    }
    try:
        coarse = run_case(
            boundary,
            diagnostic,
            n=coarse_n,
            max_end_s=max_end_s,
            rtol=rtol,
            atol=atol,
            max_step_before_s=max_step_before_s,
            max_step_after_s=max_step_after_s,
        )
        fine = run_case(
            boundary,
            diagnostic,
            n=fine_n,
            max_end_s=max_end_s,
            rtol=rtol,
            atol=atol,
            max_step_before_s=max_step_before_s,
            max_step_after_s=max_step_after_s,
        )
        spatial_event = _event_time_difference(
            {"event": coarse.get("event")}, {"event": fine.get("event")}
        )
        spatial_fields = _compare_case_fields(coarse, fine)
        spatial = {
            "coarse_n": int(coarse_n),
            "fine_n": int(fine_n),
            "event": spatial_event,
            "fields": spatial_fields,
            "monotonicity_coarse": float(coarse["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
            "monotonicity_fine": float(fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
            "monotonicity_pass": bool(
                coarse["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
                and fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
            ),
            "pass": bool(
                spatial_event.get("pass", False)
                and spatial_fields.get("pass", False)
                and coarse["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
                and fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
            ),
        }
        records["spatial"] = spatial
        _write_json(diagnostics_dir / "q3_spatial_convergence.json", spatial)

        tight = run_case(
            boundary,
            diagnostic,
            n=fine_n,
            max_end_s=max_end_s,
            rtol=rtol / 10.0,
            atol=atol / 10.0,
            max_step_before_s=max_step_before_s / 2.0,
            max_step_after_s=max_step_after_s / 2.0,
        )
        temporal_event = _event_time_difference(
            {"event": fine.get("event")}, {"event": tight.get("event")}
        )
        temporal_fields = _compare_case_fields(fine, tight)
        temporal = {
            "base": {"rtol": float(rtol), "atol": float(atol), "max_step_before_s": float(max_step_before_s), "max_step_after_s": float(max_step_after_s)},
            "tight": {"rtol": float(rtol / 10.0), "atol": float(atol / 10.0), "max_step_before_s": float(max_step_before_s / 2.0), "max_step_after_s": float(max_step_after_s / 2.0)},
            "event": temporal_event,
            "fields": temporal_fields,
            "monotonicity_base": float(fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
            "monotonicity_tight": float(tight["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
            "monotonicity_pass": bool(
                fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
                and tight["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
            ),
            "pass": bool(
                temporal_event.get("pass", False)
                and temporal_fields.get("pass", False)
                and fine["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
                and tight["monitor"]["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C
            ),
        }
        records["temporal"] = temporal
        records["status"] = "PASSED" if spatial["pass"] and temporal["pass"] else "REVIEW_REQUIRED"
        if temporal["pass"]:
            records["selected_final_settings"] = {
                "n_intervals": int(tight["n_intervals"]),
                "rtol": float(tight["rtol"]),
                "atol": float(tight["atol"]),
                "max_step_s": float(tight["max_step_s"]),
            }
            final_result = tight
        else:
            final_result = fine
        _write_json(diagnostics_dir / "q3_temporal_convergence.json", temporal)
        _write_json(output / "q3_convergence.json", records)
        return final_result, records
    except Exception as exc:
        records["status"] = "FAILED"
        records["error"] = repr(exc)
        _write_json(output / "q3_convergence.json", records)
        return None, records


def run_radau_audit(
    boundary: np.ndarray,
    diagnostic: dict[str, Any],
    *,
    n: int,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_before_s: float,
    max_step_after_s: float,
) -> dict[str, Any]:
    """Compare an independently integrated medium-grid BDF/Radau segment."""
    # All q2/q3 grids must be integer multiples of 20 so the 0.1 cm output
    # positions are exact nodes.  Use the largest moderate grid not exceeding
    # 1000 intervals for the independent time-integrator comparison.
    cross_n = min(int(n), 1000)
    cross_end = min(int(max_end_s), PLATEAU_START_S + 3600)
    if cross_end <= PLATEAU_START_S:
        return {"status": "NOT_RUN", "reason": "max_end_s does not cover a post-switch audit interval"}
    started = time.perf_counter()
    bdf = run_case(
        boundary,
        diagnostic,
        n=cross_n,
        max_end_s=cross_end,
        rtol=rtol,
        atol=atol,
        max_step_before_s=max_step_before_s,
        max_step_after_s=max_step_after_s,
    )
    if bdf.get("event") is not None:
        return {"status": "NOT_RUN", "reason": "event occurred before the comparison endpoint"}
    first_state = bdf["first_segment"]["state_final"]
    plateau = bdf["boundary_plateau"]
    radau = run_radau_crosscheck(
        first_state,
        cross_n,
        plateau,
        PLATEAU_START_S,
        cross_end,
        rtol=rtol,
        atol=atol,
        max_step=max_step_after_s,
    )
    bdf_final = np.asarray(bdf["state_final"], dtype=float)
    radau_final = np.asarray(radau["temperature_final_C"], dtype=float)
    radau_state = interleaved_state(radau["temperature_final_C"], radau["moisture_final_kg_per_kg"])
    dt = np.max(np.abs(bdf_final[0::2] - radau_state[0::2]))
    dc = np.max(np.abs(bdf_final[1::2] - radau_state[1::2]))
    report = {
        "status": "PASSED" if dt < FIELD_THRESHOLD_T_C and dc < FIELD_THRESHOLD_C else "REVIEW_REQUIRED",
        "n_intervals": int(cross_n),
        "start_s": int(PLATEAU_START_S),
        "end_s": int(cross_end),
        "temperature_max_abs_C": float(dt),
        "moisture_max_abs_kg_per_kg": float(dc),
        "temperature_threshold_C": FIELD_THRESHOLD_T_C,
        "moisture_threshold_kg_per_kg": FIELD_THRESHOLD_C,
        "bdf_nfev": int(bdf["nfev"]),
        "radau_nfev": int(radau["nfev"]),
        "radau_runtime_s": float(radau["runtime_s"]),
        "total_runtime_s": float(time.perf_counter() - started),
    }
    return report


def run_sensitivities(
    boundary: np.ndarray,
    diagnostic: dict[str, Any],
    *,
    baseline_event_time_s: float | None,
    n: int,
    max_end_s: int,
    rtol: float,
    atol: float,
    max_step_before_s: float,
    max_step_after_s: float,
    output: Path,
) -> list[dict[str, Any]]:
    """Run the requested boundary and transfer-coefficient sensitivities."""
    sensitivity_n = min(int(n), 1000)
    mean_T = float(diagnostic["temperature_C"]["mean"])
    mean_C = float(diagnostic["moisture_kg_per_kg"]["mean"])
    last_T = float(boundary[-1, 1])
    last_C = float(boundary[-1, 2])
    cases = [
        ("final_hour_mean", q2.H_BASE, q2.HM_BASE, mean_T, mean_C),
        ("last_observation", q2.H_BASE, q2.HM_BASE, last_T, last_C),
        ("h_minus20", 0.8 * q2.H_BASE, q2.HM_BASE, NOMINAL_PLATEAU_T_C, NOMINAL_PLATEAU_C),
        ("h_plus20", 1.2 * q2.H_BASE, q2.HM_BASE, NOMINAL_PLATEAU_T_C, NOMINAL_PLATEAU_C),
        ("hm_minus20", q2.H_BASE, 0.8 * q2.HM_BASE, NOMINAL_PLATEAU_T_C, NOMINAL_PLATEAU_C),
        ("hm_plus20", q2.H_BASE, 1.2 * q2.HM_BASE, NOMINAL_PLATEAU_T_C, NOMINAL_PLATEAU_C),
    ]
    # The formal baseline may use a finer grid than the sensitivity budget.
    # Compute a same-grid nominal reference so every delta in the sensitivity
    # table is interpreted against a like-for-like numerical baseline.
    reference = sensitivity_endpoint(
        boundary,
        diagnostic,
        n=sensitivity_n,
        max_end_s=max_end_s,
        rtol=rtol,
        atol=atol,
        max_step_before_s=max_step_before_s,
        max_step_after_s=max_step_after_s,
        h=q2.H_BASE,
        hm=q2.HM_BASE,
        plateau_temperature_C=NOMINAL_PLATEAU_T_C,
        plateau_moisture_kg_per_kg=NOMINAL_PLATEAU_C,
    )
    reference_time = reference.get("time_s")
    rows: list[dict[str, Any]] = []
    for name, h, hm, plateau_T, plateau_C in cases:
        try:
            row = sensitivity_endpoint(
                boundary,
                diagnostic,
                n=sensitivity_n,
                max_end_s=max_end_s,
                rtol=rtol,
                atol=atol,
                max_step_before_s=max_step_before_s,
                max_step_after_s=max_step_after_s,
                h=h,
                hm=hm,
                plateau_temperature_C=plateau_T,
                plateau_moisture_kg_per_kg=plateau_C,
            )
            row["case"] = name
            row["delta_time_s_vs_baseline"] = (
                None if baseline_event_time_s is None or row["time_s"] is None else float(row["time_s"] - baseline_event_time_s)
            )
            row["nominal_reference_time_s_same_n"] = reference_time
            row["delta_time_s_vs_nominal_same_n"] = (
                None if reference_time is None or row["time_s"] is None else float(row["time_s"] - reference_time)
            )
        except Exception as exc:
            row = {"case": name, "status": "FAILED", "error": repr(exc), "n_intervals": sensitivity_n}
        rows.append(row)
    headers = sorted({key for row in rows for key in row})
    output.mkdir(parents=True, exist_ok=True)
    with (output / "q3_sensitivity.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    _write_json(
        output / "q3_sensitivity.json",
        {
            "sensitivity_n": sensitivity_n,
            "reference": reference,
            "formal_baseline_event_time_s": baseline_event_time_s,
            "cases": rows,
            "delta_interpretation": "Use delta_time_s_vs_nominal_same_n for sensitivity; formal-baseline delta is retained only as a cross-grid context value.",
        },
    )
    return rows


def _paper_rows(fields: dict[str, np.ndarray], result: dict[str, Any]) -> tuple[np.ndarray, np.ndarray]:
    event = result.get("event")
    times = [int(t) for t in fields["times_s"] if int(t) % PAPER_STEP_S == 0]
    rows = [fields["moisture_kg_per_kg"][int(np.where(fields["times_s"] == t)[0][0])] for t in times]
    labels = [float(t) for t in times]
    if event is not None:
        labels.append(float(event["time_s"]))
        rows.append(np.asarray(event["state"])[1::2][:: max(1, result["n_intervals"] // 20)])
    return np.asarray(labels, dtype=float), np.asarray(rows, dtype=float)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--boundary", type=Path, default=DEFAULT_BOUNDARY)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--n", type=int, default=10240, help="Final grid intervals.")
    parser.add_argument("--coarse-n", type=int, default=5120, help="Coarse grid used by spatial convergence.")
    parser.add_argument("--max-end-s", type=int, default=DEFAULT_MAX_END_S)
    parser.add_argument("--rtol", type=float, default=1e-9)
    parser.add_argument("--atol", type=float, default=1e-11)
    parser.add_argument("--max-step-before", type=float, default=5.0)
    parser.add_argument("--max-step-after", type=float, default=60.0)
    parser.add_argument("--chunk-before", type=int, default=DEFAULT_CHUNK_S)
    parser.add_argument("--chunk-after", type=int, default=DEFAULT_CHUNK_S)
    parser.add_argument("--from-restart", action="store_true")
    parser.add_argument("--restart", type=Path, default=Q2_RESTART)
    parser.add_argument("--skip-convergence", action="store_true", help="Exploratory run only; cannot publish.")
    parser.add_argument("--skip-sensitivity", action="store_true", help="Exploratory run only; cannot publish.")
    parser.add_argument("--no-radau", action="store_true", help="Skip Radau audit; cannot publish.")
    parser.add_argument("--publish", action="store_true")
    args = parser.parse_args()

    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    # Invalidate any earlier exploratory PASS/REVIEW file before a new run.
    # This prevents a failed refinement from leaving stale evidence behind.
    _write_json(
        output / "q3_validation.json",
        {
            "status": "RUNNING",
            "requested_n_intervals": int(args.n),
            "requested_coarse_n_intervals": int(args.coarse_n),
            "publish_requested": bool(args.publish),
        },
    )
    boundary, boundary_sheet = q2.read_boundary(args.boundary, end=PLATEAU_START_S)
    diagnostic = platform_diagnostic(boundary)
    _write_json(output / "q3_platform_diagnostic.json", diagnostic)
    if not diagnostic["pass"]:
        failure = {"status": "PLATFORM_CHECK_FAILED", "platform": diagnostic}
        _write_json(output / "q3_failure_summary.json", failure)
        _write_json(output / "q3_validation.json", failure)
        raise SystemExit(2)

    started = time.perf_counter()
    convergence: dict[str, Any]
    if args.from_restart:
        restart = load_q2_restart(args.restart)
        result = run_from_restart(
            restart,
            diagnostic,
            measured_boundary=boundary,
            max_end_s=args.max_end_s,
            rtol=args.rtol,
            atol=args.atol,
            max_step_after_s=args.max_step_after,
            chunk_after_s=args.chunk_after,
        )
        convergence = {"status": "NOT_RUN", "reason": "restart exploration does not replace the formal 0-to-end convergence run"}
    elif args.skip_convergence:
        result = run_case(
            boundary,
            diagnostic,
            n=args.n,
            max_end_s=args.max_end_s,
            rtol=args.rtol,
            atol=args.atol,
            max_step_before_s=args.max_step_before,
            max_step_after_s=args.max_step_after,
            chunk_before_s=args.chunk_before,
            chunk_after_s=args.chunk_after,
        )
        convergence = {"status": "NOT_RUN", "reason": "--skip-convergence was supplied"}
    else:
        result, convergence = run_convergence(
            boundary,
            diagnostic,
            output=output,
            coarse_n=args.coarse_n,
            fine_n=args.n,
            max_end_s=args.max_end_s,
            rtol=args.rtol,
            atol=args.atol,
            max_step_before_s=args.max_step_before,
            max_step_after_s=args.max_step_after,
        )
        if result is None or convergence.get("status") != "PASSED":
            failure = {"status": "CONVERGENCE_FAILED", "convergence": convergence}
            _write_json(output / "q3_failure_summary.json", failure)
            _write_json(output / "q3_validation.json", failure)
            print(json.dumps(_json_safe({"status": "CONVERGENCE_FAILED", "convergence": convergence}), ensure_ascii=False))
            raise SystemExit(2)
    q2_history = _load_q2_history()

    fields = build_output_arrays(result, q2_history, boundary)
    q2_compatibility_pass = bool(
        fields.get("q2_compatibility_status") == "SAME_GRID_COMPARISON"
        and float(fields.get("q2_history_temperature_max_abs", np.array([math.inf]))[0]) < FIELD_THRESHOLD_T_C
        and float(fields.get("q2_history_moisture_max_abs", np.array([math.inf]))[0]) < FIELD_THRESHOLD_C
    )
    event_meta = _event_metadata(result)
    monitor = result["monitor"]
    monotonicity_pass = bool(monitor["max_positive_neighbor_jump_kg_per_kg"] <= MONOTONICITY_TOL_C)
    event = result.get("event")
    center_time = None if event is None else event.get("center_event_time_s")
    center_grid_time_difference = None if event is None or center_time is None else abs(float(center_time) - float(event["time_s"]))
    center_event_pass = bool(
        event is not None
        and monotonicity_pass
        and center_time is not None
        and center_grid_time_difference <= EVENT_TIME_TOL_S
        and event["center_minus_max_C"] >= -MONOTONICITY_TOL_C
    )
    grid_event_crosscheck = {
        "status": "CENTER_AND_GRID_AGREE" if center_event_pass else "GRID_MAX_REQUIRES_SEPARATE_EVENT",
        "grid_max_event_time_s": None if event is None else float(event["time_s"]),
        "center_event_time_s": None if center_time is None else float(center_time),
        "grid_center_time_difference_s": center_grid_time_difference,
        "event_center_minus_grid_max_C": None if event is None else float(event["center_minus_max_C"]),
        "center_event_within_monotonicity_tolerance": center_event_pass,
    }
    strict_event_pass = bool(
        event is not None
        and event.get("strict_before_pass", False)
        and event.get("strict_after_pass", False)
        and float(event.get("grid_event_bracket_width_s", math.inf)) <= EVENT_TIME_TOL_S
    )
    event_root_pass = bool(
        event is not None
        and abs(float(event.get("event_residual_C", math.inf))) <= EVENT_CONCENTRATION_TOL
    )
    strict_minute_end = None if event is None else int(math.floor(float(event["time_s"]) / 60.0) * 60)
    if strict_minute_end is not None:
        workbook_mask = (fields["times_s"] >= 60) & (fields["times_s"] <= strict_minute_end)
    else:
        workbook_mask = fields["times_s"] >= 60
    workbook_times = fields["times_s"][workbook_mask]
    workbook_moisture = fields["moisture_kg_per_kg"][workbook_mask]
    candidate = output / "result3_candidate.xlsx"
    write_result3_workbook(candidate, workbook_times, workbook_moisture, template=args.template)
    workbook_report = validate_result3_workbook(candidate, workbook_times, workbook_moisture, template=args.template)

    mass_balance_max_abs = float(np.max(np.abs(fields["mass_balance_residual"])))
    initial_mean_for_balance = float(result.get("initial_mean_moisture", fields["mean_moisture"][0]))
    mass_balance_relative_max = mass_balance_max_abs / max(abs(initial_mean_for_balance), 1.0e-12)
    mass_balance = {
        "max_abs_residual": mass_balance_max_abs,
        "initial_mean_moisture": initial_mean_for_balance,
        "max_relative_residual": mass_balance_relative_max,
        "relative_tolerance": MASS_BALANCE_REL_TOL,
        "pass": bool(mass_balance_relative_max <= MASS_BALANCE_REL_TOL),
        "interpretation": "effective moisture-equation boundary-flux closure; relative residual is normalized by the initial mean moisture",
    }
    if args.no_radau:
        radau = {"status": "NOT_RUN", "reason": "--no-radau was supplied"}
    else:
        try:
            radau = run_radau_audit(
                boundary,
                diagnostic,
                n=int(result["n_intervals"]),
                max_end_s=args.max_end_s,
                rtol=args.rtol,
                atol=args.atol,
                max_step_before_s=args.max_step_before,
                max_step_after_s=args.max_step_after,
            )
        except Exception as exc:
            radau = {"status": "FAILED", "error": repr(exc)}
    if args.skip_sensitivity or args.from_restart:
        sensitivity_rows: list[dict[str, Any]] = []
        sensitivity_status = {"status": "NOT_RUN", "reason": "sensitivity was skipped for exploratory/restart execution"}
    else:
        baseline_event_time = None if event is None else float(event["time_s"])
        # Sensitivity cases can be slower than nominal (especially h_m-20%).
        # Let every case use the full 168-hour ceiling rather than declaring a
        # case failed merely because it crosses the nominal endpoint later.
        sensitivity_end = args.max_end_s
        try:
            sensitivity_rows = run_sensitivities(
                boundary,
                diagnostic,
                baseline_event_time_s=baseline_event_time,
                n=int(result["n_intervals"]),
                max_end_s=sensitivity_end,
                rtol=args.rtol,
                atol=args.atol,
                max_step_before_s=args.max_step_before,
                max_step_after_s=args.max_step_after,
                output=output,
            )
            sensitivity_pass = bool(sensitivity_rows) and all(row.get("status") == "EVENT_LOCATED" for row in sensitivity_rows)
            sensitivity_status = {"status": "PASSED" if sensitivity_pass else "REVIEW_REQUIRED", "rows": sensitivity_rows}
        except Exception as exc:
            sensitivity_rows = []
            sensitivity_status = {"status": "FAILED", "error": repr(exc)}
    convergence_pass = convergence.get("status") == "PASSED"
    radau_pass = radau.get("status") == "PASSED"
    sensitivity_pass = sensitivity_status.get("status") == "PASSED"
    all_checks_pass = bool(
        event is not None
        and diagnostic["pass"]
        and monotonicity_pass
        and center_event_pass
        and strict_event_pass
        and event_root_pass
        and workbook_report["four_decimal_match"]
        and mass_balance["pass"]
        and convergence_pass
        and radau_pass
        and sensitivity_pass
        and q2_compatibility_pass
    )
    validation: dict[str, Any] = {
        "status": "NUMERICAL_CHECKS_PASSED" if all_checks_pass else "REVIEW_REQUIRED",
        "boundary_source": str(args.boundary.resolve()),
        "boundary_sheet": boundary_sheet,
        "platform": diagnostic,
        "parameters": {
            "radius_m": q2.RADIUS_M,
            "length_m": q2.LENGTH_M,
            "initial_temperature_C": q2.INITIAL_T_C,
            "initial_moisture_kg_per_kg": q2.INITIAL_C,
            "h_W_m2K": q2.H_BASE,
            "hm_m_s": q2.HM_BASE,
            "post_14400_temperature_C": NOMINAL_PLATEAU_T_C,
            "post_14400_moisture_kg_per_kg": NOMINAL_PLATEAU_C,
            "h_assumption": "Q3 continues the Q2 coefficient because no new h is supplied.",
            "hm_assumption": "Q3 continues the Q2 coefficient because no new hm is supplied.",
        },
        "model": "Q3 fixed-radius axisymmetric variable-property coupled FV+BDF; log(C) internal coordinate",
        "n_intervals": int(result["n_intervals"]),
        "dr_m": float(result["dr_m"]),
        "rtol": float(result.get("rtol", args.rtol)),
        "atol": float(result.get("atol", args.atol)),
        "max_step_before_s": float(result.get("first_segment", {}).get("max_step_s", args.max_step_before)),
        "max_step_after_s": float(result.get("max_step_s", args.max_step_after)),
        "runtime_s": float(time.perf_counter() - started),
        "integrator": {"nfev": result["nfev"], "njev": result["njev"], "nlu": result["nlu"], "positivity_retries": result["positivity_retries"]},
        "event": event_meta,
        "grid_event_crosscheck": grid_event_crosscheck,
        "strict_event_evidence": {
            "pass": strict_event_pass,
            "event_root_pass": event_root_pass,
            "event_residual_C": None if event is None else event.get("event_residual_C"),
            "event_residual_tolerance_C": EVENT_CONCENTRATION_TOL,
            "before_time_s": None if event is None else event.get("strict_before_time_s"),
            "after_time_s": None if event is None else event.get("strict_after_time_s"),
            "before_max_C": None if event is None else event.get("strict_before_max_C"),
            "after_max_C": None if event is None else event.get("strict_after_max_C"),
            "bracket_width_s": None if event is None else event.get("grid_event_bracket_width_s"),
        },
        "monotonicity": {
            "max_positive_neighbor_jump_kg_per_kg": float(monitor["max_positive_neighbor_jump_kg_per_kg"]),
            "max_center_gap_kg_per_kg": float(monitor["max_center_gap_kg_per_kg"]),
            "tolerance_kg_per_kg": MONOTONICITY_TOL_C,
            "pass": monotonicity_pass,
            "checked_accepted_states": monitor["checked_accepted_states"],
            "checked_midpoint_states": monitor["checked_midpoint_states"],
            "checked_output_states": monitor["checked_output_states"],
        },
        "result3_time_schema": "60,120,...,floor(t_dry/60)*60; no non-minute endpoint row",
        "workbook_validation": workbook_report,
        "official_workbook": str(OFFICIAL_RESULT3.resolve()),
        "convergence": convergence,
        "sensitivity": sensitivity_status,
        "independent_radau": radau,
        "mass_balance": mass_balance,
        "gate_summary": {
            "event": bool(event is not None),
            "platform": bool(diagnostic["pass"]),
            "monotonicity": monotonicity_pass,
            "center_grid_crosscheck": center_event_pass,
            "strict_event": strict_event_pass,
            "event_root": event_root_pass,
            "workbook": bool(workbook_report["four_decimal_match"]),
            "mass_balance": bool(mass_balance["pass"]),
            "convergence": convergence_pass,
            "radau": radau_pass,
            "sensitivity": sensitivity_pass,
            "q2_same_grid_compatibility": q2_compatibility_pass,
        },
        "field_output": {
            "rows": int(len(fields["times_s"])),
            "radius_points": 21,
            "mass_balance_max_abs": mass_balance_max_abs,
            "mass_balance_max_relative": mass_balance_relative_max,
            "q2_compatibility_status": fields.get("q2_compatibility_status", "UNAVAILABLE"),
            "q2_compatibility_temperature_max_abs": None if "q2_history_temperature_max_abs" not in fields else float(fields["q2_history_temperature_max_abs"][0]),
            "q2_compatibility_moisture_max_abs": None if "q2_history_moisture_max_abs" not in fields else float(fields["q2_history_moisture_max_abs"][0]),
        },
        "limitations": [
            "The post-14400 s boundary is a model extension supported by a final-hour platform screen, not future observations.",
            "The one-dimensional radial approximation and omitted latent-heat feedback are not quantified here.",
            "Mass balance is an effective moisture-equation flux audit.",
        ],
    }
    # Save all artifacts before publication.  The official workbook is never
    # touched until every requested check reaches a publishable status.
    save_result_files(output, result, fields, diagnostic, validation)
    plot_results(output, fields, result.get("event"), result["radius_output_cm"])
    if validation["status"] == "NUMERICAL_CHECKS_PASSED" and args.publish:
        shutil.copy2(candidate, OFFICIAL_RESULT3)
        validation["official_workbook_validation"] = validate_result3_workbook(OFFICIAL_RESULT3, workbook_times, workbook_moisture, template=args.template)
        _write_json(output / "q3_validation.json", validation)
    else:
        validation["official_workbook_validation"] = {"status": "NOT_OVERWRITTEN", "reason": "requires all checks and --publish"}
        _write_json(output / "q3_validation.json", validation)
    print(json.dumps(_json_safe({"status": validation["status"], "event": event_meta, "candidate": str(candidate), "official": str(OFFICIAL_RESULT3), "runtime_s": validation["runtime_s"], "n_intervals": result["n_intervals"], "gate_summary": validation["gate_summary"]}), ensure_ascii=False))
    if validation["status"] != "NUMERICAL_CHECKS_PASSED":
        raise SystemExit(2)


if __name__ == "__main__":
    main()
