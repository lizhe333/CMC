"""Q1 radial heat conduction and conservative nonlinear moisture diffusion."""
from __future__ import annotations
import argparse
import csv
import json
import platform
import sys
import time
from pathlib import Path

import numpy as np
from numba import njit
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
R, ALPHA, H, K, HM = .02, .36 / (820 * 2600), 25., .36, 8e-7
TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800])
COLS = np.array([0, 5, 10, 15, 20])


def read_boundary(path, sheet=None, end=1800):
    """Read original workbook without modification; reject gaps and extrapolation."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.values)
        header = next((i for i, row in enumerate(rows[:15]) if len(row) >= 3 and
                       all(k in str(row[j]).replace(" ", "") for j, k in
                           enumerate(("时间", "温度", "水分浓度")))), None)
        if header is None:
            raise ValueError("Expected columns: 时间(s), 温度(°C), 水分浓度(kg/kg).")
        values = []
        for index, row in enumerate(rows[header + 1:], header + 2):
            if all(v is None for v in row):
                continue
            if len(row) < 3 or any(v is None or isinstance(v, bool) for v in row[:3]):
                raise ValueError(f"Invalid boundary row {index}.")
            values.append([float(v) for v in row[:3]])
        data = np.asarray(values, dtype=float)
        if data.ndim != 2 or len(data) < 2 or not np.isfinite(data).all():
            raise ValueError("At least two finite boundary samples are required.")
        if (np.diff(data[:, 0]) <= 0).any():
            raise ValueError("Times must be strictly increasing, without duplicates.")
        if data[0, 0] > 0 or data[-1, 0] < end:
            raise ValueError("Boundary must cover requested interval; no extrapolation.")
        if (data[:, 2] < 0).any() or (data[:, 1] <= -273.15).any():
            raise ValueError("Invalid boundary temperature or moisture.")
        return data, ws.title
    finally:
        wb.close()


def stability(n, substeps):
    if n < 20 or n % 20 or substeps < 1 or int(substeps) != substeps:
        raise ValueError("n must be a multiple of 20; substeps a positive integer.")
    dr, dt = R / n, 1 / substeps
    mu, b = ALPHA * dt / dr**2, H * dr / K
    volume = (R**2 - (R - dr / 2)**2) / 2
    weights = {
        "heat_center": 4 * mu, "heat_interior": 2 * mu,
        "heat_surface": mu * (2 + b * (2 + 1 / n)),
        "moisture_center": 4 * 7e-9 * dt / dr**2,
        "moisture_interior": 2 * 7e-9 * dt / dr**2,
        "moisture_surface": dt / volume * ((R - dr / 2) * 7e-9 / dr + R * HM),
    }
    if max(weights.values()) > 1 + 1e-12:
        raise ValueError(f"Explicit step violates nonnegative weights: {weights}")
    return weights


@njit(cache=True)
def integrate(boundary, n, substeps, end, initial_t, initial_c):
    dr, dt = R / n, 1.0 / substeps
    mu, b = ALPHA * dt / dr**2, H * dr / K
    t, c = np.full(n + 1, initial_t), np.full(n + 1, initial_c)
    tn, cn = t.copy(), c.copy()
    diffusivity, flux, volumes = np.empty(n + 1), np.empty(n), np.empty(n + 1)
    for i in range(n + 1):
        volumes[i] = (min(R, (i + .5) * dr)**2 - max(0., (i - .5) * dr)**2) / 2
    stride = n // 20
    tout, cout = np.empty((end + 1, 21)), np.empty((end + 1, 21))
    balance = np.empty((end + 1, 3))
    tout[0, :], cout[0, :] = initial_t, initial_c
    balance[0, 0], balance[0, 1], balance[0, 2] = initial_c, 0., 0.
    loss, segment, total_volume = 0., 0, R**2 / 2
    for step in range(end * substeps):
        now = step * dt
        while segment + 1 < len(boundary) - 1 and now >= boundary[segment + 1, 0]:
            segment += 1
        fraction = ((now - boundary[segment, 0]) /
                    (boundary[segment + 1, 0] - boundary[segment, 0]))
        ta = boundary[segment, 1] + fraction * (boundary[segment + 1, 1] - boundary[segment, 1])
        ca = boundary[segment, 2] + fraction * (boundary[segment + 1, 2] - boundary[segment, 2])
        tn[0] = t[0] + 4 * mu * (t[1] - t[0])
        for i in range(1, n):
            tn[i] = t[i] + mu * (t[i + 1] - 2*t[i] + t[i - 1] +
                                  (t[i + 1] - t[i - 1]) / (2*i))
        tn[n] = t[n] + 2*mu*(t[n - 1] - t[n]) - mu*b*(2 + 1/n)*(t[n] - ta)
        for i in range(n + 1):
            diffusivity[i] = 7e-9 * np.exp(-.89 / c[i]) if c[i] > 0 else 0.
        for i in range(n):
            # r_(i+1/2) D_(i+1/2) (C_(i+1)-C_i) / dr; inward flux.
            flux[i] = (i + .5) * .5 * (diffusivity[i] + diffusivity[i + 1]) * (c[i + 1] - c[i])
        cn[0] = c[0] + dt * flux[0] / volumes[0]
        for i in range(1, n):
            cn[i] = c[i] + dt * (flux[i] - flux[i - 1]) / volumes[i]
        outward = R * HM * (c[n] - ca)
        cn[n] = c[n] + dt * (-outward - flux[n - 1]) / volumes[n]
        loss += dt * outward
        t, tn = tn, t
        c, cn = cn, c
        if (step + 1) % substeps == 0:
            second = (step + 1) // substeps
            for j in range(21):
                tout[second, j], cout[second, j] = t[j * stride], c[j * stride]
            mean = np.sum(c * volumes) / total_volume
            balance[second, 0] = mean
            balance[second, 1] = loss / total_volume
            balance[second, 2] = mean + loss / total_volume - initial_c
    return tout, cout, balance


def solve(boundary, n=20, substeps=1, end=1800, initial_t=28., initial_c=2.55):
    weights = stability(n, substeps)
    if end < 1 or int(end) != end or initial_c < 0:
        raise ValueError("Invalid duration or initial concentration.")
    if boundary[0, 0] > 0 or boundary[-1, 0] < end:
        raise ValueError("Boundary does not cover requested interval.")
    started = time.perf_counter()
    t, c, balance = integrate(boundary, n, substeps, end, initial_t, initial_c)
    active = boundary[(boundary[:, 0] >= 0) & (boundary[:, 0] <= end)]
    env = np.r_[initial_t, active[:, 1], np.interp([0, end], boundary[:, 0], boundary[:, 1])]
    if not np.isfinite(t).all() or not np.isfinite(c).all():
        raise ArithmeticError("Nonfinite numerical state.")
    if t.min() < env.min() - 1e-9 or t.max() > env.max() + 1e-9:
        raise ArithmeticError("Temperature maximum principle failed.")
    if c.min() < -1e-12 or c.max() > max(initial_c, boundary[:, 2].max()) + 1e-12:
        raise ArithmeticError("Moisture maximum principle failed.")
    checks = {
        "n_intervals": n, "dr_m": R / n, "dt_s": 1 / substeps,
        "runtime_s": time.perf_counter() - started, "weights": weights,
        "temperature_range_C": [float(t.min()), float(t.max())],
        "moisture_range_kg_per_kg": [float(c.min()), float(c.max())],
        "max_mass_balance_residual": float(np.abs(balance[:, 2]).max()),
    }
    return t, c, balance, checks


def differences(first, second):
    dt, dc = np.abs(first[0] - second[0]), np.abs(first[1] - second[1])
    key = np.ix_(TIMES, COLS)
    return {
        "max_temperature_C": float(dt.max()), "max_moisture_kg_per_kg": float(dc.max()),
        "paper_temperature_C": float(dt[key].max()),
        "paper_moisture_kg_per_kg": float(dc[key].max()),
        "same_rounded_paper_temperature": bool(np.array_equal(
            np.round(first[0][key], 4), np.round(second[0][key], 4))),
        "same_rounded_paper_moisture": bool(np.array_equal(
            np.round(first[1][key], 4), np.round(second[1][key], 4))),
    }


def solve_bdf(boundary, n=20, rtol=1e-10, end=1800):
    """Same spatial equations; sparse implicit integration removes the explicit CFL cost."""
    from scipy.integrate import solve_ivp
    from scipy.sparse import diags
    started = time.perf_counter()
    dr, size = R / n, n + 1
    i = np.arange(1, n, dtype=float)
    mu, b = ALPHA / dr**2, H * dr / K
    diagonal = np.full(size, -2 * mu)
    diagonal[0], diagonal[-1] = -4 * mu, -mu * (2 + b * (2 + 1 / n))
    upper = np.r_[4 * mu, mu * (1 + 1 / (2 * i))]
    lower = np.r_[mu * (1 - 1 / (2 * i)), 2 * mu]
    heat_matrix = diags([lower, diagonal, upper], [-1, 0, 1], format="csc")
    heat_source = mu * b * (2 + 1 / n)
    edge = np.arange(n + 1, dtype=float)
    volumes = (np.minimum(R, (edge + .5) * dr)**2 -
               np.maximum(0, (edge - .5) * dr)**2) / 2
    faces = np.arange(n) + .5

    def thermal_rhs(now, y):
        dy = heat_matrix @ y
        dy[-1] += heat_source * np.interp(now, boundary[:, 0], boundary[:, 1])
        return dy

    def diffusivity(c):
        positive = np.maximum(c, 1e-100)
        d = 7e-9 * np.exp(-.89 / positive)
        return d, d * .89 / positive**2

    def moisture_rhs(now, y):
        c = y[:-1]
        d, _ = diffusivity(c)
        flux = faces * .5 * (d[:-1] + d[1:]) * np.diff(c)
        ca = np.interp(now, boundary[:, 0], boundary[:, 2])
        outward = R * HM * (c[-1] - ca)
        dc = (np.r_[flux, -outward] - np.r_[0., flux]) / volumes
        return np.r_[dc, outward / (R**2 / 2)]

    def moisture_jac(now, y):
        c = y[:-1]
        d, dp = diffusivity(c)
        average = .5 * (d[:-1] + d[1:])
        delta = np.diff(c)
        left = faces * (.5 * dp[:-1] * delta - average)
        right = faces * (.5 * dp[1:] * delta + average)
        diag = (np.r_[left, -R*HM] - np.r_[0., right]) / volumes
        return diags([np.r_[-left / volumes[1:], 2*HM/R],
                      np.r_[diag, 0.], np.r_[right / volumes[:-1], 0.]],
                     [-1, 0, 1], shape=(size + 1, size + 1), format="csc")

    times = np.arange(end + 1)
    heat = solve_ivp(thermal_rhs, (0, end), np.full(size, 28.), method="BDF",
                     jac=heat_matrix, rtol=rtol, atol=rtol * .01, t_eval=times, max_step=5.)
    water = solve_ivp(moisture_rhs, (0, end), np.r_[np.full(size, 2.55), 0.],
                      method="BDF", jac=moisture_jac, rtol=rtol, atol=rtol * .01,
                      t_eval=times, max_step=5.)
    if not heat.success or not water.success:
        raise ArithmeticError(f"BDF integration failed: {heat.message}; {water.message}")
    t, c = heat.y[::n//20].T, water.y[:-1:n//20].T
    mean, loss = volumes @ water.y[:-1] / (R**2 / 2), water.y[-1]
    balance = np.column_stack([mean, loss, mean + loss - 2.55])
    active = boundary[(boundary[:, 0] >= 0) & (boundary[:, 0] <= end)]
    env = np.r_[28., active[:, 1], np.interp([0, end], boundary[:, 0], boundary[:, 1])]
    if not np.isfinite(t).all() or not np.isfinite(c).all() or c.min() < -1e-9:
        raise ArithmeticError("BDF finite/nonnegative-state check failed.")
    if t.min() < env.min() - 1e-7 or t.max() > env.max() + 1e-7 or c.max() > 2.55 + 1e-9:
        raise ArithmeticError("BDF range check failed.")
    checks = {
        "method": "BDF with analytic sparse Jacobian, same spatial discretization",
        "n_intervals": n, "dr_m": dr, "rtol": rtol, "atol": rtol * .01,
        "max_step_s": 5., "runtime_s": time.perf_counter() - started,
        "heat_nfev": heat.nfev, "moisture_nfev": water.nfev,
        "temperature_range_C": [float(t.min()), float(t.max())],
        "moisture_range_kg_per_kg": [float(c.min()), float(c.max())],
        "max_mass_balance_residual": float(np.abs(balance[:, 2]).max()),
    }
    return t, c, balance, checks


def write_csv(path, header, values):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(values)


def write_workbook(path, temperature, moisture, template=None):
    if template:
        if path.resolve() == template.resolve():
            raise ValueError("Refusing to overwrite original template.")
        wb = load_workbook(template)
        if set(wb.sheetnames) != {"温度", "水分浓度"}:
            raise ValueError("Template must have 温度 and 水分浓度 sheets.")
        for ws in wb:
            for j in range(21):
                value = ws.cell(1, j + 2).value
                if value is not None and (not isinstance(value, (int, float)) or abs(value - j / 10) > 1e-9):
                    raise ValueError("Unexpected template radius header.")
    else:
        wb = Workbook()
        wb.remove(wb.active)
        for name in ("温度", "水分浓度"):
            wb.create_sheet(name)
    for name, field in (("温度", temperature), ("水分浓度", moisture)):
        ws = wb[name]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
        ws.cell(1, 1, "时间（s）与距离（cm）")
        for j in range(21):
            ws.cell(1, j + 2, j / 10)
        for i in range(len(field)):
            ws.cell(i + 2, 1, i)
            for j in range(21):
                cell = ws.cell(i + 2, j + 2, float(f"{field[i, j]:.4f}"))
                cell.number_format = "0.0000"
        if not template:
            ws.freeze_panes = "B2"
            ws.column_dimensions["A"].width = 25
            for row in ws:
                for cell in row:
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(horizontal="center")
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="DCEAF7")
                cell.font = Font(name="Arial", size=10, bold=True)
            for j in range(2, 23):
                ws.column_dimensions[ws.cell(1, j).column_letter].width = 12
    wb.save(path)
    wb.close()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for name, field in (("温度", temperature), ("水分浓度", moisture)):
            ws = wb[name]
            if ws.max_row != len(field) + 1 or ws.max_column != 22:
                raise AssertionError("Output workbook shape mismatch.")
            rows = list(ws.values)
            np.testing.assert_allclose(rows[0][1:], np.arange(21) / 10, atol=1e-12)
            np.testing.assert_array_equal([r[0] for r in rows[1:]], np.arange(len(field)))
            np.testing.assert_allclose(np.asarray([r[1:] for r in rows[1:]], float),
                                       np.round(field, 4), atol=5e-12, rtol=0)
    finally:
        wb.close()


def tex_table(path, field, caption, label):
    lines = [r"\begin{table}[H]", r"  \centering", rf"  \caption{{{caption}}}",
             rf"  \label{{{label}}}", r"  \begin{tabular}{rrrrrr}", r"    \toprule",
             r"    时间（s） & $0\,\mathrm{cm}$ & $0.5\,\mathrm{cm}$ & $1\,\mathrm{cm}$ & $1.5\,\mathrm{cm}$ & $2\,\mathrm{cm}$ \\",
             r"    \midrule"]
    for second in TIMES:
        values = " & ".join(f"{field[second, j]:.4f}" for j in COLS)
        lines.append(f"    {second} & {values} " + r"\\")
    lines += [r"    \bottomrule", r"  \end{tabular}", r"\end{table}", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def plot_results(destination, t, c, boundary):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    available = {f.name for f in font_manager.fontManager.ttflist}
    font = next((f for f in ["Microsoft YaHei", "SimHei", "SimSun", "Noto Sans CJK SC"]
                 if f in available), None)
    if not font:
        raise RuntimeError("A Chinese font is required for paper figures.")
    plt.rcParams.update({"font.family": font, "font.size": 10, "axes.unicode_minus": False,
                         "pdf.fonttype": 42, "axes.spines.top": False, "axes.spines.right": False})
    radii, minutes = np.arange(21) / 10, np.arange(len(t)) / 60
    fig, axes = plt.subplots(2, 2, figsize=(9, 6.4), layout="constrained")
    colors = ["#174A70", "#2584A6", "#77AABD", "#D88A41", "#963E3E"]
    for field, row, ylabel in [(t, 0, "温度（°C）"), (c, 1, r"干基含水率（$\frac{kg}{kg}$）")]:
        for second, color in zip([100, 300, 600, 1200, 1800], colors):
            axes[row, 0].plot(radii, field[second], color=color, label=f"{second} s")
        for j, color in zip(COLS, colors):
            axes[row, 1].plot(minutes, field[:, j], color=color, label=f"{j / 10:g} cm")
        axes[row, 0].set(xlabel="距中心轴的径向距离（cm）", ylabel=ylabel)
        axes[row, 1].set(xlabel="时间（min）", ylabel=ylabel)
        for ax in axes[row]:
            ax.grid(alpha=.18)
            ax.legend(frameon=False, fontsize=8, ncol=2)
    env = np.interp(np.arange(len(t)), boundary[:, 0], boundary[:, 1])
    axes[0, 1].plot(minutes, env, "--", color="#333333", lw=1, label="烘房")
    axes[0, 1].legend(frameon=False, fontsize=8, ncol=2)
    fig.savefig(destination / "q1_fields.pdf", bbox_inches="tight")
    fig.savefig(destination / "q1_fields.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--boundary", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--template", type=Path)
    parser.add_argument("--method", choices=["explicit", "bdf"], default="bdf")
    parser.add_argument("--max-n", type=int, default=5120)
    parser.add_argument("--tolerance", type=float, default=5e-5)
    parser.add_argument("--output", type=Path, default=ROOT / "results" / "q1")
    parser.add_argument("--sync-paper", action="store_true")
    args = parser.parse_args()
    ratio = args.max_n // 20
    if args.max_n < 40 or args.max_n % 20 or ratio & (ratio - 1):
        parser.error("--max-n must be 40, 80, 160, 320, 640, ...")
    if args.tolerance <= 0:
        parser.error("--tolerance must be positive.")
    if args.sync_paper and args.output.resolve() != (ROOT / "results" / "q1").resolve():
        parser.error("--sync-paper requires default results directory.")
    boundary, sheet = read_boundary(args.boundary, args.sheet)
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    start, records, previous = time.perf_counter(), [], None
    n, substeps, converged = 20, 1, False
    while n <= args.max_n:
        current = (solve(boundary, n, substeps) if args.method == "explicit"
                   else solve_bdf(boundary, n))
        record = dict(current[3])
        if previous is not None:
            d = differences(current, previous)
            record["difference_from_previous"] = d
            converged = d["max_temperature_C"] < args.tolerance and d["max_moisture_kg_per_kg"] < args.tolerance
        records.append(record)
        print(json.dumps(record), flush=True)
        if converged or n == args.max_n:
            break
        previous = current
        n, substeps = n * 2, substeps * 4
    final = (solve(boundary, n, substeps * 2) if args.method == "explicit"
             else solve_bdf(boundary, n, rtol=1e-11))
    time_difference = differences(current, final)
    print("time_refinement " + json.dumps(time_difference), flush=True)
    t, c, balance, checks = final
    valid = converged and time_difference["max_temperature_C"] < args.tolerance and time_difference["max_moisture_kg_per_kg"] < args.tolerance
    if not valid:
        # Failed refinements must never replace the last successful result set.
        out = out / "diagnostics" / f"{args.method}_n{n}"
        out.mkdir(parents=True, exist_ok=True)
    metadata = {
        "status": "NUMERICAL_CHECKS_PASSED" if valid else "REFINEMENT_REQUIRED",
        "boundary_source": str(args.boundary.resolve()), "boundary_sheet": sheet,
        "boundary_samples": len(boundary), "boundary_time_range_s": boundary[[0, -1], 0].tolist(),
        "boundary_interpolation": "piecewise linear, no extrapolation",
        "model": "Q1: fixed radius, independent heat and effective dry-basis diffusion",
        "assumptions": ["ignore end transport", "constant dry skeleton",
                        "no latent-heat feedback", "boundary concentration on problem scale"],
        "parameters": {"rho": 820, "cp": 2600, "k": K, "h": H, "hm": HM,
                       "radius_m": R, "T0_C": 28, "C0_kg_per_kg": 2.55},
        "refinement": records, "temporal_refinement_difference": time_difference,
        "final_run": checks, "absolute_difference_threshold": args.tolerance,
        "accuracy_statement": "Empirical grid differences, not rigorous error bounds. Four-decimal formatting alone does not prove four-decimal accuracy.",
        "template": str(args.template.resolve()) if args.template else "Recreated from problem PDF, Appendix 1",
        "environment": {"python": sys.version, "executable": sys.executable,
                        "numpy": np.__version__, "platform": platform.platform()},
        "runtime_s": time.perf_counter() - start,
        "end_temperature_C": t[-1, COLS].tolist(), "end_moisture_kg_per_kg": c[-1, COLS].tolist(),
        "end_mean_moisture": float(balance[-1, 0]),
    }
    times, radii = np.arange(1801), np.arange(21) / 10
    np.savez_compressed(out / "q1_fields.npz", time_s=times, radius_cm=radii,
                        temperature_C=t, moisture_kg_per_kg=c, balance=balance)
    write_csv(out / "boundary_used.csv", ["time_s", "temperature_C", "moisture_kg_per_kg"],
              np.column_stack([times, np.interp(times, boundary[:, 0], boundary[:, 1]),
                               np.interp(times, boundary[:, 0], boundary[:, 2])]))
    for name, field in (("temperature", t), ("moisture", c)):
        write_csv(out / f"q1_{name}_full.csv", ["time_s"] + [f"r_{r:.1f}_cm" for r in radii],
                  np.column_stack([times, field]))
        write_csv(out / f"q1_{name}_paper.csv", ["time_s"] + [f"r_{r:g}_cm" for r in radii[COLS]],
                  np.column_stack([TIMES, field[np.ix_(TIMES, COLS)]]))
    write_csv(out / "mass_balance.csv", ["time_s", "mean_moisture", "cumulative_loss", "residual"],
              np.column_stack([times, balance]))
    workbook = out / ("result1.xlsx" if valid else "q1_provisional.xlsx")
    write_workbook(workbook, t, c, args.template)
    metadata["workbook_cell_validation"] = "PASS: every numerical output and coordinate reread"
    (out / "q1_validation.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")
    plot_results(out, t, c, boundary)
    if args.sync_paper:
        if not valid:
            raise SystemExit("Refinement failed: provisional outputs retained, paper unchanged.")
        table_dir = ROOT / "paper" / "tables"
        tex_table(table_dir / "q1_temperature.tex", t,
                  r"30 分钟内药材的温度（\({}^\circ\mathrm{C}\)）", "tab:q1-temperature")
        tex_table(table_dir / "q1_moisture.tex", c,
                  r"30 分钟内药材的干基含水率（\(\frac{\mathrm{kg}}{\mathrm{kg}}\)）", "tab:q1-moisture")
        import shutil
        shutil.copyfile(out / "q1_fields.pdf", ROOT / "paper" / "figures" / "q1_fields.pdf")
    print(json.dumps({"status": metadata["status"], "workbook": str(workbook.resolve()),
                      "T_1800": metadata["end_temperature_C"],
                      "C_1800": metadata["end_moisture_kg_per_kg"]}), flush=True)
    if not valid:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
