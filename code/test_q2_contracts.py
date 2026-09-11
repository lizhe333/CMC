"""Targeted contracts for the question 2 coupled solver.

These tests use synthetic fields for numerical contracts.  They do not assert
competition results and do not replace the full spatial/time refinement.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

CODE_DIR = Path(__file__).resolve().parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

import q2_solver as q2


class Q2Contracts(unittest.TestCase):
    def test_profile_seconds_include_all_six_paper_times(self):
        """The radial panels must include 0.5, 1, 1.5, 2, 2.5 and 3 h."""
        self.assertEqual(
            q2.PROFILE_SECONDS.tolist(),
            [1800, 3600, 5400, 7200, 9000, 10800],
        )

    def test_appendix3_initial_properties(self):
        rho, cp, k, diffusivity = q2.material_properties(2.55, 28.0)
        self.assertAlmostEqual(float(rho), 976.4, places=10)
        self.assertAlmostEqual(float(cp), 3415.295774647887, places=10)
        self.assertAlmostEqual(float(k), 0.48295774647887324, places=12)
        self.assertAlmostEqual(float(diffusivity), 5.641680373025664e-9, places=18)

    def test_uniform_equilibrium_field(self):
        boundary = np.array([[0.0, 28.0, 2.55], [100.0, 28.0, 2.55]])
        result = q2.integrate_case(boundary, n=20, end=20, chunk_seconds=20)
        np.testing.assert_allclose(result["temperature_C"], 28.0, atol=2e-11, rtol=0.0)
        np.testing.assert_allclose(result["moisture_kg_per_kg"], 2.55, atol=2e-11, rtol=0.0)
        self.assertLess(float(np.max(np.abs(result["mass_balance_residual"]))), 2e-11)

    def test_constant_property_cylindrical_manufactured_solution(self):
        """T=r^2 has the exact radial operator 4*k/(rho*cp)."""
        n = 40
        _, nodes, _, _ = q2.radial_geometry(n)
        moisture = np.full(n + 1, 2.55)
        temperature = nodes * nodes
        rho, cp, k, _ = q2.material_properties(moisture, temperature)
        surface_environment = q2.RADIUS_M**2 + 2.0 * float(k[0]) * q2.RADIUS_M / q2.H_BASE
        boundary = np.array(
            [[0.0, surface_environment, 2.55], [100.0, surface_environment, 2.55]]
        )
        state = np.empty(2 * (n + 1))
        state[0::2] = temperature
        state[1::2] = moisture
        rhs = q2.assemble_rhs(state, 0.0, boundary, n)
        expected = 4.0 * float(k[0]) / (rho * cp)
        np.testing.assert_allclose(rhs[0::2], expected, atol=2e-12, rtol=2e-12)
        np.testing.assert_allclose(rhs[1::2], 0.0, atol=2e-12, rtol=0.0)

    def test_internal_flux_closure(self):
        n = 40
        _, _, _, volumes = q2.radial_geometry(n)
        temperature = np.linspace(28.0, 42.0, n + 1)
        moisture = np.linspace(2.5, 1.2, n + 1)
        state = np.empty(2 * (n + 1))
        state[0::2] = temperature
        state[1::2] = moisture
        boundary = np.array([[0.0, 45.0, 0.05], [100.0, 45.0, 0.05]])
        rhs = q2.assemble_rhs(state, 0.0, boundary, n)
        rho, cp, _, _ = q2.material_properties(moisture, temperature)
        heat_surface = -q2.RADIUS_M * q2.H_BASE * (temperature[-1] - 45.0)
        moisture_surface = -q2.RADIUS_M * q2.HM_BASE * (moisture[-1] - 0.05)
        heat_closed = np.dot(rhs[0::2] * rho * cp, volumes)
        moisture_closed = np.dot(rhs[1::2], volumes)
        self.assertAlmostEqual(float(heat_closed), float(heat_surface), places=12)
        self.assertAlmostEqual(float(moisture_closed), float(moisture_surface), places=12)

    def test_cross_dependency_is_present(self):
        n = 20
        _, nodes, _, _ = q2.radial_geometry(n)
        state = np.empty(2 * (n + 1))
        state[0::2] = np.linspace(28.0, 42.0, n + 1)
        state[1::2] = np.linspace(2.5, 1.2, n + 1)
        boundary = np.array([[0.0, 45.0, 0.05], [100.0, 45.0, 0.05]])
        base = q2.assemble_rhs(state, 0.0, boundary, n)
        temperature_perturbed = state.copy()
        temperature_perturbed[2] += 1.0
        moisture_perturbed = state.copy()
        moisture_perturbed[1] += 0.1
        changed_by_temperature = q2.assemble_rhs(
            temperature_perturbed, 0.0, boundary, n
        )[1::2] - base[1::2]
        changed_by_moisture = q2.assemble_rhs(
            moisture_perturbed, 0.0, boundary, n
        )[0::2] - base[0::2]
        self.assertGreater(float(np.max(np.abs(changed_by_temperature))), 1e-14)
        self.assertGreater(float(np.max(np.abs(changed_by_moisture))), 1e-12)
        sparsity = q2.jacobian_sparsity(n)
        self.assertEqual(sparsity.shape, (2 * (n + 1), 2 * (n + 1)))
        # In interleaved ordering, a node-neighbour 2x2 block reaches scalar
        # offsets +/-3.  The mask must not silently drop T_i<->C_(i+1)
        # cross-dependencies.
        self.assertEqual(float(sparsity[0, 3]), 1.0)
        self.assertEqual(float(sparsity[1, 2]), 1.0)

    def test_domain_and_extrapolation_rejection(self):
        with self.assertRaises(q2.Q2DomainError):
            q2.material_properties(0.0, 28.0)
        with self.assertRaises(ValueError):
            q2.read_boundary(q2.DEFAULT_BOUNDARY, end=14401)
        boundary = np.array([[0.0, 28.0, 2.55], [100.0, 28.0, 2.55]])
        with self.assertRaises(ValueError):
            q2.environment_at(101.0, boundary)

    def test_workbook_contract_and_restart_payload(self):
        boundary = np.array([[0.0, 28.0, 2.55], [100.0, 28.0, 2.55]])
        result = q2.integrate_case(boundary, n=20, end=5, chunk_seconds=5)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            workbook = output / "candidate.xlsx"
            # The official result2.xlsx is intentionally populated after the
            # formal run.  Contract tests must always start from the preserved
            # five-row source template rather than treating that result as a
            # template with a fixed max_row.
            template = q2.TEMPLATE_BACKUP if q2.TEMPLATE_BACKUP.exists() else q2.DEFAULT_TEMPLATE
            q2.write_workbook(
                workbook,
                result["temperature_C"],
                result["moisture_kg_per_kg"],
                template,
                end=5,
            )
            report = q2.validate_workbook(
                workbook,
                result["temperature_C"],
                result["moisture_kg_per_kg"],
                end=5,
                template=template,
            )
            self.assertTrue(report["sheets"]["温度"]["four_decimal_match"])
            self.assertEqual(report["sheets"]["水分浓度"]["rows"], 5)
            restart = output / "q2_restart.npz"
            np.savez_compressed(
                restart,
                final_time_s=result["final_time_s"],
                radius_nodes_m=result["radius_nodes_m"],
                control_volume_weights_m2=result["control_volume_weights_m2"],
                temperature_final_C=result["temperature_final_C"],
                moisture_final_kg_per_kg=result["moisture_final_kg_per_kg"],
                n_intervals=result["n_intervals"],
                h_W_m2K=result["h_W_m2K"],
                hm_m_s=result["hm_m_s"],
            )
            with np.load(restart, allow_pickle=False) as payload:
                self.assertEqual(payload["temperature_final_C"].shape, (21,))
                self.assertEqual(payload["moisture_final_kg_per_kg"].shape, (21,))
                self.assertEqual(payload["radius_nodes_m"].shape, (21,))
                self.assertEqual(payload["control_volume_weights_m2"].shape, (21,))

    def test_workbook_template_style_inheritance_and_widths(self):
        """G:V inherit the template's F style and width; A:F widths stay fixed."""
        boundary = np.array([[0.0, 28.0, 2.55], [100.0, 28.0, 2.55]])
        result = q2.integrate_case(boundary, n=20, end=5, chunk_seconds=5)
        template = q2.TEMPLATE_BACKUP if q2.TEMPLATE_BACKUP.exists() else q2.DEFAULT_TEMPLATE
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / "styled_candidate.xlsx"
            q2.write_workbook(
                workbook,
                result["temperature_C"],
                result["moisture_kg_per_kg"],
                template,
                end=5,
            )
            style_report = q2.validate_workbook_template_styles(workbook, template, end=5)
            for sheet_report in style_report["sheets"].values():
                self.assertTrue(sheet_report["template_style_match"])
                self.assertTrue(sheet_report["existing_widths_preserved"])
            output_wb = load_workbook(workbook, data_only=False)
            source_wb = load_workbook(template, data_only=False)
            try:
                for sheet_name in ("温度", "水分浓度"):
                    output_ws = output_wb[sheet_name]
                    source_ws = source_wb[sheet_name]
                    for column in range(1, 7):
                        letter = output_ws.cell(1, column).column_letter
                        self.assertEqual(
                            output_ws.column_dimensions[letter].width,
                            source_ws.column_dimensions[letter].width,
                        )
                    self.assertEqual(
                        output_ws.column_dimensions["G"].width,
                        source_ws.column_dimensions["F"].width,
                    )
                    for column in range(7, 23):
                        output_cell = output_ws.cell(2, column)
                        source_cell = source_ws.cell(2, 6)
                        self.assertEqual(copy(output_cell.font), copy(source_cell.font))
                        self.assertEqual(copy(output_cell.alignment), copy(source_cell.alignment))
                        self.assertEqual(copy(output_cell.fill), copy(source_cell.fill))
                        self.assertEqual(copy(output_cell.border), copy(source_cell.border))
                        self.assertEqual(output_cell.number_format, "0.0000")
            finally:
                source_wb.close()
                output_wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
