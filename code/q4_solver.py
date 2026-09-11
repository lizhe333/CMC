"""Question 4 moving-radius heat--moisture solver and development interfaces.

The module implements the normalized material coordinate
``xi = r / R(t)`` for the axisymmetric effective model in Q4_MODELING.md.
The radius is an external input.  It is never appended to the BDF state and
there is no second advection term after the material-coordinate transform.

This file intentionally does not write the official Excel workbook.  It
produces unrounded field/restart/endpoint evidence and provides the numerical
interfaces used by the workbook and paper stages after the audits pass.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import platform
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.sparse import coo_matrix, diags


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BOUNDARY = ROOT.parent / "附件1.xlsx"
DEFAULT_RADIUS = ROOT.parent / "附件2.xlsx"
DEFAULT_OUTPUT = ROOT / "results" / "q4" / "dev"

RADIUS_INITIAL_M = 0.02
LENGTH_M = 0.25
INITIAL_T_C = 28.0
INITIAL_C = 2.55
H_BASE = 25.0
HM_BASE = 8.0e-7
SWITCH_S = 14400.0
RADIUS_MEASURED_END_S = 259200.0
RADIUS_FINAL_M = 0.01198
PLATEAU_T_C = 50.0
PLATEAU_C = 0.05
THRESHOLD_C = 0.15
OUTPUT_STEP_S = 60
PAPER_STEP_S = 6 * 3600
MAX_END_S = 14 * 24 * 3600
DEFAULT_CHUNK_S = 600.0
EVENT_BRACKET_TOL_S = 1.0
EVENT_CONCENTRATION_TOL = 1.0e-9
MONOTONICITY_TOL = 1.0e-8
FIELD_TOL_T_C = 5.0e-5
FIELD_TOL_C = 5.0e-5
RADIUS_MASK_ABS_M = 1.0e-12
RADIUS_MASK_REL = 1.0e-10

BASE_RTL = 1.0e-9
BASE_ATOL = 1.0e-11
BASE_MAX_STEP_BEFORE_S = 5.0
BASE_MAX_STEP_AFTER_S = 60.0
TIGHT_RTL = 1.0e-10
TIGHT_ATOL = 1.0e-12
TIGHT_MAX_STEP_BEFORE_S = 2.5
TIGHT_MAX_STEP_AFTER_S = 30.0
EVENT_RECOVERY_MAX_RETRIES = 6


class Q4DomainError(ValueError):
    """The numerical state is outside the empirical-property domain."""


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path: str | Path, value: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(_json_safe(value), ensure_ascii=False, indent=2), encoding="utf-8")


def _finite_vector(name: str, value: Any) -> np.ndarray:
    result = np.asarray(value, dtype=float)
    if not np.isfinite(result).all():
        raise ValueError(f"{name} contains a nonfinite value.")
    return result


@dataclass
class RadiusModel:
    """Measured-radius interpolator with explicit post-measurement behavior."""

    times_s: np.ndarray
    radii_m: np.ndarray
    interpolation: str = "pchip"
    post_mode: str = "plateau"

    def __post_init__(self) -> None:
        self.times_s = _finite_vector("radius times", self.times_s).reshape(-1)
        self.radii_m = _finite_vector("radius values", self.radii_m).reshape(-1)
        if len(self.times_s) < 2 or len(self.times_s) != len(self.radii_m):
            raise ValueError("At least two radius records with equal lengths are required.")
        if np.any(np.diff(self.times_s) <= 0.0):
            raise ValueError("Radius times must be strictly increasing.")
        if np.any(self.radii_m <= 0.0) or np.any(np.diff(self.radii_m) > 1.0e-12):
            raise ValueError("Radius must be positive and nonincreasing.")
        if self.interpolation not in {"pchip", "linear"}:
            raise ValueError("interpolation must be pchip or linear.")
        if self.post_mode not in {"plateau", "slow1pct"}:
            raise ValueError("post_mode must be plateau or slow1pct.")
        self.times_s = self.times_s.astype(float)
        self.radii_m = self.radii_m.astype(float)
        if self.interpolation == "pchip":
            self._pchip = PchipInterpolator(self.times_s, self.radii_m, extrapolate=False)
            self._pchip_derivative = self._pchip.derivative()
        else:
            self._pchip = None
            self._pchip_derivative = None

    @property
    def start_s(self) -> float:
        return float(self.times_s[0])

    @property
    def end_s(self) -> float:
        return float(self.times_s[-1])

    @property
    def final_radius_m(self) -> float:
        return float(self.radii_m[-1])

    def _inside_value(self, values: np.ndarray) -> np.ndarray:
        if self.interpolation == "pchip":
            return np.asarray(self._pchip(values), dtype=float)
        return np.interp(values, self.times_s, self.radii_m)

    def _inside_derivative(self, values: np.ndarray) -> np.ndarray:
        if self.interpolation == "pchip":
            return np.asarray(self._pchip_derivative(values), dtype=float)
        indices = np.searchsorted(self.times_s, values, side="right") - 1
        indices = np.clip(indices, 0, len(self.times_s) - 2)
        slopes = np.diff(self.radii_m) / np.diff(self.times_s)
        return slopes[indices]

    def __call__(self, t: float | np.ndarray) -> float | np.ndarray:
        values = np.asarray(t, dtype=float)
        if not np.isfinite(values).all() or np.any(values < self.start_s - 1.0e-9):
            raise ValueError("Radius requested before the measured interval.")
        # Values in the tiny input tolerance are explicitly snapped to the
        # first measured time; passing them to PCHIP would otherwise return
        # NaN despite the guard above.
        flat = np.maximum(values.reshape(-1), self.start_s)
        result = np.empty_like(flat)
        inside = flat <= self.end_s + 1.0e-9
        if np.any(inside):
            result[inside] = self._inside_value(np.minimum(flat[inside], self.end_s))
        if np.any(~inside):
            delta = flat[~inside] - self.end_s
            if self.post_mode == "plateau":
                result[~inside] = self.final_radius_m
            else:
                result[~inside] = self.final_radius_m * (
                    1.0 - 0.01 * (1.0 - np.exp(-delta / 86400.0))
                )
        result = result.reshape(values.shape)
        if not np.isfinite(result).all():
            raise ValueError("Radius interpolation returned a nonfinite value.")
        return float(result) if values.ndim == 0 else result

    def derivative(self, t: float | np.ndarray) -> float | np.ndarray:
        values = np.asarray(t, dtype=float)
        if not np.isfinite(values).all() or np.any(values < self.start_s - 1.0e-9):
            raise ValueError("Radius derivative requested before the measured interval.")
        flat = np.maximum(values.reshape(-1), self.start_s)
        result = np.empty_like(flat)
        inside = flat <= self.end_s + 1.0e-9
        if np.any(inside):
            result[inside] = self._inside_derivative(np.minimum(flat[inside], self.end_s))
        if np.any(~inside):
            if self.post_mode == "plateau":
                result[~inside] = 0.0
            else:
                delta = flat[~inside] - self.end_s
                result[~inside] = -self.final_radius_m * 0.01 / 86400.0 * np.exp(-delta / 86400.0)
        result = result.reshape(values.shape)
        if not np.isfinite(result).all():
            raise ValueError("Radius derivative returned a nonfinite value.")
        return float(result) if values.ndim == 0 else result


def read_radius(path: str | Path = DEFAULT_RADIUS, *, interpolation: str = "pchip", post_mode: str = "plateau") -> tuple[RadiusModel, dict[str, Any]]:
    """Read Attachment 2 without implicit Excel or interpolator extrapolation."""
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.values)
    finally:
        workbook.close()
    records: list[tuple[float, float]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if row is None or all(value is None for value in row):
            continue
        if len(row) < 2 or any(value is None or isinstance(value, bool) for value in row[:2]):
            raise ValueError(f"Invalid radius row {row_number}.")
        try:
            records.append((float(row[0]), float(row[1]) * 0.01))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid radius row {row_number}.") from exc
    times, radii = np.asarray(records, dtype=float).T
    model = RadiusModel(times, radii, interpolation=interpolation, post_mode=post_mode)
    dense_times = np.linspace(model.start_s, model.end_s, 2001)
    dense_radius = np.asarray(model(dense_times), dtype=float)
    dense_derivative = np.asarray(model.derivative(dense_times), dtype=float)
    metadata = {
        "path": str(path),
        "sheet": sheet.title,
        "record_count": int(len(records)),
        "start_s": model.start_s,
        "end_s": model.end_s,
        "start_radius_m": float(model.radii_m[0]),
        "end_radius_m": float(model.radii_m[-1]),
        "interpolation": interpolation,
        "post_mode": post_mode,
        "dense_min_radius_m": float(np.min(dense_radius)),
        "dense_max_radius_m": float(np.max(dense_radius)),
        "dense_min_derivative_m_s": float(np.min(dense_derivative)),
        "dense_max_derivative_m_s": float(np.max(dense_derivative)),
        "outside_evaluation": "explicit post_mode; no interpolator extrapolation",
    }
    if metadata["record_count"] != 145:
        raise ValueError(f"Attachment 2 expected 145 records, got {metadata['record_count']}.")
    return model, metadata


def read_boundary(path: str | Path = DEFAULT_BOUNDARY) -> tuple[np.ndarray, dict[str, Any]]:
    """Read Attachment 1 and require coverage through the 4-hour switch."""
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.values)
    finally:
        workbook.close()
    header = None
    for index, row in enumerate(rows[:20]):
        if row is None or len(row) < 3:
            continue
        normalized = [str(value).replace(" ", "") if value is not None else "" for value in row[:3]]
        if all(token in normalized[j] for j, token in enumerate(("时间", "温度", "水分浓度"))):
            header = index
            break
    if header is None:
        raise ValueError("Attachment 1 columns 时间, 温度, 水分浓度 were not found.")
    records: list[list[float]] = []
    for row_number, row in enumerate(rows[header + 1 :], start=header + 2):
        if row is None or all(value is None for value in row):
            continue
        if len(row) < 3 or any(value is None or isinstance(value, bool) for value in row[:3]):
            raise ValueError(f"Invalid boundary row {row_number}.")
        try:
            records.append([float(row[0]), float(row[1]), float(row[2])])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid boundary row {row_number}.") from exc
    boundary = np.asarray(records, dtype=float)
    if boundary.ndim != 2 or len(boundary) < 2 or not np.isfinite(boundary).all():
        raise ValueError("Attachment 1 has no finite boundary table.")
    if np.any(np.diff(boundary[:, 0]) <= 0.0):
        raise ValueError("Attachment 1 times must be strictly increasing.")
    if boundary[0, 0] > 0.0 or boundary[-1, 0] < SWITCH_S:
        raise ValueError("Attachment 1 does not cover the 4-hour switch.")
    if np.any(boundary[:, 1] <= -273.15) or np.any(boundary[:, 2] < 0.0):
        raise ValueError("Attachment 1 contains a nonphysical boundary value.")
    metadata = {
        "path": str(path),
        "sheet": sheet.title,
        "record_count": int(len(boundary)),
        "start_s": float(boundary[0, 0]),
        "end_s": float(boundary[-1, 0]),
        "interpolation": "piecewise linear on measured interval; no extrapolation",
        "post_switch": {"temperature_C": PLATEAU_T_C, "moisture_kg_per_kg": PLATEAU_C},
    }
    return boundary, metadata


def measured_environment(boundary: np.ndarray) -> Callable[[float], tuple[float, float]]:
    boundary = np.asarray(boundary, dtype=float)

    def evaluate(t: float) -> tuple[float, float]:
        if t < boundary[0, 0] - 1.0e-8 or t > boundary[-1, 0] + 1.0e-8:
            raise ValueError("Measured boundary requested outside its explicit range.")
        checked = min(max(float(t), float(boundary[0, 0])), float(boundary[-1, 0]))
        return (
            float(np.interp(checked, boundary[:, 0], boundary[:, 1])),
            float(np.interp(checked, boundary[:, 0], boundary[:, 2])),
        )

    return evaluate


def plateau_environment(*, temperature_C: float = PLATEAU_T_C, moisture_kg_per_kg: float = PLATEAU_C) -> Callable[[float], tuple[float, float]]:
    if temperature_C <= -273.15 or moisture_kg_per_kg < 0.0:
        raise ValueError("Invalid plateau environment.")

    def evaluate(_t: float) -> tuple[float, float]:
        return float(temperature_C), float(moisture_kg_per_kg)

    return evaluate


def material_properties(C: np.ndarray | float, T_C: np.ndarray | float, model: str = "appendix4") -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Evaluate all four local properties for Appendix 3 or Appendix 4."""
    c = np.asarray(C, dtype=float)
    t = np.asarray(T_C, dtype=float)
    if model not in {"appendix3", "appendix4"}:
        raise ValueError("model must be appendix3 or appendix4.")
    if not np.isfinite(c).all() or not np.isfinite(t).all():
        raise Q4DomainError("Nonfinite temperature or moisture state.")
    if np.any(c <= 0.0):
        raise Q4DomainError("Moisture must remain strictly positive.")
    T_K = t + 273.15
    if np.any(T_K <= 0.0):
        raise Q4DomainError("Absolute temperature must remain positive.")
    if model == "appendix4":
        rho = 760.0 + 90.0 * c
        cp = 1850.0 + 2150.0 * c / (c + 1.0)
        k = 0.12 + 0.20 * c / (c + 1.0)
        D = 4.2e-4 * np.exp(-0.30 / c) * np.exp(-3850.0 / T_K)
    else:
        rho = 650.0 + 128.0 * c
        cp = 1450.0 + 2736.0 * c / (c + 1.0)
        k = 0.21 + 0.38 * c / (c + 1.0)
        D = 2.4e-3 * np.exp(-0.45 / c) * np.exp(-3850.0 / T_K)
    if not all(np.isfinite(value).all() for value in (rho, cp, k, D)):
        raise Q4DomainError("Empirical property evaluation is nonfinite.")
    return rho, cp, k, D


def property_derivatives(C: np.ndarray | float, T_C: np.ndarray | float, model: str = "appendix4") -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    rho, cp, k, D = material_properties(C, T_C, model=model)
    c = np.asarray(C, dtype=float)
    T_K = np.asarray(T_C, dtype=float) + 273.15
    if model == "appendix4":
        rho_C = np.full_like(c, 90.0, dtype=float)
        cp_C = 2150.0 / (c + 1.0) ** 2
        k_C = 0.20 / (c + 1.0) ** 2
        D_C = 0.30 * D / c**2
    else:
        rho_C = np.full_like(c, 128.0, dtype=float)
        cp_C = 2736.0 / (c + 1.0) ** 2
        k_C = 0.38 / (c + 1.0) ** 2
        D_C = 0.45 * D / c**2
    D_T = 3850.0 * D / T_K**2
    return rho_C, cp_C, k_C, D_C, D_T


def radial_geometry(n: int) -> tuple[float, np.ndarray, np.ndarray, np.ndarray]:
    if int(n) != n or n < 4:
        raise ValueError("n must be an integer of at least 4 intervals.")
    n = int(n)
    dxi = 1.0 / n
    nodes = np.arange(n + 1, dtype=float) * dxi
    faces = (np.arange(n, dtype=float) + 0.5) * dxi
    edge = np.arange(n + 1, dtype=float)
    left = np.maximum(0.0, (edge - 0.5) * dxi)
    right = np.minimum(1.0, (edge + 0.5) * dxi)
    weights = (right**2 - left**2) / 2.0
    if not np.isclose(np.sum(weights), 0.5, rtol=0.0, atol=1.0e-15):
        raise ArithmeticError("Normalized cylindrical weights do not close.")
    return dxi, nodes, faces, weights


def jacobian_sparsity(n: int):
    size = 2 * (n + 1)
    offsets = list(range(-3, 4))
    diagonals = [np.ones(size - abs(offset), dtype=float) for offset in offsets]
    return diags(diagonals, offsets=offsets, shape=(size, size), format="csc")


def interleaved_state(temperature_C: np.ndarray, moisture: np.ndarray) -> np.ndarray:
    temperature_C = np.asarray(temperature_C, dtype=float).reshape(-1)
    moisture = np.asarray(moisture, dtype=float).reshape(-1)
    if temperature_C.shape != moisture.shape:
        raise ValueError("Temperature and moisture restart vectors must have equal lengths.")
    result = np.empty(2 * len(temperature_C), dtype=float)
    result[0::2] = temperature_C
    result[1::2] = moisture
    return result


def initial_state(n: int, temperature_C: float = INITIAL_T_C, moisture: float = INITIAL_C) -> np.ndarray:
    if moisture <= 0.0 or temperature_C <= -273.15:
        raise Q4DomainError("Invalid initial state.")
    material_properties(np.array([moisture]), np.array([temperature_C]), model="appendix4")
    return interleaved_state(np.full(n + 1, temperature_C), np.full(n + 1, moisture))


def _physical_from_log(z_state: np.ndarray, n: int) -> tuple[np.ndarray, np.ndarray]:
    z_state = np.asarray(z_state, dtype=float)
    if z_state.size != 2 * (n + 1) or not np.isfinite(z_state).all():
        raise Q4DomainError("Invalid transformed state size or finiteness.")
    with np.errstate(over="raise", invalid="raise", under="ignore"):
        try:
            moisture = np.exp(z_state[1::2])
        except FloatingPointError as exc:
            raise Q4DomainError("log(C) overflowed; no clipping was applied.") from exc
    if not np.isfinite(moisture).all() or np.any(moisture <= 0.0):
        raise Q4DomainError("log(C) did not map to positive moisture.")
    physical = np.asarray(z_state, dtype=float).copy()
    physical[1::2] = moisture
    if np.any(physical[0::2] + 273.15 <= 0.0):
        raise Q4DomainError("Transformed temperature is outside the physical domain.")
    return physical, moisture


def _to_log_state(state: np.ndarray, n: int, model: str = "appendix4") -> np.ndarray:
    state = np.asarray(state, dtype=float).copy()
    if state.size != 2 * (n + 1):
        raise ValueError("State size and grid disagree.")
    material_properties(state[1::2], state[0::2], model=model)
    with np.errstate(divide="raise", invalid="raise"):
        try:
            state[1::2] = np.log(state[1::2])
        except FloatingPointError as exc:
            raise Q4DomainError("Initial moisture cannot be represented as log(C).") from exc
    return state


def _source_values(source: Callable[..., tuple[np.ndarray, np.ndarray]] | None, t: float, radius_m: float, nodes: np.ndarray) -> tuple[np.ndarray, np.ndarray] | None:
    if source is None:
        return None
    values = source(float(t), float(radius_m), nodes)
    if not isinstance(values, (tuple, list)) or len(values) != 2:
        raise ValueError("Source must return (temperature_source, moisture_source).")
    source_T = np.asarray(values[0], dtype=float)
    source_C = np.asarray(values[1], dtype=float)
    if source_T.shape != nodes.shape or source_C.shape != nodes.shape:
        raise ValueError("Manufactured source arrays must match the normalized nodes.")
    if not np.isfinite(source_T).all() or not np.isfinite(source_C).all():
        raise Q4DomainError("Manufactured source is nonfinite.")
    return source_T, source_C


def make_rhs(
    environment: Callable[[float], tuple[float, float]],
    radius: Callable[[float], float],
    n: int,
    *,
    model: str = "appendix4",
    h: float = H_BASE,
    hm: float = HM_BASE,
    source: Callable[..., tuple[np.ndarray, np.ndarray]] | None = None,
    omit_r2: bool = False,
    omit_r1: bool = False,
) -> tuple[Callable[[float, np.ndarray], np.ndarray], tuple[float, np.ndarray, np.ndarray, np.ndarray]]:
    if h < 0.0 or hm < 0.0:
        raise ValueError("h and hm must be nonnegative.")
    dxi, nodes, faces, weights = radial_geometry(n)

    def rhs(t: float, state: np.ndarray) -> np.ndarray:
        state = np.asarray(state, dtype=float)
        if state.size != 2 * (n + 1):
            raise ValueError("Physical state has the wrong size.")
        T = state[0::2]
        C = state[1::2]
        rho, cp, k, D = material_properties(C, T, model=model)
        T_env, C_env = environment(float(t))
        R = float(radius(float(t)))
        if not np.isfinite(R) or R <= 0.0:
            raise Q4DomainError("Radius is nonpositive or nonfinite.")
        k_face = 0.5 * (k[:-1] + k[1:])
        D_face = 0.5 * (D[:-1] + D[1:])
        heat_flux = faces * k_face * np.diff(T) / dxi
        moisture_flux = faces * D_face * np.diff(C) / dxi
        heat_surface = -(1.0 if omit_r1 else R) * h * (T[-1] - T_env)
        moisture_surface = -(1.0 if omit_r1 else R) * hm * (C[-1] - C_env)
        heat_div = np.empty(n + 1, dtype=float)
        moisture_div = np.empty(n + 1, dtype=float)
        heat_div[0] = heat_flux[0] / weights[0]
        moisture_div[0] = moisture_flux[0] / weights[0]
        if n > 1:
            heat_div[1:-1] = (heat_flux[1:] - heat_flux[:-1]) / weights[1:-1]
            moisture_div[1:-1] = (moisture_flux[1:] - moisture_flux[:-1]) / weights[1:-1]
        heat_div[-1] = (heat_surface - heat_flux[-1]) / weights[-1]
        moisture_div[-1] = (moisture_surface - moisture_flux[-1]) / weights[-1]
        scale = 1.0 if omit_r2 else 1.0 / (R * R)
        result = np.empty_like(state)
        result[0::2] = scale * heat_div / (rho * cp)
        result[1::2] = scale * moisture_div
        source_values = _source_values(source, float(t), R, nodes)
        if source_values is not None:
            source_T, source_C = source_values
            result[0::2] += source_T / (rho * cp)
            result[1::2] += source_C
        if not np.isfinite(result).all():
            raise Q4DomainError("Nonfinite physical RHS.")
        return result

    return rhs, (dxi, nodes, faces, weights)


def make_log_rhs(
    environment: Callable[[float], tuple[float, float]],
    radius: Callable[[float], float],
    n: int,
    *,
    model: str = "appendix4",
    h: float = H_BASE,
    hm: float = HM_BASE,
    source: Callable[..., tuple[np.ndarray, np.ndarray]] | None = None,
    omit_r2: bool = False,
    omit_r1: bool = False,
) -> tuple[Callable[[float, np.ndarray], np.ndarray], Callable[[float, np.ndarray], Any], tuple[float, np.ndarray, np.ndarray, np.ndarray]]:
    """Build the physical and analytic-log-coordinate BDF right hand sides."""
    physical_rhs, geometry = make_rhs(
        environment,
        radius,
        n,
        model=model,
        h=h,
        hm=hm,
        source=source,
        omit_r2=omit_r2,
        omit_r1=omit_r1,
    )
    dxi, nodes, faces, weights = geometry

    def rhs_log(t: float, transformed: np.ndarray) -> np.ndarray:
        physical, moisture = _physical_from_log(transformed, n)
        physical_derivative = physical_rhs(float(t), physical)
        result = np.empty_like(transformed)
        result[0::2] = physical_derivative[0::2]
        result[1::2] = physical_derivative[1::2] / moisture
        if not np.isfinite(result).all():
            raise Q4DomainError("Nonfinite RHS in log(C) coordinates.")
        return result

    def jac_log(t: float, transformed: np.ndarray):
        """Closed-form sparse Jacobian in interleaved (T,z) coordinates."""
        physical, moisture = _physical_from_log(transformed, n)
        T = physical[0::2]
        rho, cp, k, D = material_properties(moisture, T, model=model)
        rho_C, cp_C, k_C, D_C, D_T = property_derivatives(moisture, T, model=model)
        T_env, C_env = environment(float(t))
        R = float(radius(float(t)))
        if not np.isfinite(R) or R <= 0.0:
            raise Q4DomainError("Radius is nonpositive or nonfinite.")
        k_face = 0.5 * (k[:-1] + k[1:])
        D_face = 0.5 * (D[:-1] + D[1:])
        delta_T = np.diff(T)
        delta_C = np.diff(moisture)
        heat_flux = faces * k_face * delta_T / dxi
        moisture_flux = faces * D_face * delta_C / dxi
        surface_radius = 1.0 if omit_r1 else R
        heat_surface = -surface_radius * h * (T[-1] - T_env)
        moisture_surface = -surface_radius * hm * (moisture[-1] - C_env)
        heat_div = np.empty(n + 1, dtype=float)
        moisture_div = np.empty(n + 1, dtype=float)
        heat_div[0] = heat_flux[0] / weights[0]
        moisture_div[0] = moisture_flux[0] / weights[0]
        if n > 1:
            heat_div[1:-1] = (heat_flux[1:] - heat_flux[:-1]) / weights[1:-1]
            moisture_div[1:-1] = (moisture_flux[1:] - moisture_flux[:-1]) / weights[1:-1]
        heat_div[-1] = (heat_surface - heat_flux[-1]) / weights[-1]
        moisture_div[-1] = (moisture_surface - moisture_flux[-1]) / weights[-1]
        scale = 1.0 if omit_r2 else 1.0 / (R * R)

        # Each tuple is (node, dH/dT_node, dH/dC_node), where H is the
        # signed flux divergence before division by the control-volume weight.
        heat_T_left = -faces * k_face / dxi
        heat_T_right = -heat_T_left
        heat_C_left = faces * 0.5 * k_C[:-1] * delta_T / dxi
        heat_C_right = faces * 0.5 * k_C[1:] * delta_T / dxi
        moisture_T_left = faces * 0.5 * D_T[:-1] * delta_C / dxi
        moisture_T_right = faces * 0.5 * D_T[1:] * delta_C / dxi
        moisture_C_left = faces * (0.5 * D_C[:-1] * delta_C - D_face) / dxi
        moisture_C_right = faces * (0.5 * D_C[1:] * delta_C + D_face) / dxi

        rows: list[int] = []
        cols: list[int] = []
        data: list[float] = []

        def add(row: int, col: int, value: float) -> None:
            if value != 0.0:
                rows.append(int(row))
                cols.append(int(col))
                data.append(float(value))

        A = rho * cp
        A_C = rho_C * cp + rho * cp_C
        for i in range(n + 1):
            if i == 0:
                heat_terms = [(0, heat_T_left[0], heat_C_left[0]), (1, heat_T_right[0], heat_C_right[0])]
                moisture_terms = [(0, moisture_T_left[0], moisture_C_left[0]), (1, moisture_T_right[0], moisture_C_right[0])]
            elif i < n:
                heat_terms = [
                    (i, heat_T_left[i], heat_C_left[i]),
                    (i + 1, heat_T_right[i], heat_C_right[i]),
                    (i - 1, -heat_T_left[i - 1], -heat_C_left[i - 1]),
                    (i, -heat_T_right[i - 1], -heat_C_right[i - 1]),
                ]
                moisture_terms = [
                    (i, moisture_T_left[i], moisture_C_left[i]),
                    (i + 1, moisture_T_right[i], moisture_C_right[i]),
                    (i - 1, -moisture_T_left[i - 1], -moisture_C_left[i - 1]),
                    (i, -moisture_T_right[i - 1], -moisture_C_right[i - 1]),
                ]
            else:
                heat_terms = [
                    (n - 1, -heat_T_left[-1], -heat_C_left[-1]),
                    (n, -surface_radius * h - heat_T_right[-1], -heat_C_right[-1]),
                ]
                moisture_terms = [
                    (n - 1, -moisture_T_left[-1], -moisture_C_left[-1]),
                    (n, -moisture_T_right[-1], -surface_radius * hm - moisture_C_right[-1]),
                ]

            volume = weights[i]
            heat_capacity = A[i] * volume
            for j, dH_dT, dH_dC in heat_terms:
                add(2 * i, 2 * j, scale * dH_dT / heat_capacity)
                add(2 * i, 2 * j + 1, scale * dH_dC * moisture[j] / heat_capacity)
            # f_T = scale * H/(w A); the following is the A(C) quotient rule
            # in z_i, with H/w represented by heat_div[i].
            add(2 * i, 2 * i + 1, -scale * heat_div[i] * A_C[i] * moisture[i] / (A[i] * A[i]))

            moisture_scale = volume * moisture[i]
            for j, dM_dT, dM_dC in moisture_terms:
                add(2 * i + 1, 2 * j, scale * dM_dT / moisture_scale)
                add(2 * i + 1, 2 * j + 1, scale * dM_dC * moisture[j] / moisture_scale)
            # f_z = f_C/C, hence the local -f_z term from differentiating C.
            add(2 * i + 1, 2 * i + 1, -scale * moisture_div[i] / moisture[i])

        jacobian = coo_matrix(
            (np.asarray(data), (np.asarray(rows), np.asarray(cols))),
            shape=(2 * (n + 1), 2 * (n + 1)),
        ).tocsc()
        if not np.isfinite(jacobian.data).all():
            raise Q4DomainError("Analytic Jacobian contains a nonfinite entry.")
        return jacobian

    return rhs_log, jac_log, geometry


def formal_time_set(end_time_s: float, step_s: int = OUTPUT_STEP_S) -> np.ndarray:
    """Return 60-second rows from 60 s plus one unique exact endpoint."""
    end_time_s = float(end_time_s)
    if not np.isfinite(end_time_s) or end_time_s < 0.0:
        raise ValueError("End time must be finite and nonnegative.")
    if step_s <= 0:
        raise ValueError("Time step must be positive.")
    last = int(math.floor(end_time_s / step_s + 1.0e-12))
    values = [float(step_s * i) for i in range(1, last + 1)]
    if not values or abs(values[-1] - end_time_s) > 1.0e-9:
        values.append(end_time_s)
    return np.asarray(values, dtype=float)


def paper_time_set(end_time_s: float, step_s: int = PAPER_STEP_S) -> np.ndarray:
    return formal_time_set(end_time_s, step_s=step_s)


def radius_mask_tolerance(radius_m: float) -> float:
    return max(RADIUS_MASK_ABS_M, RADIUS_MASK_REL * abs(float(radius_m)))


def map_fixed_positions(
    xi_nodes: np.ndarray,
    scalar_states: np.ndarray,
    radii_m: np.ndarray | float,
    positions_cm: np.ndarray | None = None,
) -> dict[str, np.ndarray]:
    """Map any scalar field to fixed physical positions.

    The temperature and moisture fields must go through this same routine so
    that interpolation, radius tolerance, and exterior masking cannot drift
    apart.  ``moisture`` and ``surface_moisture`` are retained as aliases for
    older development callers; new code should use ``values`` and
    ``surface_values``.
    """
    xi_nodes = np.asarray(xi_nodes, dtype=float).reshape(-1)
    if len(xi_nodes) < 2 or not np.isfinite(xi_nodes).all():
        raise ValueError("Normalized nodes must contain at least two finite values.")
    if np.any(np.diff(xi_nodes) <= 0.0) or not np.isclose(xi_nodes[0], 0.0, atol=1.0e-12) or not np.isclose(xi_nodes[-1], 1.0, atol=1.0e-12):
        raise ValueError("Normalized nodes must be strictly increasing from 0 to 1.")
    states = np.asarray(scalar_states, dtype=float)
    if states.ndim == 1:
        states = states[None, :]
    if states.shape[1] != len(xi_nodes):
        raise ValueError("Scalar state and normalized grid disagree.")
    if not np.isfinite(states).all():
        raise ValueError("Scalar states must be finite before physical mapping.")
    radii = np.asarray(radii_m, dtype=float).reshape(-1)
    if len(radii) == 1 and len(states) != 1:
        radii = np.full(len(states), float(radii[0]))
    if len(radii) != len(states):
        raise ValueError("Radius history and state history disagree.")
    if not np.isfinite(radii).all() or np.any(radii <= 0.0):
        raise ValueError("Output radii must be positive and finite.")
    positions = np.arange(21, dtype=float) * 0.1 if positions_cm is None else np.asarray(positions_cm, dtype=float).reshape(-1)
    if len(positions) == 0 or not np.isfinite(positions).all() or np.any(positions < 0.0):
        raise ValueError("Output positions must be finite and nonnegative.")
    output = np.full((len(states), len(positions)), np.nan, dtype=float)
    surface = states[:, -1].copy()
    domain_mask = np.zeros_like(output, dtype=bool)
    for row, R in enumerate(radii):
        eps = radius_mask_tolerance(R)
        for col, position_cm in enumerate(positions):
            position_m = float(position_cm) * 0.01
            if position_m > R + eps:
                continue
            xi = position_m / R
            if abs(position_m - R) <= eps:
                xi = 1.0
            output[row, col] = float(np.interp(min(max(xi, 0.0), 1.0), xi_nodes, states[row]))
            domain_mask[row, col] = True
    return {
        "positions_cm": positions,
        "values": output,
        "surface_values": surface,
        "domain_mask": domain_mask,
        "moisture": output,
        "surface_moisture": surface,
    }


def _audit_state(log_state: np.ndarray, n: int, weights: np.ndarray) -> dict[str, float]:
    physical, moisture = _physical_from_log(log_state, n)
    return {
        "min_temperature_C": float(np.min(physical[0::2])),
        "max_temperature_C": float(np.max(physical[0::2])),
        "min_moisture_kg_per_kg": float(np.min(moisture)),
        "max_moisture_kg_per_kg": float(np.max(moisture)),
        "max_minus_center_kg_per_kg": float(np.max(moisture) - moisture[0]),
        "max_positive_neighbor_jump_kg_per_kg": float(np.max(np.diff(moisture))),
        "mean_moisture_kg_per_kg": float(2.0 * np.dot(weights, moisture)),
    }


def _event_value(dense: Callable[[np.ndarray], np.ndarray], t: float, n: int, threshold: float, center: bool = False) -> float:
    physical, moisture = _physical_from_log(np.asarray(dense(np.array([float(t)]))).reshape(-1), n)
    return float((moisture[0] if center else np.max(moisture)) - threshold)


def _normalize_event_monitor_g(value: float) -> float:
    """Represent a certified near-threshold root as zero in sign audits.

    The physical event certificate retains the unmodified residual.  This
    normalization is only for the sampled sign sequence: a dense-output
    endpoint can be a tiny positive number after the event callback has
    already identified the first crossing, and treating that round-off as a
    positive sign would erase the required positive-to-nonpositive transition.
    """
    value = float(value)
    if np.isfinite(value) and abs(value) <= EVENT_CONCENTRATION_TOL:
        return 0.0
    return value


def locate_crossing(
    dense: Callable[[np.ndarray], np.ndarray],
    lo: float,
    hi: float,
    n: int,
    *,
    threshold: float = THRESHOLD_C,
    center: bool = False,
) -> tuple[float, np.ndarray, dict[str, Any]]:
    """Refine a downward crossing and retain a <=1 s sign-change bracket."""
    value = lambda t: _event_value(dense, t, n, threshold, center=center)
    vlo, vhi = value(lo), value(hi)
    near_zero_endpoint = bool(0.0 < vhi <= EVENT_CONCENTRATION_TOL)
    if vlo <= 0.0 or (vhi > 0.0 and not near_zero_endpoint):
        raise ValueError("Crossing bracket does not have a downward sign change.")
    bracket_lo, bracket_hi = float(lo), float(hi)
    while bracket_hi - bracket_lo > EVENT_BRACKET_TOL_S:
        mid = 0.5 * (bracket_lo + bracket_hi)
        if value(mid) > 0.0:
            bracket_lo = mid
        else:
            bracket_hi = mid
    endpoint_value = float(value(bracket_hi))
    if endpoint_value > 0.0:
        # A solve_ivp event callback may return its endpoint with a strictly
        # positive residual caused only by floating-point round-off.  First
        # refine the original bracket to <=1 s, then accept only the explicit
        # near-zero case.  Never pass that positive endpoint to brentq.
        if not (near_zero_endpoint and endpoint_value <= EVENT_CONCENTRATION_TOL):
            raise ValueError("Crossing bracket does not have a downward sign change.")
        root = float(bracket_hi)
    else:
        root = float(brentq(value, bracket_lo, bracket_hi, xtol=1.0e-9, rtol=1.0e-13))
    root_log = np.asarray(dense(np.array([root]))).reshape(-1)
    root_state, root_moisture = _physical_from_log(root_log, n)
    return root, root_state, {
        "bracket_s": [float(bracket_lo), float(bracket_hi)],
        "bracket_width_s": float(bracket_hi - bracket_lo),
        "root_residual_kg_per_kg": float((root_moisture[0] if center else np.max(root_moisture)) - threshold),
        "center_C": float(root_moisture[0]),
        "max_C": float(np.max(root_moisture)),
    }


def _physical_from_transformed_matrix(values: np.ndarray, n: int) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    if values.ndim != 2:
        raise ValueError("Expected a (state, time) matrix.")
    result = values.copy()
    with np.errstate(over="raise", invalid="raise", under="ignore"):
        try:
            result[1::2, :] = np.exp(values[1::2, :])
        except FloatingPointError as exc:
            raise Q4DomainError("Log-state output overflowed; no clipping was applied.") from exc
    if not np.isfinite(result).all() or np.any(result[1::2, :] <= 0.0):
        raise Q4DomainError("Log-state output did not map to positive moisture.")
    if np.any(result[0::2, :] + 273.15 <= 0.0):
        raise Q4DomainError("Log-state output has nonphysical temperature.")
    return result


def _sample_output(
    transformed_values: np.ndarray,
    n: int,
    nodes: np.ndarray,
    weights: np.ndarray,
    radii_m: np.ndarray,
    positions_cm: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    physical = _physical_from_transformed_matrix(transformed_values, n)
    temperatures = physical[0::2, :].T
    moistures = physical[1::2, :].T
    mapped = map_fixed_positions(nodes, moistures, radii_m, positions_cm=positions_cm)
    mapped_temperature = map_fixed_positions(nodes, temperatures, radii_m, positions_cm=positions_cm)
    if mapped["values"].shape != mapped_temperature["values"].shape or not np.array_equal(mapped["domain_mask"], mapped_temperature["domain_mask"]):
        raise ArithmeticError("Temperature and moisture physical-position mappings disagree.")
    mean = 2.0 * (moistures @ weights)
    return (
        mapped["values"],
        mapped_temperature["values"],
        mapped["surface_values"],
        mapped_temperature["surface_values"],
        mean,
        physical,
        mapped["domain_mask"],
    )


def _save_block_evidence(
    evidence_dir: Path,
    run_id: str,
    segment_name: str,
    block_index: int,
    accepted_times: np.ndarray,
    accepted_physical: np.ndarray,
    accepted_radius: np.ndarray,
    accepted_g: np.ndarray,
    midpoint_times: np.ndarray,
    midpoint_physical: np.ndarray,
    midpoint_radius: np.ndarray,
    midpoint_g: np.ndarray,
    output_times: np.ndarray,
    output_g: np.ndarray,
) -> Path:
    """Persist one chunk's complete states without retaining all chunks."""
    evidence_dir.mkdir(parents=True, exist_ok=True)
    safe_run_id = "".join(char if char.isalnum() or char in "-_" else "_" for char in str(run_id))
    safe_segment = "".join(char if char.isalnum() or char in "-_" else "_" for char in str(segment_name))
    path = evidence_dir / f"{safe_run_id}_{safe_segment}_block_{block_index:04d}.npz"
    np.savez_compressed(
        path,
        accepted_time_s=np.asarray(accepted_times, dtype=float),
        accepted_temperature_C=np.asarray(accepted_physical[0::2, :].T, dtype=float),
        accepted_moisture_kg_per_kg=np.asarray(accepted_physical[1::2, :].T, dtype=float),
        accepted_radius_m=np.asarray(accepted_radius, dtype=float),
        accepted_g_N=np.asarray(accepted_g, dtype=float),
        midpoint_time_s=np.asarray(midpoint_times, dtype=float),
        midpoint_temperature_C=np.asarray(midpoint_physical[0::2, :].T, dtype=float),
        midpoint_moisture_kg_per_kg=np.asarray(midpoint_physical[1::2, :].T, dtype=float),
        midpoint_radius_m=np.asarray(midpoint_radius, dtype=float),
        midpoint_g_N=np.asarray(midpoint_g, dtype=float),
        output_time_s=np.asarray(output_times, dtype=float),
        output_g_N=np.asarray(output_g, dtype=float),
    )
    return path


def _event_sequence_audit(times: np.ndarray, values: np.ndarray, *, event_time: float) -> dict[str, Any]:
    """Audit sampled accepted/midpoint g values for missed re-crossings."""
    times = np.asarray(times, dtype=float).reshape(-1)
    values = np.asarray(values, dtype=float).reshape(-1)
    if len(times) != len(values) or len(times) == 0:
        return {"status": "FAIL", "reason": "empty_or_mismatched_sequence"}
    order = np.argsort(times, kind="mergesort")
    times, values = times[order], values[order]
    prior = times < float(event_time) - 1.0e-10
    prior_values = values[prior]
    after_or_root = values[~prior]
    downward = int(np.sum((values[:-1] > 0.0) & (values[1:] <= 0.0)))
    upward = int(np.sum((values[:-1] <= 0.0) & (values[1:] > 0.0)))
    return {
        "status": "CHECKED",
        "sample_count": int(len(values)),
        "time_start_s": float(times[0]),
        "time_end_s": float(times[-1]),
        "min_prior_g_N": float(np.min(prior_values)) if len(prior_values) else None,
        "max_prior_g_N": float(np.max(prior_values)) if len(prior_values) else None,
        "min_root_or_later_g_N": float(np.min(after_or_root)) if len(after_or_root) else None,
        "downward_sign_changes": downward,
        "upward_sign_changes": upward,
        "prior_positive": bool(len(prior_values) and np.all(prior_values > 0.0)),
        "no_upward_recrossing": bool(upward == 0),
    }


def _merge_monitor_samples(*samples: tuple[np.ndarray, np.ndarray]) -> tuple[np.ndarray, np.ndarray]:
    """Merge accepted, midpoint, and requested-output monitor samples.

    Samples at a shared time are collapsed conservatively by retaining the
    smallest g value.  This makes a nonpositive output sample visible even
    when an accepted-step or midpoint sample at the same time is rounded by a
    different floating-point path.
    """
    pieces = [
        (np.asarray(times, dtype=float).reshape(-1), np.asarray(values, dtype=float).reshape(-1))
        for times, values in samples
        if len(np.asarray(times).reshape(-1))
    ]
    if not pieces:
        return np.empty(0, dtype=float), np.empty(0, dtype=float)
    times = np.concatenate([piece[0] for piece in pieces])
    values = np.concatenate([piece[1] for piece in pieces])
    if len(times) != len(values) or not np.isfinite(times).all() or not np.isfinite(values).all():
        raise ValueError("Event monitor samples must be finite and have matching lengths.")
    order = np.argsort(times, kind="mergesort")
    times, values = times[order], values[order]
    merged_times: list[float] = [float(times[0])]
    merged_values: list[float] = [float(values[0])]
    for time_value, g_value in zip(times[1:], values[1:]):
        if abs(float(time_value) - merged_times[-1]) <= 1.0e-9:
            merged_values[-1] = min(merged_values[-1], float(g_value))
        else:
            merged_times.append(float(time_value))
            merged_values.append(float(g_value))
    return np.asarray(merged_times, dtype=float), np.asarray(merged_values, dtype=float)


def _first_downward_bracket(times: np.ndarray, values: np.ndarray) -> tuple[float, float] | None:
    """Return the earliest sampled positive-to-nonpositive bracket."""
    merged_times, merged_values = _merge_monitor_samples((times, values))
    for index in range(1, len(merged_times)):
        if merged_values[index - 1] > 0.0 and merged_values[index] <= 0.0:
            return float(merged_times[index - 1]), float(merged_times[index])
    return None


def _monitor_recovery_audit(
    times: np.ndarray,
    values: np.ndarray,
    *,
    event_time: float | None,
) -> dict[str, Any]:
    """Decide whether a block must be reintegrated to protect event detection."""
    merged_times, merged_values = _merge_monitor_samples((times, values))
    if len(merged_times) == 0:
        return {
            "status": "FAIL",
            "reason": "empty_monitor_sequence",
            "needs_recovery": True,
            "nonpositive_before_event": True,
            "positive_negative_positive": False,
            "first_downward_bracket_s": None,
        }
    if event_time is None:
        before = np.ones(len(merged_times), dtype=bool)
    else:
        before = merged_times < float(event_time) - 1.0e-10
    prior_values = merged_values[before]
    nonpositive_before = bool(np.any(prior_values <= 0.0))
    positive_negative_positive = False
    for index in range(1, len(merged_values) - 1):
        if merged_values[index] <= 0.0 and merged_values[index - 1] > 0.0 and np.any(merged_values[index + 1:] > 0.0):
            positive_negative_positive = True
            break
    any_nonpositive = bool(np.any(merged_values <= 0.0))
    needs_recovery = bool(nonpositive_before or positive_negative_positive or (event_time is None and any_nonpositive))
    bracket = _first_downward_bracket(merged_times, merged_values)
    return {
        "status": "CHECKED",
        "sample_count": int(len(merged_times)),
        "time_start_s": float(merged_times[0]),
        "time_end_s": float(merged_times[-1]),
        "prior_positive": bool(len(prior_values) and np.all(prior_values > 0.0)),
        "nonpositive_before_event": nonpositive_before,
        "any_nonpositive": any_nonpositive,
        "positive_negative_positive": positive_negative_positive,
        "first_downward_bracket_s": None if bracket is None else [float(bracket[0]), float(bracket[1])],
        "needs_recovery": needs_recovery,
    }


def _failed_monitor_event(
    state: np.ndarray,
    time_s: float,
    *,
    monitor_audit: dict[str, Any],
    sequence_audit: dict[str, Any],
    recovery_attempts: int,
    reason: str,
) -> dict[str, Any]:
    """Create an explicit failed event record when monitor recovery is ambiguous."""
    state = np.asarray(state, dtype=float).reshape(-1).copy()
    bracket = monitor_audit.get("first_downward_bracket_s")
    if not isinstance(bracket, (list, tuple)) or len(bracket) != 2:
        bracket = [float(time_s), float(time_s)]
    bracket_width = float(bracket[1]) - float(bracket[0])
    nan_state = np.full_like(state, np.nan, dtype=float)
    nan_nodes = np.full(state.size // 2, np.nan, dtype=float)
    return {
        "time_s": float(time_s),
        "state": state,
        "center_C": float(state[1]) if state.size > 1 else math.nan,
        "max_C": float(np.max(state[1::2])) if state.size > 1 else math.nan,
        "root_residual_kg_per_kg": float(monitor_audit.get("min_root_or_later_g_N", math.nan)),
        "bracket_s": [float(bracket[0]), float(bracket[1])],
        "bracket_width_s": bracket_width,
        "strict_before_time_s": math.nan,
        "strict_after_time_s": math.nan,
        "strict_before_max_C": math.nan,
        "strict_after_max_C": math.nan,
        "strict_before_pass": False,
        "strict_after_pass": False,
        "strict_before_state": nan_state.copy(),
        "strict_after_state": nan_state.copy(),
        "strict_before_g_N": math.nan,
        "strict_after_g_N": math.nan,
        "strict_before_mapped_C": nan_nodes.copy(),
        "strict_after_mapped_C": nan_nodes.copy(),
        "strict_before_mapped_T_C": nan_nodes.copy(),
        "strict_after_mapped_T_C": nan_nodes.copy(),
        "strict_before_domain_mask": np.zeros(state.size // 2, dtype=bool),
        "strict_after_domain_mask": np.zeros(state.size // 2, dtype=bool),
        "refined_time_s": None,
        "refinement_difference_s": math.inf,
        "refinement_pass": False,
        "before_state_error": reason,
        "after_state_error": reason,
        "sequence_audit": sequence_audit,
        "monitor_recovery_audit": monitor_audit,
        "monitor_recovery_attempts": int(recovery_attempts),
        "certification_status": "FAIL",
    }


def integrate_segment(
    state: np.ndarray,
    n: int,
    start_s: float,
    end_s: float,
    environment: Callable[[float], tuple[float, float]],
    radius: Callable[[float], float],
    *,
    model: str = "appendix4",
    rtol: float = BASE_RTL,
    atol: float = BASE_ATOL,
    max_step: float = BASE_MAX_STEP_BEFORE_S,
    chunk_seconds: float = DEFAULT_CHUNK_S,
    output_times: np.ndarray | None = None,
    output_positions_cm: np.ndarray | None = None,
    detect_event: bool = False,
    h: float = H_BASE,
    hm: float = HM_BASE,
    source: Callable[..., tuple[np.ndarray, np.ndarray]] | None = None,
    omit_r2: bool = False,
    omit_r1: bool = False,
    evidence_dir: str | Path | None = None,
    run_id: str = "q4_run",
    segment_name: str = "segment",
) -> dict[str, Any]:
    """Integrate one explicit environment segment in chunks.

    Only requested physical-distance rows and compact audit statistics are
    retained in memory.  When ``evidence_dir`` is supplied, every chunk also
    writes its complete accepted and midpoint physical states, radius, and
    full-grid event function to a separate NPZ file.
    """
    start_s, end_s = float(start_s), float(end_s)
    if end_s <= start_s:
        raise ValueError("Segment end must be later than its start.")
    if rtol <= 0.0 or atol <= 0.0 or max_step <= 0.0 or chunk_seconds <= 0.0:
        raise ValueError("Tolerances, max_step, and chunk_seconds must be positive.")
    state = np.asarray(state, dtype=float).copy()
    if state.size != 2 * (n + 1):
        raise ValueError("Initial state has the wrong size.")
    transformed = _to_log_state(state, n, model=model)
    rhs_log, jac_log, geometry = make_log_rhs(
        environment,
        radius,
        n,
        model=model,
        h=h,
        hm=hm,
        source=source,
        omit_r2=omit_r2,
        omit_r1=omit_r1,
    )
    domain_errors: tuple[type[Exception], ...] = (Q4DomainError,)
    dxi, nodes, faces, weights = geometry
    positions_cm = np.arange(21, dtype=float) * 0.1 if output_positions_cm is None else np.asarray(output_positions_cm, dtype=float).reshape(-1)
    requested = formal_time_set(end_s) if output_times is None else np.asarray(output_times, dtype=float).reshape(-1)
    requested = requested[(requested >= start_s - 1.0e-8) & (requested <= end_s + 1.0e-8) & (requested > 0.0)]
    requested = np.unique(np.round(requested, 10))
    out_times: list[float] = []
    out_moisture: list[np.ndarray] = []
    out_temperature: list[np.ndarray] = []
    out_surface: list[float] = []
    out_surface_temperature: list[float] = []
    out_mean: list[float] = []
    out_radius: list[float] = []
    out_domain_mask: list[np.ndarray] = []
    flux_times: list[np.ndarray] = []
    flux_rates: list[np.ndarray] = []
    accepted_audit: list[dict[str, float]] = []
    midpoint_audit: list[dict[str, float]] = []
    accepted_monitor_times: list[np.ndarray] = []
    accepted_monitor_g: list[np.ndarray] = []
    midpoint_monitor_times: list[np.ndarray] = []
    midpoint_monitor_g: list[np.ndarray] = []
    output_monitor_times: list[np.ndarray] = []
    output_monitor_g: list[np.ndarray] = []
    checked_output_states = 0
    block_records: list[dict[str, Any]] = []
    evidence_path = None if evidence_dir is None else Path(evidence_dir).resolve()
    if evidence_path is not None:
        evidence_path.mkdir(parents=True, exist_ok=True)
    block_index = 0
    previous_dense: Callable[[np.ndarray], np.ndarray] | None = None
    previous_interval: tuple[float, float] | None = None
    retry_step_override: float | None = None
    retry_attempts_carry = 0
    retry_history_carry: list[dict[str, Any]] = []

    def store_row(
        value: float,
        mapped_C: np.ndarray,
        mapped_T: np.ndarray,
        surface_C: float,
        surface_T: float,
        mean_C: float,
        radius_value: float,
        domain_mask: np.ndarray,
    ) -> None:
        """Append or replace one unrounded output row by its physical time."""
        existing = [index for index, old in enumerate(out_times) if abs(float(old) - float(value)) <= 1.0e-8]
        if existing:
            index = existing[0]
            out_times[index] = float(value)
            out_moisture[index] = np.asarray(mapped_C, dtype=float).copy()
            out_temperature[index] = np.asarray(mapped_T, dtype=float).copy()
            out_surface[index] = float(surface_C)
            out_surface_temperature[index] = float(surface_T)
            out_mean[index] = float(mean_C)
            out_radius[index] = float(radius_value)
            out_domain_mask[index] = np.asarray(domain_mask, dtype=bool).copy()
            return
        out_times.append(float(value))
        out_moisture.append(np.asarray(mapped_C, dtype=float).copy())
        out_temperature.append(np.asarray(mapped_T, dtype=float).copy())
        out_surface.append(float(surface_C))
        out_surface_temperature.append(float(surface_T))
        out_mean.append(float(mean_C))
        out_radius.append(float(radius_value))
        out_domain_mask.append(np.asarray(domain_mask, dtype=bool).copy())
    nfev = njev = nlu = 0
    positivity_retries = 0
    cursor = start_s
    state_final = state.copy()
    event: dict[str, Any] | None = None
    started = time.perf_counter()

    def event_function(t: float, z_state: np.ndarray) -> float:
        _physical, moisture = _physical_from_log(z_state, n)
        return float(np.max(moisture) - THRESHOLD_C)

    event_function.terminal = True  # type: ignore[attr-defined]
    event_function.direction = -1.0  # type: ignore[attr-defined]

    while cursor < end_s - 1.0e-11:
        segment_end = min(cursor + chunk_seconds, end_s)
        step_used = float(max_step if retry_step_override is None else retry_step_override)
        recovery_attempts = int(retry_attempts_carry)
        recovery_history: list[dict[str, Any]] = list(retry_history_carry)
        recovery_failed = False
        while True:
            try:
                solution = solve_ivp(
                    rhs_log,
                    (cursor, segment_end),
                    transformed,
                    method="BDF",
                    jac=jac_log if source is None else None,
                    jac_sparsity=jacobian_sparsity(n),
                    rtol=rtol,
                    atol=atol,
                    max_step=step_used,
                    dense_output=True,
                    events=event_function if detect_event else None,
                )
                break
            except domain_errors:
                if step_used <= max_step / 64.0:
                    raise
                step_used /= 2.0
                positivity_retries += 1
        if not solution.success:
            raise RuntimeError(f"BDF failed on segment {cursor}--{segment_end}: {solution.message}")
        nfev += int(solution.nfev)
        njev += int(solution.njev or 0)
        nlu += int(solution.nlu or 0)
        actual_end = float(solution.t[-1])
        accepted_phys = _physical_from_transformed_matrix(solution.y, n)
        accepted_C = accepted_phys[1::2, :]
        accepted_radius = np.asarray([float(radius(float(t))) for t in solution.t], dtype=float)
        accepted_g = np.max(accepted_C, axis=0) - THRESHOLD_C
        midpoint_times = np.empty(0, dtype=float)
        midpoint_states = np.empty((2 * (n + 1), 0), dtype=float)
        midpoint_phys = np.empty((2 * (n + 1), 0), dtype=float)
        midpoint_radius = np.empty(0, dtype=float)
        midpoint_g = np.empty(0, dtype=float)
        if len(solution.t) > 1:
            midpoint_times = 0.5 * (solution.t[:-1] + solution.t[1:])
            midpoint_states = solution.sol(midpoint_times)
            midpoint_phys = _physical_from_transformed_matrix(midpoint_states, n)
            midpoint_radius = np.asarray([float(radius(float(t))) for t in midpoint_times], dtype=float)
            midpoint_g = np.max(midpoint_phys[1::2, :], axis=0) - THRESHOLD_C
        eligible = requested[(requested >= cursor - 1.0e-8) & (requested <= actual_end + 1.0e-8)]
        output_monitor_times_block = np.asarray(eligible, dtype=float)
        output_monitor_g_block = np.empty(0, dtype=float)
        if len(output_monitor_times_block):
            output_monitor_states = solution.sol(output_monitor_times_block)
            output_monitor_phys = _physical_from_transformed_matrix(output_monitor_states, n)
            output_monitor_g_block = np.max(output_monitor_phys[1::2, :], axis=0) - THRESHOLD_C
        raw_event = None
        if detect_event and solution.t_events and len(solution.t_events[0]):
            raw_event = float(solution.t_events[0][0])
        monitor_times_block, monitor_g_block = _merge_monitor_samples(
            (np.asarray(solution.t, dtype=float), accepted_g),
            (midpoint_times, midpoint_g),
            (output_monitor_times_block, output_monitor_g_block),
        )
        recovery_audit = _monitor_recovery_audit(
            monitor_times_block,
            monitor_g_block,
            event_time=raw_event,
        )
        recovery_needed = bool(detect_event and recovery_audit.get("needs_recovery", False))
        if recovery_needed and recovery_attempts < EVENT_RECOVERY_MAX_RETRIES:
            recovery_history.append({
                "attempt": int(recovery_attempts),
                "max_step_s": float(step_used),
                "event_callback_time_s": raw_event,
                "monitor_audit": recovery_audit,
            })
            recovery_attempts += 1
            step_used = max(float(step_used) / 2.0, 1.0e-6)
            retry_step_override = step_used
            retry_attempts_carry = recovery_attempts
            retry_history_carry = list(recovery_history)
            continue
        recovery_failed = recovery_needed
        retry_step_override = None
        retry_attempts_carry = 0
        retry_history_carry = []

        # The statements below are kept outside the recovery loop so that a
        # coarse solve which exposed a hidden crossing never contaminates the
        # accepted-step evidence of the retried solve.
        accepted_monitor_times.append(np.asarray(solution.t, dtype=float))
        accepted_monitor_g.append(accepted_g.copy())
        for column in range(accepted_C.shape[1]):
            accepted_audit.append(_audit_state(solution.y[:, column], n, weights))
        if len(midpoint_times):
            midpoint_monitor_times.append(midpoint_times.copy())
            midpoint_monitor_g.append(midpoint_g.copy())
            for column in range(midpoint_states.shape[1]):
                midpoint_audit.append(_audit_state(midpoint_states[:, column], n, weights))
        if len(output_monitor_times_block):
            output_monitor_times.append(output_monitor_times_block.copy())
            output_monitor_g.append(output_monitor_g_block.copy())
            checked_output_states += int(len(output_monitor_times_block))
        if len(solution.t):
            env_C = np.asarray([environment(float(t))[1] for t in solution.t], dtype=float)
            flux_times.append(np.asarray(solution.t, dtype=float))
            flux_rates.append(2.0 * hm / accepted_radius * (accepted_C[-1, :] - env_C))

        block_path = None
        if evidence_path is not None:
            block_path = _save_block_evidence(
                evidence_path,
                run_id,
                segment_name,
                block_index,
                solution.t,
                accepted_phys,
                accepted_radius,
                accepted_g,
                midpoint_times,
                midpoint_phys,
                midpoint_radius,
                midpoint_g,
                output_monitor_times_block,
                output_monitor_g_block,
            )
        block_records.append({
            "run_id": str(run_id),
            "segment": str(segment_name),
            "block_index": int(block_index),
            "start_s": float(cursor),
            "requested_end_s": float(segment_end),
            "actual_end_s": actual_end,
            "accepted_states": int(len(solution.t)),
            "midpoint_states": int(len(midpoint_times)),
            "max_step_used_s": float(step_used),
            "accepted_g_min": float(np.min(accepted_g)),
            "accepted_g_max": float(np.max(accepted_g)),
            "midpoint_g_min": float(np.min(midpoint_g)) if len(midpoint_g) else None,
            "midpoint_g_max": float(np.max(midpoint_g)) if len(midpoint_g) else None,
            "output_monitor_states": int(len(output_monitor_times_block)),
            "output_monitor_g_min": float(np.min(output_monitor_g_block)) if len(output_monitor_g_block) else None,
            "output_monitor_g_max": float(np.max(output_monitor_g_block)) if len(output_monitor_g_block) else None,
            "monitor_sequence_states": int(len(monitor_times_block)),
            "monitor_sequence_audit": _event_sequence_audit(
                monitor_times_block,
                monitor_g_block,
                event_time=raw_event if raw_event is not None else segment_end,
            ),
            "monitor_recovery_audit": recovery_audit,
            "monitor_recovery_attempts": int(recovery_attempts),
            "monitor_recovery_history": recovery_history,
            "event_detected": bool(raw_event is not None),
            "evidence_file": None if block_path is None else str(block_path),
        })
        block_index += 1

        if len(eligible):
            z_values = solution.sol(eligible.astype(float))
            r_values = np.asarray([float(radius(float(t))) for t in eligible], dtype=float)
            mapped, temperatures, surfaces, temperature_surfaces, means, _, masks = _sample_output(
                z_values, n, nodes, weights, r_values, positions_cm
            )
            for index, value in enumerate(eligible):
                store_row(
                    float(value),
                    mapped[index],
                    temperatures[index],
                    float(surfaces[index]),
                    float(temperature_surfaces[index]),
                    float(means[index]),
                    float(r_values[index]),
                    masks[index],
                )

        if recovery_failed:
            event = _failed_monitor_event(
                accepted_phys[:, -1],
                actual_end,
                monitor_audit=recovery_audit,
                sequence_audit=block_records[-1]["monitor_sequence_audit"],
                recovery_attempts=recovery_attempts,
                reason="nonpositive_or_positive_negative_positive_monitor_after_step_halving",
            )
            state_final = accepted_phys[:, -1].copy()
            transformed = solution.y[:, -1].copy()
            cursor = actual_end
            break

        if detect_event and raw_event is not None:
            before_indices = np.where(solution.t < raw_event - 1.0e-10)[0]
            bracket_lo = float(solution.t[before_indices[-1]]) if len(before_indices) else float(cursor)
            if _event_value(solution.sol, bracket_lo, n, THRESHOLD_C) <= 0.0:
                if previous_dense is not None and previous_interval is not None and previous_interval[1] >= cursor - 1.0e-8:
                    bracket_lo = float(previous_interval[0])
                    while bracket_lo < cursor - 1.0e-8:
                        previous_value = _event_value(previous_dense, bracket_lo, n, THRESHOLD_C)
                        if previous_value > 0.0:
                            break
                        bracket_lo = min(cursor, bracket_lo + 0.5 * (cursor - bracket_lo))
                else:
                    bracket_lo = float(cursor)
            event_time, event_state, event_meta = locate_crossing(
                solution.sol,
                bracket_lo,
                raw_event,
                n,
            )
            # A second solve with half of the accepted-step ceiling checks
            # that the root is not an artefact of one large BDF step.
            refined_time = math.nan
            refinement_error = None
            try:
                # The current dense solution is valid only on this chunk;
                # for a cross-block event restart the refinement from the
                # current chunk's accepted initial state.
                refine_start = max(float(cursor), float(bracket_lo))
                refine_state = np.asarray(solution.sol(np.array([refine_start]))).reshape(-1)
                refine_end = max(raw_event + 1.0, refine_start + 1.0)
                refined_solution = solve_ivp(
                    rhs_log,
                    (refine_start, refine_end),
                    refine_state,
                    method="BDF",
                    jac=jac_log if source is None else None,
                    jac_sparsity=jacobian_sparsity(n),
                    rtol=rtol,
                    atol=atol,
                    max_step=max(min(step_used / 2.0, max_step), 0.125),
                    dense_output=True,
                    events=event_function,
                )
                if refined_solution.success and refined_solution.t_events and len(refined_solution.t_events[0]):
                    raw_refined = float(refined_solution.t_events[0][0])
                    refined_time, _refined_state, _refined_meta = locate_crossing(
                        refined_solution.sol,
                        refine_start,
                        raw_refined,
                        n,
                    )
                    refinement_error = abs(float(refined_time) - float(event_time))
                else:
                    refinement_error = math.inf
            except Exception as exc:  # noqa: BLE001 - preserve audit cause
                refinement_error = math.inf

            before_time = float(event_time - 0.5)
            after_time = float(event_time + 0.5)
            before_state = None
            before_error = None
            if before_time >= float(cursor) - 1.0e-8:
                before_state = np.asarray(solution.sol(np.array([before_time]))).reshape(-1)
            elif previous_dense is not None and previous_interval is not None and before_time >= previous_interval[0] - 1.0e-8 and before_time <= previous_interval[1] + 1.0e-8:
                before_state = np.asarray(previous_dense(np.array([before_time]))).reshape(-1)
            else:
                before_error = "no retained dense solution covers t_d-0.5 s"

            after_state = event_state.copy()
            after_error = None
            try:
                after_solution = solve_ivp(
                    rhs_log,
                    (event_time, after_time),
                    _to_log_state(event_state, n, model=model),
                    method="BDF",
                    jac=jac_log if source is None else None,
                    jac_sparsity=jacobian_sparsity(n),
                    rtol=rtol,
                    atol=atol,
                    max_step=min(step_used, 0.25),
                    dense_output=True,
                )
                if after_solution.success:
                    after_state, _ = _physical_from_log(after_solution.y[:, -1], n)
                else:
                    after_error = str(after_solution.message)
            except Exception as exc:  # noqa: BLE001 - failure is evidence
                after_error = repr(exc)
            before_witness_available = before_state is not None and before_error is None
            if before_witness_available:
                _before_physical, before_C = _physical_from_log(before_state, n)
                before_state = _before_physical
            else:
                # Do not substitute the root state for a missing strict-before
                # witness: doing so would silently turn an uncertified event
                # into apparently valid evidence (and would also request a
                # negative radius when an event occurs in the first 0.5 s).
                before_state = np.full_like(event_state, np.nan, dtype=float)
                before_C = np.full(n + 1, np.nan, dtype=float)
            after_C = after_state[1::2]
            sequence_times, sequence_values = _merge_monitor_samples(
                (np.concatenate(accepted_monitor_times) if accepted_monitor_times else np.empty(0), np.concatenate(accepted_monitor_g) if accepted_monitor_g else np.empty(0)),
                (np.concatenate(midpoint_monitor_times) if midpoint_monitor_times else np.empty(0), np.concatenate(midpoint_monitor_g) if midpoint_monitor_g else np.empty(0)),
                (np.concatenate(output_monitor_times) if output_monitor_times else np.empty(0), np.concatenate(output_monitor_g) if output_monitor_g else np.empty(0)),
                (np.array([float(event_time)]), np.array([_normalize_event_monitor_g(event_meta["root_residual_kg_per_kg"])])),
            )
            sequence_meta = _event_sequence_audit(sequence_times, sequence_values, event_time=event_time)
            root_mapped_C, root_mapped_T, root_surface_C, root_surface_T, root_mean_C, _root_physical, root_mask = _sample_output(
                _to_log_state(event_state, n, model=model)[:, None],
                n,
                nodes,
                weights,
                np.array([float(radius(event_time))]),
                positions_cm,
            )
            store_row(
                float(event_time),
                root_mapped_C[0],
                root_mapped_T[0],
                float(root_surface_C[0]),
                float(root_surface_T[0]),
                float(root_mean_C[0]),
                float(radius(event_time)),
                root_mask[0],
            )
            if before_witness_available:
                before_mapped_C, before_mapped_T, before_surface_C, before_surface_T, before_mean_C, _before_physical_full, before_mask = _sample_output(
                    _to_log_state(before_state, n, model=model)[:, None],
                    n,
                    nodes,
                    weights,
                    np.array([float(radius(before_time))]),
                    positions_cm,
                )
                before_radius = float(radius(before_time))
            else:
                before_mapped_C = np.full((1, len(positions_cm)), np.nan, dtype=float)
                before_mapped_T = np.full((1, len(positions_cm)), np.nan, dtype=float)
                before_surface_C = np.array([np.nan], dtype=float)
                before_surface_T = np.array([np.nan], dtype=float)
                before_mean_C = np.array([np.nan], dtype=float)
                before_mask = np.zeros((1, len(positions_cm)), dtype=bool)
                before_radius = math.nan
            after_mapped_C, after_mapped_T, after_surface_C, after_surface_T, after_mean_C, _after_physical_full, after_mask = _sample_output(
                _to_log_state(after_state, n, model=model)[:, None],
                n,
                nodes,
                weights,
                np.array([float(radius(after_time))]),
                positions_cm,
            )
            certification = (
                event_meta["bracket_width_s"] <= EVENT_BRACKET_TOL_S
                and abs(event_meta["root_residual_kg_per_kg"]) <= EVENT_CONCENTRATION_TOL
                and bool(np.max(before_C) > THRESHOLD_C)
                and bool(np.max(after_C) < THRESHOLD_C)
                and before_error is None
                and after_error is None
                and refinement_error is not None
                and refinement_error <= EVENT_BRACKET_TOL_S
                and sequence_meta.get("downward_sign_changes", 0) == 1
                and sequence_meta.get("prior_positive", False)
                and sequence_meta.get("no_upward_recrossing", False)
            )
            event = {
                "time_s": float(event_time),
                "state": event_state.copy(),
                "center_C": float(event_state[1]),
                "max_C": float(np.max(event_state[1::2])),
                "root_residual_kg_per_kg": float(event_meta["root_residual_kg_per_kg"]),
                "bracket_s": event_meta["bracket_s"],
                "bracket_width_s": float(event_meta["bracket_width_s"]),
                "strict_before_time_s": float(before_time),
                "strict_after_time_s": float(after_time),
                "strict_before_max_C": float(np.max(before_C)),
                "strict_after_max_C": float(np.max(after_C)),
                "strict_before_pass": bool(np.max(before_C) > THRESHOLD_C),
                "strict_after_pass": bool(np.max(after_C) < THRESHOLD_C),
                "strict_before_state": before_state.copy(),
                "strict_after_state": after_state.copy(),
                "strict_before_g_N": float(np.max(before_C) - THRESHOLD_C),
                "strict_after_g_N": float(np.max(after_C) - THRESHOLD_C),
                "strict_before_mapped_C": before_mapped_C[0].copy(),
                "strict_after_mapped_C": after_mapped_C[0].copy(),
                "strict_before_mapped_T_C": before_mapped_T[0].copy(),
                "strict_after_mapped_T_C": after_mapped_T[0].copy(),
                "strict_before_domain_mask": before_mask[0].copy(),
                "strict_after_domain_mask": after_mask[0].copy(),
                "root_mapped_C": root_mapped_C[0].copy(),
                "root_mapped_T_C": root_mapped_T[0].copy(),
                "root_domain_mask": root_mask[0].copy(),
                "root_radius_m": float(radius(event_time)),
                "strict_before_radius_m": before_radius,
                "strict_after_radius_m": float(radius(after_time)),
                "refined_time_s": None if not np.isfinite(refined_time) else float(refined_time),
                "refinement_difference_s": refinement_error,
                "refinement_pass": bool(refinement_error is not None and refinement_error <= EVENT_BRACKET_TOL_S),
                "before_state_error": before_error,
                "after_state_error": after_error,
                "sequence_audit": sequence_meta,
                "monitor_recovery_audit": recovery_audit,
                "monitor_recovery_attempts": int(recovery_attempts),
                "monitor_recovery_history": recovery_history,
                "certification_status": "PASS" if certification else "FAIL",
            }
            state_final = event_state.copy()
            transformed = _to_log_state(event_state, n, model=model)
            cursor = event_time
            break

        transformed = solution.y[:, -1].copy()
        state_final, _ = _physical_from_log(transformed, n)
        previous_dense = solution.sol
        previous_interval = (float(cursor), actual_end)
        cursor = actual_end

    monitor: dict[str, Any] = {
        "checked_accepted_states": int(len(accepted_audit)),
        "checked_midpoint_states": int(len(midpoint_audit)),
        "checked_output_states": int(checked_output_states),
        "min_temperature_C": float(min(item["min_temperature_C"] for item in accepted_audit + midpoint_audit)),
        "max_temperature_C": float(max(item["max_temperature_C"] for item in accepted_audit + midpoint_audit)),
        "min_moisture_kg_per_kg": float(min(item["min_moisture_kg_per_kg"] for item in accepted_audit + midpoint_audit)),
        "max_moisture_kg_per_kg": float(max(item["max_moisture_kg_per_kg"] for item in accepted_audit + midpoint_audit)),
        "max_max_minus_center_kg_per_kg": float(max(item["max_minus_center_kg_per_kg"] for item in accepted_audit + midpoint_audit)),
        "max_positive_neighbor_jump_kg_per_kg": float(max(item["max_positive_neighbor_jump_kg_per_kg"] for item in accepted_audit + midpoint_audit)),
    }
    return {
        "start_s": float(start_s),
        "end_s_requested": float(end_s),
        "final_time_s": float(event["time_s"] if event else end_s),
        "state_final": state_final,
        "times_s": np.asarray(out_times, dtype=float),
        "temperature_C": np.asarray(out_temperature, dtype=float).reshape((-1, len(positions_cm))) if out_temperature else np.empty((0, len(positions_cm))),
        "moisture_kg_per_kg": np.asarray(out_moisture, dtype=float).reshape((-1, len(positions_cm))) if out_moisture else np.empty((0, len(positions_cm))),
        "surface_moisture_kg_per_kg": np.asarray(out_surface, dtype=float),
        "surface_temperature_C": np.asarray(out_surface_temperature, dtype=float),
        "mean_moisture": np.asarray(out_mean, dtype=float),
        "radius_output_m": np.asarray(out_radius, dtype=float),
        "domain_mask": np.asarray(out_domain_mask, dtype=bool).reshape((-1, len(positions_cm))) if out_domain_mask else np.empty((0, len(positions_cm)), dtype=bool),
        "positions_cm": positions_cm,
        "flux_times_s": np.concatenate(flux_times) if flux_times else np.empty(0, dtype=float),
        "flux_rate_mean_per_s": np.concatenate(flux_rates) if flux_rates else np.empty(0, dtype=float),
        "accepted_monitor_times_s": np.concatenate(accepted_monitor_times) if accepted_monitor_times else np.empty(0, dtype=float),
        "accepted_monitor_g_N": np.concatenate(accepted_monitor_g) if accepted_monitor_g else np.empty(0, dtype=float),
        "midpoint_monitor_times_s": np.concatenate(midpoint_monitor_times) if midpoint_monitor_times else np.empty(0, dtype=float),
        "midpoint_monitor_g_N": np.concatenate(midpoint_monitor_g) if midpoint_monitor_g else np.empty(0, dtype=float),
        "output_monitor_times_s": np.concatenate(output_monitor_times) if output_monitor_times else np.empty(0, dtype=float),
        "output_monitor_g_N": np.concatenate(output_monitor_g) if output_monitor_g else np.empty(0, dtype=float),
        "segment_records": block_records,
        "monitor": monitor,
        "event": event,
        "radius_nodes_xi": nodes,
        "control_volume_weights": weights,
        "n_intervals": int(n),
        "dxi": float(dxi),
        "runtime_s": float(time.perf_counter() - started),
        "nfev": int(nfev),
        "njev": int(njev),
        "nlu": int(nlu),
        "positivity_retries": int(positivity_retries),
        "rtol": float(rtol),
        "atol": float(atol),
        "max_step_s": float(max_step),
        "model": model,
        "h_W_m2K": float(h),
        "hm_m_s": float(hm),
    }


def _merge_segment_outputs(first: dict[str, Any], second: dict[str, Any]) -> dict[str, Any]:
    """Merge two continuous pieces and remove their shared split row."""
    first_times = np.asarray(first["times_s"], dtype=float)
    second_times = np.asarray(second["times_s"], dtype=float)
    keep_second = np.ones(len(second_times), dtype=bool)
    if len(first_times):
        keep_second &= second_times > first_times[-1] + 1.0e-8
    merged = dict(second)
    for key in ("times_s", "temperature_C", "moisture_kg_per_kg", "surface_moisture_kg_per_kg", "surface_temperature_C", "mean_moisture", "radius_output_m", "domain_mask"):
        merged[key] = np.concatenate([np.asarray(first[key]), np.asarray(second[key])[keep_second]], axis=0)
    merged["flux_times_s"] = np.concatenate([first["flux_times_s"], second["flux_times_s"]])
    merged["flux_rate_mean_per_s"] = np.concatenate([first["flux_rate_mean_per_s"], second["flux_rate_mean_per_s"]])
    merged["accepted_monitor_times_s"] = np.concatenate([first["accepted_monitor_times_s"], second["accepted_monitor_times_s"]])
    merged["accepted_monitor_g_N"] = np.concatenate([first["accepted_monitor_g_N"], second["accepted_monitor_g_N"]])
    merged["midpoint_monitor_times_s"] = np.concatenate([first["midpoint_monitor_times_s"], second["midpoint_monitor_times_s"]])
    merged["midpoint_monitor_g_N"] = np.concatenate([first["midpoint_monitor_g_N"], second["midpoint_monitor_g_N"]])
    merged["output_monitor_times_s"] = np.concatenate([first.get("output_monitor_times_s", np.empty(0)), second.get("output_monitor_times_s", np.empty(0))])
    merged["output_monitor_g_N"] = np.concatenate([first.get("output_monitor_g_N", np.empty(0)), second.get("output_monitor_g_N", np.empty(0))])
    merged["segment_records"] = list(first.get("segment_records", [])) + list(second.get("segment_records", []))
    merged["monitor"] = {
        "checked_accepted_states": first["monitor"]["checked_accepted_states"] + second["monitor"]["checked_accepted_states"],
        "checked_midpoint_states": first["monitor"]["checked_midpoint_states"] + second["monitor"]["checked_midpoint_states"],
        "checked_output_states": first["monitor"].get("checked_output_states", 0) + second["monitor"].get("checked_output_states", 0),
        "min_temperature_C": min(first["monitor"]["min_temperature_C"], second["monitor"]["min_temperature_C"]),
        "max_temperature_C": max(first["monitor"]["max_temperature_C"], second["monitor"]["max_temperature_C"]),
        "min_moisture_kg_per_kg": min(first["monitor"]["min_moisture_kg_per_kg"], second["monitor"]["min_moisture_kg_per_kg"]),
        "max_moisture_kg_per_kg": max(first["monitor"]["max_moisture_kg_per_kg"], second["monitor"]["max_moisture_kg_per_kg"]),
        "max_max_minus_center_kg_per_kg": max(first["monitor"]["max_max_minus_center_kg_per_kg"], second["monitor"]["max_max_minus_center_kg_per_kg"]),
        "max_positive_neighbor_jump_kg_per_kg": max(first["monitor"]["max_positive_neighbor_jump_kg_per_kg"], second["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
    }
    merged["runtime_s"] = float(first["runtime_s"] + second["runtime_s"])
    merged["nfev"] = int(first["nfev"] + second["nfev"])
    merged["njev"] = int(first["njev"] + second["njev"])
    merged["nlu"] = int(first["nlu"] + second["nlu"])
    merged["positivity_retries"] = int(first["positivity_retries"] + second["positivity_retries"])
    merged["first_segment"] = first
    merged["second_segment"] = second
    return merged


def _g_evidence_summary(result: dict[str, Any]) -> dict[str, Any]:
    """Summarize the accepted-step and midpoint threshold evidence by segment."""
    def segment_summary(segment: dict[str, Any]) -> dict[str, Any]:
        accepted_t = np.asarray(segment.get("accepted_monitor_times_s", []), dtype=float)
        accepted_g = np.asarray(segment.get("accepted_monitor_g_N", []), dtype=float)
        midpoint_t = np.asarray(segment.get("midpoint_monitor_times_s", []), dtype=float)
        midpoint_g = np.asarray(segment.get("midpoint_monitor_g_N", []), dtype=float)
        output_t = np.asarray(segment.get("output_monitor_times_s", []), dtype=float)
        output_g = np.asarray(segment.get("output_monitor_g_N", []), dtype=float)
        all_t, all_g = _merge_monitor_samples(
            (accepted_t, accepted_g),
            (midpoint_t, midpoint_g),
            (output_t, output_g),
        )
        return {
            "accepted_count": int(len(accepted_t)),
            "midpoint_count": int(len(midpoint_t)),
            "output_count": int(len(output_t)),
            "time_start_s": float(np.min(all_t)) if len(all_t) else None,
            "time_end_s": float(np.max(all_t)) if len(all_t) else None,
            "g_min_N": float(np.min(all_g)) if len(all_g) else None,
            "g_max_N": float(np.max(all_g)) if len(all_g) else None,
            "block_count": int(len(segment.get("segment_records", []))),
        }

    accepted_t = np.asarray(result.get("accepted_monitor_times_s", []), dtype=float)
    accepted_g = np.asarray(result.get("accepted_monitor_g_N", []), dtype=float)
    midpoint_t = np.asarray(result.get("midpoint_monitor_times_s", []), dtype=float)
    midpoint_g = np.asarray(result.get("midpoint_monitor_g_N", []), dtype=float)
    output_t = np.asarray(result.get("output_monitor_times_s", []), dtype=float)
    output_g = np.asarray(result.get("output_monitor_g_N", []), dtype=float)
    all_t, all_g = _merge_monitor_samples(
        (accepted_t, accepted_g),
        (midpoint_t, midpoint_g),
        (output_t, output_g),
    )
    event = result.get("event")
    if (
        isinstance(event, dict)
        and np.isfinite(float(event.get("time_s", math.nan)))
        and np.isfinite(float(event.get("root_residual_kg_per_kg", math.nan)))
    ):
        all_t, all_g = _merge_monitor_samples(
            (all_t, all_g),
            (np.array([float(event["time_s"])]), np.array([_normalize_event_monitor_g(event.get("root_residual_kg_per_kg", math.nan))])),
        )
    reference_time = float(event["time_s"]) if event is not None else float(result.get("final_time_s", result.get("end_s_requested", 0.0)))
    sequence = _event_sequence_audit(all_t, all_g, event_time=reference_time) if len(all_t) else {"status": "NOT_AVAILABLE"}
    segments: dict[str, Any] = {}
    first = result.get("first_segment")
    second = result.get("second_segment")
    if isinstance(first, dict):
        segments["measured"] = segment_summary(first)
    if isinstance(second, dict):
        segments["plateau"] = segment_summary(second)
    return {
        "threshold_kg_per_kg": float(THRESHOLD_C),
        "accepted_count": int(len(accepted_t)),
        "midpoint_count": int(len(midpoint_t)),
        "output_count": int(len(output_t)),
        "time_start_s": float(np.min(all_t)) if len(all_t) else None,
        "time_end_s": float(np.max(all_t)) if len(all_t) else None,
        "g_min_N": float(np.min(all_g)) if len(all_g) else None,
        "g_max_N": float(np.max(all_g)) if len(all_g) else None,
        "full_run_sequence_audit": sequence,
        "segments": segments,
    }


def _apply_full_event_audit(result: dict[str, Any]) -> None:
    """Attach the merged two-segment monitor audit to the event certificate."""
    event = result.get("event")
    if not isinstance(event, dict):
        return
    sequence = result.get("g_evidence_summary", {}).get("full_run_sequence_audit", {})
    event["full_run_sequence_audit"] = sequence
    event["full_prior_positive"] = bool(sequence.get("prior_positive", False))
    event["full_downward_sign_changes"] = int(sequence.get("downward_sign_changes", 0))
    event["full_upward_sign_changes"] = int(sequence.get("upward_sign_changes", 0))
    event["full_no_upward_recrossing"] = bool(sequence.get("no_upward_recrossing", False))
    event_audit_pass = (
        sequence.get("status") == "CHECKED"
        and event["full_prior_positive"]
        and event["full_downward_sign_changes"] == 1
        and event["full_no_upward_recrossing"]
    )
    if event.get("certification_status") == "PASS" and not event_audit_pass:
        event["certification_status"] = "FAIL"
        event["full_audit_failure_reason"] = "merged_monitor_sequence_failed_event_uniqueness"


def _effective_balance(result: dict[str, Any], initial_mean: float) -> dict[str, Any]:
    """Check the dry-basis closure, integrating each environment side alone.

    The measured boundary and the nominal post-4-hour platform have distinct
    one-sided fluxes at 14400 s.  Combining them by timestamp would silently
    select the measured-side value for the first trapezoid of the second
    segment, so the two cumulative integrals are deliberately kept separate.
    """
    def segment_cumulative(segment: dict[str, Any]) -> tuple[np.ndarray, np.ndarray]:
        times = np.asarray(segment.get("flux_times_s", []), dtype=float)
        rates = np.asarray(segment.get("flux_rate_mean_per_s", []), dtype=float)
        if len(times) == 0:
            return np.empty(0, dtype=float), np.empty(0, dtype=float)
        order = np.argsort(times, kind="mergesort")
        times, rates = times[order], rates[order]
        unique = np.concatenate(([True], np.diff(times) > 1.0e-9))
        times, rates = times[unique], rates[unique]
        cumulative = np.zeros(len(times), dtype=float)
        if len(times) > 1:
            cumulative[1:] = np.cumsum(0.5 * (rates[1:] + rates[:-1]) * np.diff(times))
        return times, cumulative

    first = result.get("first_segment")
    second = result.get("second_segment")
    if isinstance(first, dict) and isinstance(second, dict):
        first_times, first_cumulative = segment_cumulative(first)
        second_times, second_cumulative = segment_cumulative(second)
        if len(first_times) == 0 or len(second_times) == 0:
            return {"status": "NOT_AVAILABLE"}
        split_mean = initial_mean - float(np.interp(float(first["end_s_requested"]), first_times, first_cumulative))
        output_times = np.asarray(result["times_s"], dtype=float)
        output_mean = np.asarray(result["mean_moisture"], dtype=float)
        predicted = np.empty_like(output_mean)
        first_mask = output_times <= float(first["end_s_requested"]) + 1.0e-8
        predicted[first_mask] = initial_mean - np.interp(output_times[first_mask], first_times, first_cumulative)
        predicted[~first_mask] = split_mean - np.interp(output_times[~first_mask], second_times, second_cumulative)
        residual = output_mean - predicted
        combined_cumulative = np.concatenate([first_cumulative, first_cumulative[-1] + second_cumulative])
        left_rate = float(first["flux_rate_mean_per_s"][-1])
        right_rate = float(second["flux_rate_mean_per_s"][0])
        return {
            "status": "CHECKED",
            "max_abs_residual": float(np.max(np.abs(residual))) if len(residual) else None,
            "max_relative_residual": float(np.max(np.abs(residual)) / max(abs(initial_mean), 1.0e-12)) if len(residual) else None,
            "times_s": np.concatenate([first_times, second_times]),
            "cumulative_loss": combined_cumulative,
            "output_residual": residual,
            "segment_checks": [
                {"segment": "measured", "start_s": float(first_times[0]), "end_s": float(first_times[-1]), "max_abs_residual": float(np.max(np.abs(residual[first_mask]))) if np.any(first_mask) else None},
                {"segment": "plateau", "start_s": float(second_times[0]), "end_s": float(second_times[-1]), "max_abs_residual": float(np.max(np.abs(residual[~first_mask]))) if np.any(~first_mask) else None},
            ],
            "split_flux_one_sided": {"left_rate_mean_per_s": left_rate, "right_rate_mean_per_s": right_rate, "jump_mean_per_s": right_rate - left_rate},
        }
    times = np.asarray(result.get("flux_times_s", []), dtype=float)
    rates = np.asarray(result.get("flux_rate_mean_per_s", []), dtype=float)
    if len(times) == 0:
        return {"status": "NOT_AVAILABLE"}
    order = np.argsort(times, kind="mergesort")
    times, rates = times[order], rates[order]
    unique = np.concatenate(([True], np.diff(times) > 1.0e-9))
    times, rates = times[unique], rates[unique]
    cumulative = np.zeros(len(times), dtype=float)
    if len(times) > 1:
        cumulative[1:] = np.cumsum(0.5 * (rates[1:] + rates[:-1]) * np.diff(times))
    output_times = np.asarray(result["times_s"], dtype=float)
    output_mean = np.asarray(result["mean_moisture"], dtype=float)
    predicted = initial_mean - np.interp(output_times, times, cumulative) if len(output_times) else np.empty(0, dtype=float)
    residual = output_mean - predicted
    return {
        "status": "CHECKED",
        "max_abs_residual": float(np.max(np.abs(residual))) if len(residual) else None,
        "max_relative_residual": float(np.max(np.abs(residual)) / max(abs(initial_mean), 1.0e-12)) if len(residual) else None,
        "times_s": times,
        "cumulative_loss": cumulative,
        "output_residual": residual,
    }


def _output_schema(result: dict[str, Any]) -> dict[str, Any]:
    """Return the single time/field contract consumed by later stages."""
    event = result.get("event")
    endpoint = float(event["time_s"]) if event is not None else float(result["end_s_requested"])
    expected = formal_time_set(endpoint, step_s=OUTPUT_STEP_S)
    actual = np.asarray(result.get("times_s", []), dtype=float)
    sorted_unique = np.unique(np.round(actual, 10))
    same = len(actual) == len(sorted_unique) and len(actual) == len(expected) and np.allclose(sorted_unique, expected, rtol=0.0, atol=1.0e-8)
    fields_ok = all(
        np.asarray(result.get(key, np.empty((0, 0)))).shape[:1] == (len(actual),)
        for key in ("temperature_C", "moisture_kg_per_kg", "surface_moisture_kg_per_kg", "surface_temperature_C", "mean_moisture", "radius_output_m", "domain_mask")
    )
    monotonic = bool(len(actual) <= 1 or np.all(np.diff(actual) > 0.0))
    return {
        "expected_time_count": int(len(expected)),
        "actual_time_count": int(len(actual)),
        "expected_last_time_s": float(expected[-1]) if len(expected) else None,
        "actual_last_time_s": float(actual[-1]) if len(actual) else None,
        "time_set_pass": bool(same),
        "field_first_dimension_pass": bool(fields_ok),
        "unique_time_pass": bool(len(actual) == len(np.unique(actual))),
        "monotonic_time_pass": monotonic,
        "pass": bool(same and fields_ok and monotonic and len(actual) == len(np.unique(actual))),
    }


def _claim_evidence_directory(base_directory: str | Path, run_id: str) -> Path:
    """Claim a never-before-used evidence directory for one numerical run."""
    base = Path(base_directory).expanduser().resolve()
    safe_run_id = "".join(char if char.isalnum() or char in "-_" else "_" for char in str(run_id)).strip("._")
    if not safe_run_id:
        raise ValueError("run_id must contain at least one usable character.")
    base.mkdir(parents=True, exist_ok=True)
    claimed = (base / safe_run_id).resolve()
    if claimed.parent != base:
        raise ValueError("run_id escaped the evidence directory.")
    if claimed.exists():
        raise FileExistsError(f"Evidence run directory already exists: {claimed}")
    claimed.mkdir(parents=False, exist_ok=False)
    return claimed


def run_case(
    boundary: np.ndarray,
    radius: RadiusModel | Callable[[float], float],
    *,
    n: int,
    model: str = "appendix4",
    max_end_s: float = MAX_END_S,
    rtol: float = BASE_RTL,
    atol: float = BASE_ATOL,
    max_step_before_s: float = BASE_MAX_STEP_BEFORE_S,
    max_step_after_s: float = BASE_MAX_STEP_AFTER_S,
    chunk_before_s: float = DEFAULT_CHUNK_S,
    chunk_after_s: float = DEFAULT_CHUNK_S,
    h: float = H_BASE,
    hm: float = HM_BASE,
    initial: np.ndarray | None = None,
    plateau_temperature_C: float = PLATEAU_T_C,
    plateau_moisture_kg_per_kg: float = PLATEAU_C,
    output_positions_cm: np.ndarray | None = None,
    source: Callable[..., tuple[np.ndarray, np.ndarray]] | None = None,
    omit_r2: bool = False,
    omit_r1: bool = False,
    case_name: str | None = None,
    geometry_label: str | None = None,
    radius_metadata: dict[str, Any] | None = None,
    evidence_dir: str | Path | None = None,
    run_id: str = "q4_run",
) -> dict[str, Any]:
    """Run from t=0 through the four-hour boundary and detect the event."""
    max_end_s = float(max_end_s)
    if max_end_s <= SWITCH_S:
        raise ValueError("max_end_s must exceed the four-hour split.")
    boundary = np.asarray(boundary, dtype=float)
    if boundary.ndim != 2 or boundary.shape[1] < 3 or len(boundary) < 2 or not np.isfinite(boundary[:, :3]).all():
        raise ValueError("Boundary must be a finite (time, temperature, moisture) table.")
    if np.any(np.diff(boundary[:, 0]) <= 0.0) or boundary[0, 0] > 0.0 or boundary[-1, 0] < SWITCH_S:
        raise ValueError("Measured boundary does not cover the four-hour split.")
    if initial is None:
        initial = initial_state(n)
    initial = np.asarray(initial, dtype=float)
    positions_cm = np.arange(21, dtype=float) * 0.1 if output_positions_cm is None else np.asarray(output_positions_cm, dtype=float).reshape(-1)
    minute_times = formal_time_set(max_end_s, step_s=OUTPUT_STEP_S)
    measured = measured_environment(boundary)
    plateau = plateau_environment(temperature_C=plateau_temperature_C, moisture_kg_per_kg=plateau_moisture_kg_per_kg)
    inferred_geometry = geometry_label
    if inferred_geometry is None:
        if isinstance(radius, RadiusModel) and np.allclose(radius.radii_m, radius.radii_m[0], rtol=0.0, atol=1.0e-14) and np.isclose(radius.radii_m[0], RADIUS_INITIAL_M, rtol=0.0, atol=1.0e-14):
            inferred_geometry = "fixed_R0"
        else:
            inferred_geometry = "external_radius"
    inferred_radius_metadata = radius_metadata or {
        "source": "external callable" if not isinstance(radius, RadiusModel) else "RadiusModel",
        "geometry": inferred_geometry,
        "interpolation": getattr(radius, "interpolation", "callable"),
        "post_mode": getattr(radius, "post_mode", "external"),
        "start_radius_m": float(radius(0.0)),
        "end_radius_m": float(radius(max_end_s)),
    }
    claimed_evidence_dir = None if evidence_dir is None else _claim_evidence_directory(evidence_dir, run_id)
    run_config = {
        "run_id": str(run_id),
        "case": case_name or "single_case",
        "property_model": model,
        "backend": "q4_native",
        "geometry": inferred_geometry,
        "radius": inferred_radius_metadata,
        "environment": {
            "measured_interval_s": [float(boundary[0, 0]), float(boundary[-1, 0])],
            "measured_until_s": float(SWITCH_S),
            "post_switch": {"temperature_C": float(plateau_temperature_C), "moisture_kg_per_kg": float(plateau_moisture_kg_per_kg)},
        },
        "exchange": {"h_W_m2K": float(h), "hm_m_s": float(hm)},
        "initial": {"temperature_C": float(initial[0]), "moisture_kg_per_kg": float(initial[1])},
        "n_intervals": int(n),
        "max_end_s": float(max_end_s),
        "segments": {
            "measured": {"start_s": 0.0, "end_s": float(SWITCH_S), "rtol": float(rtol), "atol": float(atol), "max_step_s": float(max_step_before_s), "chunk_s": float(chunk_before_s)},
            "plateau": {"start_s": float(SWITCH_S), "end_s": float(max_end_s), "rtol": float(rtol), "atol": float(atol), "max_step_s": float(max_step_after_s), "chunk_s": float(chunk_after_s)},
        },
        "output": {"time_step_s": int(OUTPUT_STEP_S), "positions_cm": positions_cm.tolist(), "exact_endpoint_unique": True},
        "evidence_directory": None if claimed_evidence_dir is None else str(claimed_evidence_dir),
        "termination": {"criterion": "first full-grid max moisture downward crossing", "threshold_kg_per_kg": float(THRESHOLD_C)},
    }
    first_evidence = None if claimed_evidence_dir is None else claimed_evidence_dir / "blocks"
    first = integrate_segment(
        initial,
        n,
        0.0,
        SWITCH_S,
        measured,
        radius,
        model=model,
        rtol=rtol,
        atol=atol,
        max_step=max_step_before_s,
        chunk_seconds=chunk_before_s,
        output_times=minute_times,
        output_positions_cm=positions_cm,
        detect_event=True,
        h=h,
        hm=hm,
        source=source,
        omit_r2=omit_r2,
        omit_r1=omit_r1,
        evidence_dir=first_evidence,
        run_id=run_id,
        segment_name="measured",
    )
    if first["event"] is not None:
        result = dict(first)
        result["first_segment"] = first
        result["second_segment"] = None
    else:
        second_evidence = None if claimed_evidence_dir is None else claimed_evidence_dir / "blocks"
        second = integrate_segment(
            first["state_final"],
            n,
            SWITCH_S,
            max_end_s,
            plateau,
            radius,
            model=model,
            rtol=rtol,
            atol=atol,
            max_step=max_step_after_s,
            chunk_seconds=chunk_after_s,
            output_times=minute_times,
            output_positions_cm=positions_cm,
            detect_event=True,
            h=h,
            hm=hm,
            source=source,
            omit_r2=omit_r2,
            omit_r1=omit_r1,
            evidence_dir=second_evidence,
            run_id=run_id,
            segment_name="plateau",
        )
        result = _merge_segment_outputs(first, second)
        result["event"] = second["event"]
        result["state_final"] = second["state_final"]
        result["final_time_s"] = second["final_time_s"]
    initial_mean = float(2.0 * np.dot(first["control_volume_weights"], initial[1::2]))
    result["initial_mean_moisture"] = initial_mean
    result["effective_balance"] = _effective_balance(result, initial_mean)
    result["radius_interpolation"] = getattr(radius, "interpolation", "callable")
    result["radius_post_mode"] = getattr(radius, "post_mode", "external")
    result["boundary_switch_s"] = SWITCH_S
    result["plateau_temperature_C"] = float(plateau_temperature_C)
    result["plateau_moisture_kg_per_kg"] = float(plateau_moisture_kg_per_kg)
    result["run_config"] = run_config
    result["radius_metadata"] = inferred_radius_metadata
    result["g_evidence_summary"] = _g_evidence_summary(result)
    _apply_full_event_audit(result)
    result["output_schema"] = _output_schema(result)
    return result


def _deduplicate_rows(result: dict[str, Any]) -> dict[str, Any]:
    times = np.asarray(result["times_s"], dtype=float)
    if len(times) == 0:
        return result
    order = np.argsort(times, kind="mergesort")
    keep = np.ones(len(times), dtype=bool)
    sorted_times = times[order]
    keep[1:] = np.diff(sorted_times) > 1.0e-8
    selected = order[keep]
    for key in ("times_s", "temperature_C", "moisture_kg_per_kg", "surface_moisture_kg_per_kg", "surface_temperature_C", "mean_moisture", "radius_output_m", "domain_mask"):
        result[key] = np.asarray(result[key])[selected]
    return result


def save_result_files(output: str | Path, result: dict[str, Any], *, boundary_metadata: dict[str, Any] | None = None, radius_metadata: dict[str, Any] | None = None, validation: dict[str, Any] | None = None) -> None:
    """Save unrounded fields, restart state, and auditable endpoint evidence.

    A failed event certificate is never written under the requested result
    directory.  It is redirected to a clearly named development subdirectory
    and raises after the failure evidence has been preserved.
    """
    output = Path(output)
    event = result.get("event")
    full_audit = result.get("g_evidence_summary", {}).get("full_run_sequence_audit", {})
    full_audit_failed = (
        not isinstance(full_audit, dict)
        or full_audit.get("status") != "CHECKED"
        or not bool(full_audit.get("prior_positive", False))
        or (event is not None and (
            int(full_audit.get("downward_sign_changes", 0)) != 1
            or int(full_audit.get("upward_sign_changes", 0)) != 0
            or not bool(full_audit.get("no_upward_recrossing", False))
        ))
    )
    event_failed = (event is not None and event.get("certification_status") != "PASS") or full_audit_failed
    requested_output = output
    if event_failed:
        output = output / "failed_event_evidence"
    output.mkdir(parents=True, exist_ok=True)
    result = _deduplicate_rows(dict(result))
    np.savez_compressed(
        output / "q4_fields.npz",
        time_s=np.asarray(result["times_s"], dtype=float),
        positions_cm=np.asarray(result["positions_cm"], dtype=float),
        temperature_C=np.asarray(result["temperature_C"], dtype=float),
        moisture_kg_per_kg=np.asarray(result["moisture_kg_per_kg"], dtype=float),
        surface_moisture_kg_per_kg=np.asarray(result["surface_moisture_kg_per_kg"], dtype=float),
        surface_temperature_C=np.asarray(result["surface_temperature_C"], dtype=float),
        radius_m=np.asarray(result["radius_output_m"], dtype=float),
        mean_moisture=np.asarray(result["mean_moisture"], dtype=float),
        domain_mask=np.asarray(result["domain_mask"], dtype=bool),
    )
    np.savez_compressed(
        output / "q4_restart.npz",
        final_time_s=np.array(result["final_time_s"], dtype=float),
        n_intervals=np.array(result["n_intervals"], dtype=int),
        radius_nodes_xi=np.asarray(result["radius_nodes_xi"], dtype=float),
        control_volume_weights=np.asarray(result["control_volume_weights"], dtype=float),
        temperature_final_C=np.asarray(result["state_final"][0::2], dtype=float),
        moisture_final_kg_per_kg=np.asarray(result["state_final"][1::2], dtype=float),
        h_W_m2K=np.array(result["h_W_m2K"], dtype=float),
        hm_m_s=np.array(result["hm_m_s"], dtype=float),
    )
    # Keep the discrete boundary-flux history separate from the workbook
    # fields.  It is the numerical evidence used by the dry-basis closure
    # check and avoids reconstructing a flux from rounded output rows.
    np.savez_compressed(
        output / "q4_flux.npz",
        time_s=np.asarray(result["flux_times_s"], dtype=float),
        rate_mean_per_s=np.asarray(result["flux_rate_mean_per_s"], dtype=float),
        accepted_monitor_time_s=np.asarray(result.get("accepted_monitor_times_s", []), dtype=float),
        accepted_monitor_g_N=np.asarray(result.get("accepted_monitor_g_N", []), dtype=float),
        midpoint_monitor_time_s=np.asarray(result.get("midpoint_monitor_times_s", []), dtype=float),
        midpoint_monitor_g_N=np.asarray(result.get("midpoint_monitor_g_N", []), dtype=float),
        output_monitor_time_s=np.asarray(result.get("output_monitor_times_s", []), dtype=float),
        output_monitor_g_N=np.asarray(result.get("output_monitor_g_N", []), dtype=float),
    )
    if event is not None:
        np.savez_compressed(
            output / "q4_event.npz",
            time_s=np.array(event["time_s"], dtype=float),
            temperature_C=np.asarray(event["state"][0::2], dtype=float),
            moisture_kg_per_kg=np.asarray(event["state"][1::2], dtype=float),
            root_radius_m=np.array(event.get("root_radius_m", np.nan), dtype=float),
            root_mapped_C=np.asarray(event.get("root_mapped_C", []), dtype=float),
            root_mapped_T_C=np.asarray(event.get("root_mapped_T_C", []), dtype=float),
            root_domain_mask=np.asarray(event.get("root_domain_mask", []), dtype=bool),
            strict_before_time_s=np.array(event.get("strict_before_time_s", np.nan), dtype=float),
            strict_after_time_s=np.array(event.get("strict_after_time_s", np.nan), dtype=float),
            strict_before_state=np.asarray(event.get("strict_before_state", []), dtype=float),
            strict_after_state=np.asarray(event.get("strict_after_state", []), dtype=float),
            strict_before_radius_m=np.array(event.get("strict_before_radius_m", np.nan), dtype=float),
            strict_after_radius_m=np.array(event.get("strict_after_radius_m", np.nan), dtype=float),
            strict_before_mapped_C=np.asarray(event.get("strict_before_mapped_C", []), dtype=float),
            strict_after_mapped_C=np.asarray(event.get("strict_after_mapped_C", []), dtype=float),
            strict_before_mapped_T_C=np.asarray(event.get("strict_before_mapped_T_C", []), dtype=float),
            strict_after_mapped_T_C=np.asarray(event.get("strict_after_mapped_T_C", []), dtype=float),
            strict_before_domain_mask=np.asarray(event.get("strict_before_domain_mask", []), dtype=bool),
            strict_after_domain_mask=np.asarray(event.get("strict_after_domain_mask", []), dtype=bool),
            bracket_s=np.asarray(event.get("bracket_s", [np.nan, np.nan]), dtype=float),
            root_residual_kg_per_kg=np.array(event.get("root_residual_kg_per_kg", np.nan), dtype=float),
            bracket_width_s=np.array(event.get("bracket_width_s", np.nan), dtype=float),
        )
    write_json(output / "q4_blocks_manifest.json", {
        "run_id": result.get("run_config", {}).get("run_id"),
        "run_config": result.get("run_config"),
        "segment_records": result.get("segment_records", []),
        "g_evidence_summary": result.get("g_evidence_summary"),
        "block_evidence_required": True,
    })
    summary = {
        "model": result.get("model"),
        "n_intervals": result.get("n_intervals"),
        "final_time_s": result.get("final_time_s"),
        "event": event,
        "event_certification_status": None if event is None else event.get("certification_status"),
        "radius_interpolation": result.get("radius_interpolation"),
        "radius_post_mode": result.get("radius_post_mode"),
        "boundary_switch_s": result.get("boundary_switch_s"),
        "rtol": result.get("rtol"),
        "atol": result.get("atol"),
        "max_step_s": result.get("max_step_s"),
        "runtime_s": result.get("runtime_s"),
        "nfev": result.get("nfev"),
        "njev": result.get("njev"),
        "nlu": result.get("nlu"),
        "positivity_retries": result.get("positivity_retries"),
        "monitor": result.get("monitor"),
        "g_evidence_summary": result.get("g_evidence_summary"),
        "effective_balance": result.get("effective_balance"),
        "run_config": result.get("run_config"),
        "output_directory_requested": str(requested_output),
        "output_directory_written": str(output),
        "write_mode": "DEVELOPMENT_FAILURE_EVIDENCE" if event_failed else "UNROUNDED_RESULT",
        "boundary_metadata": boundary_metadata,
        "radius_metadata": radius_metadata,
        "validation": validation,
    }
    # Numpy histories are deliberately kept out of JSON; they are in q4_fields.
    write_json(output / "q4_case_summary.json", summary)
    if event_failed:
        raise RuntimeError(f"Event certificate failed; evidence was redirected to {output}")


def case_summary(result: dict[str, Any]) -> dict[str, Any]:
    event = result.get("event")
    return {
        "case": result.get("run_config", {}).get("case"),
        "model": result.get("model"),
        "geometry": result.get("run_config", {}).get("geometry"),
        "n_intervals": int(result.get("n_intervals", 0)),
        "final_time_s": float(result.get("final_time_s", math.nan)),
        "event_time_s": None if event is None else float(event["time_s"]),
        "event_root_residual": None if event is None else float(event["root_residual_kg_per_kg"]),
        "event_bracket_width_s": None if event is None else float(event["bracket_width_s"]),
        "strict_before_pass": None if event is None else bool(event["strict_before_pass"]),
        "event_certification_status": None if event is None else event.get("certification_status"),
        "output_schema": result.get("output_schema"),
        "run_config": result.get("run_config"),
        "g_evidence_summary": result.get("g_evidence_summary"),
        "strict_after_pass": None if event is None else bool(event["strict_after_pass"]),
        "min_moisture_kg_per_kg": float(result["monitor"]["min_moisture_kg_per_kg"]),
        "max_positive_neighbor_jump_kg_per_kg": float(result["monitor"]["max_positive_neighbor_jump_kg_per_kg"]),
    }


def make_radius_function(
    path: str | Path = DEFAULT_RADIUS,
    *,
    interpolation: str = "pchip",
    post_mode: str = "plateau",
) -> tuple[RadiusModel, dict[str, Any]]:
    return read_radius(path, interpolation=interpolation, post_mode=post_mode)


def run_four_cases(
    boundary: np.ndarray,
    radius: RadiusModel,
    *,
    n: int,
    max_end_s: float = MAX_END_S,
    rtol: float = BASE_RTL,
    atol: float = BASE_ATOL,
    max_step_before_s: float = BASE_MAX_STEP_BEFORE_S,
    max_step_after_s: float = BASE_MAX_STEP_AFTER_S,
    chunk_before_s: float = DEFAULT_CHUNK_S,
    chunk_after_s: float = DEFAULT_CHUNK_S,
    h: float = H_BASE,
    hm: float = HM_BASE,
    radius_metadata: dict[str, Any] | None = None,
    evidence_dir: str | Path | None = None,
) -> dict[str, dict[str, Any]]:
    """Run the A--D factorial configurations with shared initial/boundary data."""
    fixed_radius = RadiusModel(
        np.array([0.0, max(float(max_end_s), SWITCH_S + 1.0)]),
        np.array([RADIUS_INITIAL_M, RADIUS_INITIAL_M]),
        interpolation="linear",
        post_mode="plateau",
    )
    results: dict[str, dict[str, Any]] = {}
    fixed_metadata = {
        "source": "constant initial radius R0",
        "geometry": "fixed_R0",
        "interpolation": "not_applicable",
        "post_mode": "not_applicable",
        "start_radius_m": float(RADIUS_INITIAL_M),
        "end_radius_m": float(RADIUS_INITIAL_M),
    }
    for name, model, geometry in (
        ("A", "appendix3", fixed_radius),
        ("B", "appendix3", radius),
        ("C", "appendix4", fixed_radius),
        ("D", "appendix4", radius),
    ):
        results[name] = run_case(
            boundary,
            geometry,
            n=n,
            model=model,
            max_end_s=max_end_s,
            rtol=rtol,
            atol=atol,
            max_step_before_s=max_step_before_s,
            max_step_after_s=max_step_after_s,
            chunk_before_s=chunk_before_s,
            chunk_after_s=chunk_after_s,
            h=h,
            hm=hm,
            case_name=name,
            geometry_label="fixed_R0" if name in {"A", "C"} else "attachment2_radius",
            radius_metadata=fixed_metadata if name in {"A", "C"} else radius_metadata,
            evidence_dir=None if evidence_dir is None else Path(evidence_dir) / f"case_{name}",
            run_id=f"case_{name}",
        )
    return results


def factorial_differences(results: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Compute the four planned effects only from certified event times."""
    event_times: dict[str, float | None] = {}
    for name in ("A", "B", "C", "D"):
        event = results.get(name, {}).get("event")
        event_times[name] = None if event is None or event.get("certification_status") != "PASS" else float(event["time_s"])
    if any(event_times[name] is None for name in ("A", "B", "C", "D")):
        return {"status": "NOT_COMPARABLE", "event_time_s": event_times, "reason": "one_or_more_cases_without_certified_event"}
    a, b, c, d = (float(event_times[name]) for name in ("A", "B", "C", "D"))
    return {
        "status": "CHECKED",
        "event_time_s": event_times,
        "delta_property_s": c - a,
        "delta_shrink_3_s": b - a,
        "delta_shrink_4_s": d - c,
        "delta_interaction_s": (d - c) - (b - a),
    }


def _build_cli() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Solve the Q4 moving-radius heat--moisture model and save development evidence."
    )
    parser.add_argument("--boundary", type=Path, default=DEFAULT_BOUNDARY)
    parser.add_argument("--radius", type=Path, default=DEFAULT_RADIUS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--n", type=int, default=40, help="number of normalized radial intervals")
    parser.add_argument("--model", choices=("appendix3", "appendix4"), default="appendix4")
    parser.add_argument("--max-end-s", type=float, default=MAX_END_S)
    parser.add_argument("--radius-interpolation", choices=("pchip", "linear"), default="pchip")
    parser.add_argument("--radius-post-mode", choices=("plateau", "slow1pct"), default="plateau")
    parser.add_argument("--all-cases", action="store_true", help="run the A--D factorial comparison")
    parser.add_argument("--rtol", type=float, default=BASE_RTL)
    parser.add_argument("--atol", type=float, default=BASE_ATOL)
    parser.add_argument("--max-step-before-s", type=float, default=BASE_MAX_STEP_BEFORE_S)
    parser.add_argument("--max-step-after-s", type=float, default=BASE_MAX_STEP_AFTER_S)
    parser.add_argument("--chunk-before-s", type=float, default=DEFAULT_CHUNK_S)
    parser.add_argument("--chunk-after-s", type=float, default=DEFAULT_CHUNK_S)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_cli().parse_args(argv)
    boundary, boundary_metadata = read_boundary(args.boundary)
    radius, radius_metadata = read_radius(
        args.radius,
        interpolation=args.radius_interpolation,
        post_mode=args.radius_post_mode,
    )
    if args.all_cases:
        results = run_four_cases(
            boundary,
            radius,
            n=args.n,
            max_end_s=args.max_end_s,
            rtol=args.rtol,
            atol=args.atol,
            max_step_before_s=args.max_step_before_s,
            max_step_after_s=args.max_step_after_s,
            chunk_before_s=args.chunk_before_s,
            chunk_after_s=args.chunk_after_s,
            radius_metadata=radius_metadata,
            evidence_dir=args.output / "evidence",
        )
        for name, result in results.items():
            case_dir = args.output / f"case_{name}"
            save_result_files(
                case_dir,
                result,
                boundary_metadata=boundary_metadata,
                radius_metadata=result.get("radius_metadata", radius_metadata),
            )
        write_json(args.output / "q4_four_cases_summary.json", {
            "cases": {name: case_summary(result) for name, result in results.items()},
            "factorial_differences": factorial_differences(results),
        })
        return 0
    result = run_case(
        boundary,
        radius,
        n=args.n,
        model=args.model,
        max_end_s=args.max_end_s,
        rtol=args.rtol,
        atol=args.atol,
        max_step_before_s=args.max_step_before_s,
        max_step_after_s=args.max_step_after_s,
        chunk_before_s=args.chunk_before_s,
        chunk_after_s=args.chunk_after_s,
        evidence_dir=args.output / "evidence",
        run_id=f"single_{args.model}",
    )
    save_result_files(
        args.output,
        result,
        boundary_metadata=boundary_metadata,
        radius_metadata=result.get("radius_metadata", radius_metadata),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
